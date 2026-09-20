"""Freeze REAL official documents into tests/fixtures/cb_docs/ (cut to what the parsers read) with a manifest url -> file.

    python scripts/cb_freeze_docs_fixture.py [--out tests/fixtures/cb_docs] [--append boj-followups|video-pages]

Fetched politely (src/cb_docs/http.py: robots.txt, rate limit). HTML pages are cut to their container (the extraction of the cut page is
asserted to equal the extraction of the full page), feeds to their newest items, the BoE workbook to the last two years of votes. Run once
per refresh of the fixtures; the tests never touch the network.
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from src.cb_docs import extract as X, sources as S            # noqa: E402
from src.cb_docs.http import Fetcher                            # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
MEETINGS = {
    "USD": ["2026-04-29", "2026-06-17", "2026-07-29", "2026-09-16"], "EUR": ["2026-04-30", "2026-06-11", "2026-07-23", "2026-09-10"],
    "GBP": ["2026-04-30", "2026-06-18", "2026-07-30", "2026-09-17"], "JPY": ["2026-04-28", "2026-06-16", "2026-07-31", "2026-09-18"],
    "CAD": ["2026-04-29", "2026-06-10", "2026-07-15", "2026-09-02"], "AUD": ["2026-03-17", "2026-05-05", "2026-06-16", "2026-08-11"],
    "CHF": ["2025-09-25", "2025-12-11", "2026-03-19", "2026-06-18"],
}
FEEDS = {
    "fed_speeches.xml": "https://www.federalreserve.gov/feeds/speeches.xml", "fed_testimony.xml": "https://www.federalreserve.gov/feeds/testimony.xml",
    "ecb_press.xml": "https://www.ecb.europa.eu/rss/press.html", "boe_speeches.xml": "https://www.bankofengland.co.uk/rss/speeches",
    "boc_speeches.xml": "https://www.bankofcanada.ca/content_type/speeches/feed/", "snb_speeches.xml": "https://www.snb.ch/public/rss/en/speeches",
    "bis_cbspeeches.xml": "https://www.bis.org/doclist/cbspeeches.rss", "boc_summary_feed.xml": "https://www.bankofcanada.ca/content_type/summary-of-deliberations/feed/",
}


def cut_container(html: str, rule: dict) -> str:
    """The outer HTML of the largest <tag> matching the rule's selector, balanced on the tag name."""
    tag, id_, cls = rule["sel"]
    best = ""
    for m in re.finditer(rf"<{tag}\b[^>]*>", html):
        head = m.group(0)
        if id_ and not re.search(rf"""\bid=["']{re.escape(id_)}["']""", head):
            continue
        if cls and not re.search(rf"""class=(["'])(?:[^"']*\s)?{re.escape(cls)}(?:\s[^"']*)?\1""", head):
            continue
        depth, i = 0, m.start()
        for t in re.finditer(rf"<(/?){tag}\b[^>]*>", html[m.start():]):
            depth += -1 if t.group(1) else 1
            if depth == 0:
                i = m.start() + t.end()
                break
        if i - m.start() > len(best):
            best = html[m.start():i]
    return "<!DOCTYPE html><html><head><meta charset=\"utf-8\"></head><body>" + best + "</body></html>"


def cut_feed(xml: str, n: int) -> str:
    items = re.findall(r"<item\b.*?</item>|<entry\b.*?</entry>", xml, flags=re.S)
    if len(items) <= n:
        return xml
    keep = items[:n]
    first = xml.find(items[0])
    last = xml.rfind(items[-1]) + len(items[-1])
    return xml[:first] + "\n".join(keep) + xml[last:]


def cut_list(html: str, pattern: str) -> str:
    links = re.findall(pattern, html, flags=re.S)
    return "<!DOCTYPE html><html><body><ul>\n" + "\n".join(f"<li>{l}</li>" for l in links) + "\n</ul></body></html>"


def cut_workbook(content: bytes, since: date) -> bytes:
    import openpyxl
    src = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)["Bank Rate Decisions"]
    rows = list(src.iter_rows(values_only=True))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bank Rate Decisions"
    hdr = next(i for i, r in enumerate(rows) if r and r[2] == "Current members")
    for r in rows[:hdr + 1]:
        ws.append([c for c in r[:30]])
    for r in rows[hdr + 1:]:
        if r and isinstance(r[1], datetime) and r[1].date() >= since:
            ws.append([c for c in r[:30]])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def first_pages(pdf: bytes, n: int = 1) -> bytes:
    """The first n pages of a real PDF, re-written (the long BoJ documents only need to prove that the text is extracted and hashed)."""
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    for page in PdfReader(io.BytesIO(pdf)).pages[:n]:
        w.add_page(page)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def cut_speech_page(html: str) -> str:
    """<title> + the first paragraphs of a bank speech page: what the collector reads (byline in the title, first paragraph)."""
    title = re.search(r"<title>.*?</title>", html, flags=re.S)
    body = re.findall(r"<p\b[^>]*>.*?</p>", html, flags=re.S)
    long = [p for p in body if len(re.sub(r"<[^>]+>", "", p)) > 80][:3]
    return f"<!DOCTYPE html><html><head><meta charset=\"utf-8\">{title.group(0) if title else ''}</head><body><main>\n" + "\n".join(long) + "\n</main></body></html>"


def append_boj_followups(out: Path, f: Fetcher) -> int:
    """JPY summary of opinions / minutes (first PDF page) for the meetings whose files exist, and the pages of the listed BoJ speeches."""
    manifest = json.loads((out / "manifest.json").read_text())
    for iso in MEETINGS["JPY"][:3]:
        for typ, url, due, rule in S.minutes_urls("JPY", date.fromisoformat(iso)):
            r = f.get(url, conditional=False)
            if not r.ok:
                print(f"skip {typ} {iso}: {r.error}")
                continue
            rel = f"minutes/JPY_{typ}_{iso}.pdf"
            (out / rel).write_bytes(first_pages(r.content))
            manifest[url] = {"file": rel, "content_type": "application/pdf"}
    listing = (out / "lists" / "boj_koen.html").read_text()
    for path in re.findall(r'href="(/en/about/press/koen_2026/ko\d{6}a\.htm)"', listing)[:3]:
        url = "https://www.boj.or.jp" + path
        r = f.get(url, conditional=False)
        if r.ok:
            rel = f"lists/boj_{path.rsplit('/', 1)[-1]}"
            (out / rel).write_text(cut_speech_page(r.text))
            manifest[url] = {"file": rel, "content_type": "text/html; charset=utf-8"}
    (out / "manifest.json").write_text(json.dumps(manifest, indent=1, sort_keys=True) + "\n")
    print(f"{len(manifest)} files in the manifest, {f.requests_made} requests")
    return 0


def append_video_pages(out: Path, f: Fetcher) -> int:
    """The pages that carry the press-conference video, cut to the element the parser reads (a page without the video keeps only its <title>)."""
    manifest = json.loads((out / "manifest.json").read_text())

    def keep(url: str, rel: str, html: str) -> None:
        (out / rel).write_text(html)
        manifest[url] = {"file": rel, "content_type": "text/html; charset=utf-8"}

    def shell(inner: str, title: str = "") -> str:
        return f"<!DOCTYPE html><html><head><meta charset=\"utf-8\">{title}</head><body>\n{inner}\n</body></html>"

    for ccy in ("USD", "CAD", "AUD", "GBP"):
        for iso in MEETINGS[ccy]:
            d = date.fromisoformat(iso)
            url = S.video_page(ccy, d)
            r = f.get(url, conditional=False)
            if not r.ok:
                print(f"skip {ccy} {iso}: {r.error}")
                continue
            title = (re.search(r"<title>.*?</title>", r.text, flags=re.S) or [""])[0]
            if ccy == "USD":
                inner = "\n".join(re.findall(r"<[^>]*data-video-id=\"\d+\"[^>]*>", r.text)[:1])
            elif ccy == "CAD":
                inner = "\n".join(re.findall(r'<a [^>]*href="https://www\.bankofcanada\.ca/multimedia/press-conference[^"]+"[^>]*>.*?</a>', r.text, flags=re.S)[:1])
            elif ccy == "AUD":
                inner = "\n".join(re.findall(r'<a [^>]*class="[^"]*video-placeholder[^"]*"[^>]*>.*?</a>', r.text, flags=re.S)[:1])
            else:
                i = r.text.find('id="press-conference"')
                inner = r.text[max(0, i - 120):i + 900] if i >= 0 else ""
            if ccy == "CAD":                                        # same URL as the statement: the link is added to that fixture, not a second file
                rel = f"statements/CAD_{iso}.html"
                page = (out / rel).read_text()
                page = re.sub(r"<a [^>]*multimedia/press-conference[^>]*>.*?</a>\n?", "", page, flags=re.S)
                (out / rel).write_text(page.replace("</body>", inner + "\n</body>", 1))
                manifest[url] = {"file": rel, "content_type": "text/html; charset=utf-8"}
                continue
            keep(url, f"lists/video_{ccy}_{iso}.html", shell(inner, title))
    r = f.get(S.ECB_PRESS_LANDING, conditional=False)
    if r.ok:
        box = re.search(r'<div class="jumbo-box" id="youtube">.*?<div data-video="[^"]+"></div>', r.text, flags=re.S)
        link = re.search(r'<a [^>]*ecb\.is\d{6}~[^>]*>', r.text)
        keep(S.ECB_PRESS_LANDING, "lists/video_EUR_landing.html", shell((link.group(0) + "</a>\n" if link else "") + (box.group(0) if box else "")))
    (out / "manifest.json").write_text(json.dumps(manifest, indent=1, sort_keys=True) + "\n")
    print(f"{len(manifest)} files in the manifest, {f.requests_made} requests")
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", type=Path, default=ROOT / "tests" / "fixtures" / "cb_docs")
    ap.add_argument("--append", choices=["boj-followups", "video-pages"], help="add these fixtures to an existing set instead of re-freezing everything")
    a = ap.parse_args(argv)
    out = a.out
    if a.append == "boj-followups":
        return append_boj_followups(out, Fetcher(min_interval=1.2))
    if a.append == "video-pages":
        return append_video_pages(out, Fetcher(min_interval=1.2))
    for sub in ("statements", "minutes", "feeds", "lists", "misc"):
        (out / sub).mkdir(parents=True, exist_ok=True)
    f = Fetcher(min_interval=1.2)
    manifest: dict = {}

    def keep(url: str, rel: str, data: bytes, ctype: str) -> None:
        (out / rel).write_bytes(data if isinstance(data, bytes) else data.encode())
        manifest[url] = {"file": rel, "content_type": ctype}

    # links that cannot be derived: RBA media releases from the decisions page, ECB statements from the seeds
    rba = f.get(f"https://www.rba.gov.au/monetary-policy/int-rate-decisions/{date.today().year}/", conditional=False)
    rba_links = {}
    for li in re.findall(r"<li[^>]*>(.*?)</li>", rba.text, flags=re.S):
        m = re.search(r"href=\"(/media-releases/[^\"]+)\"[^>]*>\s*(\d{1,2} [A-Z][a-z]+ \d{4})", li, flags=re.S)
        if m:
            rba_links[datetime.strptime(m.group(2), "%d %B %Y").date()] = "https://www.rba.gov.au" + m.group(1)
    keep(rba.url, "lists/rba_decisions.html", cut_list(rba.text, r'(<a[^>]+href="/media-releases/[^"]+"[^>]*>\s*\d{1,2} [A-Z][a-z]+ \d{4}.*?</a>)'), "text/html")
    for ccy, days in MEETINGS.items():
        for iso_day in days:
            d = date.fromisoformat(iso_day)
            url = S.statement_url(ccy, d, {"AUD": rba_links})
            r = f.get(url, conditional=False)
            assert r.ok, (ccy, d, r.error)
            if url.endswith(".pdf"):
                keep(url, f"statements/{ccy}_{iso_day}.pdf", r.content, "application/pdf")
                continue
            rule = X.STATEMENT_RULES[ccy]
            cut = cut_container(r.text, rule)
            assert X.html_paragraphs(cut, rule) == X.html_paragraphs(r.text, rule), (ccy, d)
            keep(url, f"statements/{ccy}_{iso_day}.html", cut.encode(), "text/html; charset=utf-8")
    for d in MEETINGS["AUD"]:                                              # RBA minutes (votes)
        for typ, url, due, rule in S.minutes_urls("AUD", date.fromisoformat(d)):
            r = f.get(url, conditional=False)
            assert r.ok, url
            keep(url, f"minutes/AUD_{d}.html", cut_container(r.text, rule).encode(), "text/html; charset=utf-8")
    for d in ("2026-04-29", "2026-06-17"):                                 # Fed minutes (one long HTML, cut)
        for typ, url, due, rule in S.minutes_urls("USD", date.fromisoformat(d)):
            r = f.get(url, conditional=False)
            if r.ok:
                keep(url, f"minutes/USD_{d}.html", cut_container(r.text, rule).encode(), "text/html; charset=utf-8")
    for ccy in ("CAD", "CHF"):                                             # introductory statements (short: their text is committed)
        for iso_day in MEETINGS[ccy]:
            for typ, url, mode, rule in S.transcript_targets(ccy, date.fromisoformat(iso_day)):
                r = f.get(url, conditional=False)
                if r.ok and rule:
                    keep(url, f"statements/{ccy}_{iso_day}_opening.html", cut_container(r.text, rule).encode(), "text/html; charset=utf-8")
    for name, url in FEEDS.items():
        r = f.get(url, conditional=False)
        assert r.ok, url
        keep(url, f"feeds/{name}", cut_feed(r.text, 14).encode(), "application/xml")
    r = f.get(S.HTML_LISTS["AUD"], conditional=False)
    keep(S.HTML_LISTS["AUD"], "lists/rba_speeches.html", cut_list(r.text, r'(<a[^>]+href="/speeches/\d{4}/[a-z\-]+-\d{4}-\d{2}-\d{2}[^"]*\.html"[^>]*>.*?</a>)').encode(), "text/html")
    r = f.get(S.HTML_LISTS["JPY"].format(year=2026), conditional=False)
    keep(S.HTML_LISTS["JPY"].format(year=2026), "lists/boj_koen.html", cut_list(r.text, r'(<a[^>]+href="/en/about/press/koen_\d{4}/ko\d{6}a\.htm"[^>]*>.*?</a>)').encode(), "text/html")
    r = f.get(S.BOE_VOTING_XLSX, conditional=False)
    keep(S.BOE_VOTING_XLSX, "misc/boe_mpcvoting.xlsx", cut_workbook(r.content, date(2025, 1, 1)), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    (out / "manifest.json").write_text(json.dumps(manifest, indent=1, sort_keys=True) + "\n")
    print(f"{len(manifest)} files -> {out} ({sum(p.stat().st_size for p in out.rglob('*') if p.is_file()) // 1024} KB), {f.requests_made} requests")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
