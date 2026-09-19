"""Market-data adapters for the Central Banks collector (phase 1A).

Extracted from the diagnostic probes in `src/cb_probe.py`; that file stays as the spike record.
Every adapter follows the contract in `base.py`: values are stored exactly as published (price or
rate, with a unit), the reference window is explicit (`ref_start`, `ref_end` exclusive) or the row is
a tenor (`tenor_months`), the exchange's naming convention is resolved HERE, and `fetch` returns a
`FetchResult` or None with `last_status` set - it never raises. Price->rate conversion, spread
adjustment and meeting extraction are phase 1B.

Each adapter splits I/O (`_fetch`) from a pure `parse*` method so fixtures give deterministic rows.
"""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from datetime import date, datetime, timedelta
from typing import Optional
from urllib.parse import urljoin

from ..rate_sources import _safe_float
from .base import (FetchResult, MarketSource, NON_BROWSER_UA, Quote, add_months, first_wednesday_after_9th,
                   imm_window, month_bounds, third_wednesday)

MONTHS = {m: i for i, m in enumerate(
    ["january", "february", "march", "april", "may", "june", "july", "august", "september", "october",
     "november", "december"], start=1)}
FUTURES_CODE = dict(zip("FGHJKMNQUVXZ", range(1, 13)))       # CME/ASX/MX month letters


def tenor_label(months: float) -> str:
    return f"{months:g}M"


def contract_month(y: int, m: int) -> str:
    return f"{y}{m:02d}"


def _as_date(v) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and re.match(r"\d{4}-\d{2}-\d{2}", v):
        return date.fromisoformat(v[:10])
    return None


def _sheet_rows(content: bytes, sheet: str):
    """Row iterator over an xlsx sheet. `reset_dimensions` because some publishers (Atlanta Fed) write a
    stale <dimension> that makes read-only mode stop after the first cell."""
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(content), data_only=True, read_only=True)[sheet]
    ws.reset_dimensions()
    return ws.iter_rows(values_only=True)


def _norm_months(x: float) -> float:
    r = round(float(x))
    return float(r) if abs(float(x) - r) < 1e-3 else round(float(x), 3)


# ---------------------------------------------------------------------------
# Official-history sources
# ---------------------------------------------------------------------------

class AtlantaMpt(MarketSource):
    """Atlanta Fed Market Probability Tracker: mean of the SOFR-option-implied 3M SOFR distribution (bp)
    for each 3M reference quarter (explicit `reference_start` = 3rd Wednesday of the start month)."""
    id = "atlantafed_mpt"
    currency = "USD"
    INSTRUMENT = "sofr_3m_ref_quarter_mean"
    UNIT = "bp"
    FIELD_IN = "Rate: mean"

    def parse(self, content: bytes, since: date) -> list[Quote]:
        it = _sheet_rows(content, "DATA")
        ix = {str(h).strip(): i for i, h in enumerate(next(it)) if h is not None}
        miss = [c for c in ("date", "reference_start", "field", "value") if c not in ix]
        if miss:
            raise ValueError(f"MPT DATA sheet lacks columns {miss}")
        out: list[Quote] = []
        for r in it:
            if r[ix["field"]] != self.FIELD_IN:
                continue
            d, rs, v = _as_date(r[ix["date"]]), _as_date(r[ix["reference_start"]]), _safe_float(r[ix["value"]])
            if d is None or rs is None or v is None or d < since or not self.within_horizon(rs, d):
                continue
            y2, m2 = add_months(rs.year, rs.month, 3)
            out.append(self.quote(self.INSTRUMENT, contract_month(rs.year, rs.month), "mean", v, self.UNIT, d,
                                  ref_start=rs, ref_end=third_wednesday(y2, m2), tenor_months=3.0))
        out.sort(key=lambda q: (q.asof, q.ref_start))
        return out

    def _fetch(self, since, state):
        r, not_modified, v = self.get_url(self.cfg["url"], state)
        if not_modified:
            return FetchResult("not_modified", state=v, note="HTTP 304")
        if r is None:
            return None
        qs = self.parse(r.content, since)
        return FetchResult("ok", qs, state=v, note="" if qs else f"no rows since {since}")


class BoeOis(MarketSource):
    """Bank of England OIS instantaneous forward curve (percent), monthly grid; published by noon on d+1."""
    id = "boe_ois"
    currency = "GBP"
    INSTRUMENT = "sonia_instantaneous_forward"
    UNIT = "percent"
    SHEET = "1. fwds, short end"

    def parse(self, blob: bytes, member_rx: str, since: date) -> dict[date, list[tuple[float, float]]]:
        z = zipfile.ZipFile(io.BytesIO(blob))
        name = next((n for n in z.namelist() if re.search(member_rx, n)), None)
        if name is None:
            raise ValueError(f"no workbook matching {member_rx!r} in {z.namelist()}")
        rows = list(_sheet_rows(z.read(name), self.SHEET))
        mrow = next((r for r in rows if r and str(r[0]).strip().lower().startswith("months")), None)
        if mrow is None:
            raise ValueError("no 'months:' row")
        months = [_norm_months(x) for x in mrow[1:] if x is not None]
        curves: dict[date, list[tuple[float, float]]] = {}
        for r in rows:
            d = _as_date(r[0])
            if d is None or d < since:
                continue
            vals = r[1:1 + len(months)]
            if vals and all(isinstance(x, (int, float)) for x in vals):
                curves[d] = [(m, float(x)) for m, x in zip(months, vals) if m <= self.horizon_months]
        return curves

    def quotes(self, curves: dict) -> list[Quote]:
        return [self.quote(self.INSTRUMENT, tenor_label(m), "fwd", v, self.UNIT, d, tenor_months=m)
                for d, pts in sorted(curves.items()) for m, v in pts]

    def _fetch(self, since, state):
        st = {k: dict(v) for k, v in state.items() if isinstance(v, dict)}
        r1, nm1, v1 = self.get_url(self.cfg["urls"]["latest"], st.get("latest"))
        if nm1:
            return FetchResult("not_modified", state=st, note="HTTP 304")
        if r1 is None:
            return None
        everything = self.parse(r1.content, r"OIS daily data current month", date.min)
        curves = {d: v for d, v in everything.items() if d >= since}
        st["latest"] = v1
        note = ""
        # the current-month workbook covers every business day of ITS month; earlier days need the history workbook
        month_start = date(max(everything).year, max(everything).month, 1) if everything else None
        if month_start is None or since < month_start:
            r2, nm2, v2 = self.get_url(self.cfg["urls"]["history"], st.get("history"))
            if r2 is not None:
                older = self.parse(r2.content, r"2025 to present", since)
                curves = {**older, **curves}
                st["history"] = v2
            elif not nm2:
                note = f"history workbook unavailable ({self.last_status} {self.last_note})".strip()
        return FetchResult("ok", self.quotes(curves), state=st, note=note)


class TreasuryBills(MarketSource):
    """US Treasury daily par yield curve, bill tenors (percent). PROXY: bills, not OIS."""
    id = "ust_bills"
    currency = "USD"
    INSTRUMENT = "ust_par_curve"
    UNIT = "percent"
    TENORS = {"1 Mo": 1, "1.5 Month": 1.5, "2 Mo": 2, "3 Mo": 3, "4 Mo": 4, "6 Mo": 6,
              "1 Yr": 12, "2 Yr": 24, "3 Yr": 36}

    def parse(self, text: str, since: date) -> list[Quote]:
        out: list[Quote] = []
        for row in csv.DictReader(io.StringIO(text.lstrip("﻿"))):
            try:
                d = datetime.strptime(row["Date"], "%m/%d/%Y").date()
            except (KeyError, ValueError):
                continue
            if d < since:
                continue
            for col, m in self.TENORS.items():
                v = _safe_float(row.get(col))
                if v is not None and m <= self.horizon_months:
                    out.append(self.quote(self.INSTRUMENT, tenor_label(m), "yield", v, self.UNIT, d,
                                          tenor_months=float(m)))
        return out

    def _fetch(self, since, state):
        out: list[Quote] = []
        for yr in range(since.year, self._now().year + 1):
            r, _, _ = self.get_url(self.cfg["url"].format(year=yr))
            if r is None or self._check_botwall(r):
                if yr == self._now().year:
                    return None
                continue
            out += self.parse(r.text, since)
        return FetchResult("ok", out, note="" if out else f"no rows since {since}")


class BocTbills(MarketSource):
    """Bank of Canada Valet: Government of Canada T-bill mid yields 1/2/3/6M (percent). PROXY."""
    id = "boc_tbills"
    currency = "CAD"
    INSTRUMENT = "gov_tbill_yield"
    UNIT = "percent"
    SERIES = {"TB.CDN.30D.MID": 1, "TB.CDN.60D.MID": 2, "TB.CDN.90D.MID": 3, "TB.CDN.180D.MID": 6}

    def parse(self, payload: dict, since: date) -> list[Quote]:
        out: list[Quote] = []
        for o in payload.get("observations") or []:
            d = _as_date(o.get("d"))
            if d is None or d < since:
                continue
            for sid, m in self.SERIES.items():
                v = _safe_float((o.get(sid) or {}).get("v"))
                if v is not None:
                    out.append(self.quote(self.INSTRUMENT, tenor_label(m), "yield", v, self.UNIT, d,
                                          tenor_months=float(m)))
        return sorted(out, key=lambda q: (q.asof, q.tenor_months))

    def _fetch(self, since, state):
        r, _, _ = self.get_url(f"{self.cfg['url']}?start_date={since.isoformat()}")
        if r is None:
            return None
        qs = self.parse(r.json(), since)
        return FetchResult("ok", qs, note="" if qs else f"no rows since {since}")


class EcbAaaForward(MarketSource):
    """ECB euro-area AAA government instantaneous forwards IF_* at monthly tenors 3..36M (percent). PROXY:
    government curve, not the EUR STR OIS path."""
    id = "ecb_aaa_fwd"
    currency = "EUR"
    INSTRUMENT = "aaa_govt_instantaneous_forward"
    UNIT = "percent"
    TENOR_RX = re.compile(r"IF_(?:(\d+)Y)?(?:(\d+)M)?$")

    @classmethod
    def tenor_months(cls, code: str) -> Optional[int]:
        m = cls.TENOR_RX.match(code)
        return (int(m.group(1) or 0) * 12 + int(m.group(2) or 0)) if m and (m.group(1) or m.group(2)) else None

    def keys(self) -> list[str]:
        out = []
        for m in range(3, self.horizon_months + 1):
            y, mo = divmod(m, 12)
            out.append("IF_" + (f"{y}Y" if y else "") + (f"{mo}M" if mo else ""))
        return out

    def parse(self, text: str, since: date) -> list[Quote]:
        out: list[Quote] = []
        for row in csv.DictReader(io.StringIO(text.lstrip("﻿"))):
            m = self.tenor_months(row["KEY"].rsplit(".", 1)[-1])
            d, v = _as_date(row.get("TIME_PERIOD")), _safe_float(row.get("OBS_VALUE"))
            if m is None or d is None or v is None or d < since or m > self.horizon_months:
                continue
            out.append(self.quote(self.INSTRUMENT, tenor_label(m), "fwd", v, self.UNIT, d, tenor_months=float(m)))
        return sorted(out, key=lambda q: (q.asof, q.tenor_months))

    def _fetch(self, since, state):
        url = self.cfg["url"].format(keys="+".join(self.keys())) + f"?startPeriod={since.isoformat()}&format=csvdata"
        r, _, _ = self.get_url(url)
        if r is None or self._check_botwall(r):
            return None
        qs = self.parse(r.text, since)
        return FetchResult("ok", qs, note="" if qs else f"no rows since {since}")


class RbaBankBills(MarketSource):
    """RBA statistical table F1: EOD bank-accepted bills/NCDs 1/3/6M (percent, source ASX). PROXY."""
    id = "rba_bank_bills"
    currency = "AUD"
    ua = NON_BROWSER_UA                      # a Chrome UA earns a 403 here
    INSTRUMENT = "bank_bill_eod_yield"
    UNIT = "percent"
    SERIES = {"FIRMMBAB30D": 1, "FIRMMBAB90D": 3, "FIRMMBAB180D": 6}

    def parse(self, text: str, since: date) -> list[Quote]:
        rows = list(csv.reader(io.StringIO(text.lstrip("﻿"))))
        sid_i = next((i for i, x in enumerate(rows) if x and x[0] == "Series ID"), None)
        if sid_i is None:
            raise ValueError("no 'Series ID' row in F1")
        col = {sid: j for j, sid in enumerate(rows[sid_i]) if sid}
        miss = [s for s in self.SERIES if s not in col]
        if miss:
            raise ValueError(f"F1 ids missing: {miss}")
        out: list[Quote] = []
        for x in rows[sid_i + 1:]:
            if not x or not re.match(r"\d{2}-[A-Za-z]{3}-\d{4}$", x[0]):
                continue
            d = datetime.strptime(x[0], "%d-%b-%Y").date()
            if d < since:
                continue
            for sid, m in self.SERIES.items():
                v = _safe_float(x[col[sid]]) if len(x) > col[sid] else None
                if v is not None:
                    out.append(self.quote(self.INSTRUMENT, tenor_label(m), "yield", v, self.UNIT, d,
                                          tenor_months=float(m)))
        return out

    def _fetch(self, since, state):
        r, not_modified, v = self.get_url(self.cfg["url"], state)
        if not_modified:
            return FetchResult("not_modified", state=v, note="HTTP 304")
        if r is None or self._check_botwall(r):
            return None
        qs = self.parse(r.text, since)
        return FetchResult("ok", qs, state=v, note="" if qs else f"no rows since {since}")


# ---------------------------------------------------------------------------
# Snapshot sources: only the latest day is exposed => history exists from the first collector run
# ---------------------------------------------------------------------------

class JpxTona(MarketSource):
    """JPX daily settlement-price CSV (cp932, latest business day only); 3M TONA futures rows.
    Contract month = START month of the reference period (E2): 202612 -> 2026-12-16 .. 2027-03-17."""
    id = "jpx_tona"
    currency = "JPY"
    INSTRUMENT = "tona_3m_futures"
    UNIT = "index_points"
    LINK_RX = re.compile(r'href="([^"]*/rb_e(\d{8})\.csv)"')
    ROW_RX = re.compile(r"\d+,FUT_TOA3M_\d{6},,(\d{6}),,([\d.]+)")

    def find_csv(self, page_html: str) -> Optional[tuple[str, date]]:
        m = self.LINK_RX.search(page_html)
        return (m.group(1), datetime.strptime(m.group(2), "%Y%m%d").date()) if m else None

    def parse(self, text: str, asof: date) -> tuple[list[Quote], str]:
        out: list[Quote] = []
        raw: list[str] = []
        for ln in text.splitlines():
            m = self.ROW_RX.match(ln)
            if not m:
                continue
            y, mo = int(m.group(1)[:4]), int(m.group(1)[4:])
            rs, re_ = imm_window(y, mo)
            if not self.within_horizon(rs, asof):
                continue
            raw.append(ln)
            out.append(self.quote(self.INSTRUMENT, m.group(1), "settle", float(m.group(2)), self.UNIT, asof,
                                  ref_start=rs, ref_end=re_, tenor_months=3.0))
        return out, "\n".join(raw)

    def _fetch(self, since, state):
        page, _, _ = self.get_url(self.cfg["url"])
        if page is None:
            return None
        found = self.find_csv(page.text)
        if not found:
            self._parse_fail("no rb_e*.csv link (JS-only page?)")
            return None
        href, asof = found
        r, not_modified, v = self.get_url(urljoin(self.cfg["url"], href), state)
        if not_modified:
            return FetchResult("not_modified", state=v, note="HTTP 304")
        if r is None:
            return None
        qs, raw = self.parse(r.content.decode("cp932", errors="replace"), asof)
        if not qs:
            self._parse_fail("no FUT_TOA3M rows")
            return None
        return FetchResult("ok", qs, raw={"asof": asof, "text": raw}, state=v)


class MxCorra(MarketSource):
    """Montreal Exchange 'Canadian Interest Rate Expectations': CRA (3M CORRA, contract month = START month of
    the reference quarter) and COA (1M CORRA, calendar month) settlement prices. The page has NO as-of stamp and
    no validators: as-of is inferred from fetch time + cutoff (asof_inferred) and fetches inside the intraday
    window are skipped."""
    id = "mx_corra"
    currency = "CAD"
    ROW_RX = re.compile(r'<td class="text-left">([A-Za-z]+) (\d{4}) \((CRA|COA)([A-Z])(\d{2})\)</td>\s*<td>([\d.]+)</td>')
    INSTR = {"CRA": ("corra_3m_futures", 3.0), "COA": ("corra_1m_futures", 1.0)}
    UNIT = "index_points"

    def parse(self, html: str, asof: date, inferred: bool = True) -> tuple[list[Quote], str]:
        seen: set = set()
        out: list[Quote] = []
        raw: list[str] = []
        for label_m, yr, fam, letter, yy, px in self.ROW_RX.findall(html):
            mo = MONTHS.get(label_m.lower())
            y = int(yr)
            if mo is None or FUTURES_CODE.get(letter) != mo or int(yy) != y % 100:
                raise ValueError(f"MX row label/code mismatch: {label_m} {yr} ({fam}{letter}{yy})")
            if (fam, y, mo) in seen:
                continue
            seen.add((fam, y, mo))
            rs, re_ = imm_window(y, mo) if fam == "CRA" else month_bounds(y, mo)
            if not self.within_horizon(rs, asof):
                continue
            instr, tenor = self.INSTR[fam]
            raw.append(f"{label_m} {yr},{fam}{letter}{yy},{px}")
            out.append(self.quote(instr, contract_month(y, mo), "settle", float(px), self.UNIT, asof,
                                  ref_start=rs, ref_end=re_, tenor_months=tenor, inferred=inferred))
        return out, "\n".join(raw)

    def _fetch(self, since, state):
        asof = self.infer_asof()
        if asof is None:
            return FetchResult("skipped", note="inside intraday window (values may be live)")
        r, _, _ = self.get_url(self.cfg["url"])
        if r is None:
            return None
        qs, raw = self.parse(r.text, asof)
        if not any(q.instrument == "corra_3m_futures" for q in qs):
            self._parse_fail("no CRA rows - layout changed?")
            return None
        return FetchResult("ok", qs, raw={"asof": asof, "text": raw})


class AsxFutures(MarketSource):
    """ASX futures via the JSON API behind asx.com.au (undocumented). Every item carries the settlement date of
    its price (`datePreviousSettlement`)."""
    CODE = ""
    INSTRUMENT = ""
    UNIT = "index_points"
    TENOR = 1.0

    def window(self, y: int, m: int) -> tuple[date, date]:
        raise NotImplementedError

    def parse(self, payload: dict) -> tuple[list[Quote], str]:
        out: list[Quote] = []
        raw: list[str] = []
        sym_rx = re.compile(rf"{self.CODE}([FGHJKMNQUVXZ])(\d{{4}})$")
        for it in (payload.get("data") or {}).get("items") or []:
            m = sym_rx.match(str(it.get("symbol", "")))
            px, sd = _safe_float(it.get("pricePreviousSettlement")), _as_date(it.get("datePreviousSettlement"))
            if not m or px is None or sd is None:
                continue
            y, mo = int(m.group(2)), FUTURES_CODE[m.group(1)]
            rs, re_ = self.window(y, mo)
            if not self.within_horizon(rs, sd):
                continue
            raw.append(f"{it['symbol']},{it.get('dateExpiry')},{px},{sd}")
            out.append(self.quote(self.INSTRUMENT, contract_month(y, mo), "settle", px, self.UNIT, sd,
                                  ref_start=rs, ref_end=re_, tenor_months=self.TENOR))
        return out, "\n".join(raw)

    def _fetch(self, since, state):
        r, _, _ = self.get_url(self.cfg["url"])
        if r is None:
            return None
        qs, raw = self.parse(r.json())
        if not qs:
            self._parse_fail("no settled contracts")
            return None
        return FetchResult("ok", qs, raw={"asof": max(q.asof for q in qs), "text": raw})


class AsxIb(AsxFutures):
    """30-day interbank cash rate futures (AUD): contract = calendar month, window = that month."""
    id = "asx_ib"
    currency = "AUD"
    CODE = "IB"
    INSTRUMENT = "ib_30d_interbank_futures"
    TENOR = 1.0

    def window(self, y, m):
        return month_bounds(y, m)


class AsxBb(AsxFutures):
    """90-day NZ bank bill futures (NZD): contract = SETTLEMENT month; expiry = first Wednesday after the 9th;
    the underlying is the 90-day BKBM forward from expiry (window = expiry .. +91 days). The API's dateExpiry
    (2 days earlier) is not used."""
    id = "asx_bb"
    currency = "NZD"
    CODE = "BB"
    INSTRUMENT = "bb_90d_nz_bank_bill_futures"
    TENOR = 3.0

    def window(self, y, m):
        exp = first_wednesday_after_9th(y, m)
        return exp, exp + timedelta(days=91)


ADAPTERS: dict[str, type[MarketSource]] = {c.id: c for c in (
    AtlantaMpt, BoeOis, JpxTona, MxCorra, AsxIb, AsxBb, TreasuryBills, BocTbills, EcbAaaForward, RbaBankBills)}
