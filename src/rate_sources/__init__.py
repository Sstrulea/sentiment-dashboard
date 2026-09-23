"""Keyless sovereign-yield source adapters (~2y) for the Rate-Expectations engine.

This is the validated source layer extracted from the spike. Each adapter
implements the `RateSource` contract and returns a `YieldSeries` (or None,
never raising). `ALL_SOURCES` + `SOURCE_PREFERENCE` + `assess()` are the public
surface reused by both `src.rates_probe` (diagnostics) and `src.rate_fetch`
(parquet build). No I/O of project data here — only outbound keyless HTTP.

Adapter status (verified): USD→FRED DGS2 (DBnomics-FED SVENY02 fallback),
EUR→ECB SR_2Y, GBP→BoE GLC nominal spot 2.0y, JPY→MoF jgbcm_all.csv + jgbcm.csv (col 2年),
CAD→BoC BD.CDN.2YR.DQ.YLD, AUD→RBA F2 (Stooq fallback), NZD→RBNZ B2 xlsx
(Stooq fallback), CHF→SNB rendoblid D0=2J (often stale → Stooq fallback).

Known FX foreign-2y gaps (as of 2026-06, NOT trivially fixable):
  JPY  — MoF resolves but lags (stale).
  AUD  — RBA F2 403s behind a bot wall despite a browser UA + Referer + Accept
         (already set, see RbaSource); not a header tweak. Runs on Stooq cache.
  NZD  — RBNZ B2 xlsx 403s the same way (browser UA + Referer already set).
  CHF  — SNB stale; no live primary.
  Stooq fallback is itself bot-walled intermittently.
These are real bot walls, not the IPv6 routing issue fixed above (that affected
connection-level hangs, e.g. FRED — these return an HTTP 403 instead). Left as-is;
the cross-asset scoring already excludes stale/missing sub-scores from averages.
"""
from __future__ import annotations

import io
import logging
import re
import socket
import subprocess
import time
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Optional, Protocol, runtime_checkable

import numpy as np
import pandas as pd
import requests
import urllib3.util.connection

log = logging.getLogger(__name__)

# Force IPv4 for all outbound fetches. FRED (and some other endpoints) advertise
# AAAA records whose IPv6 route is broken from many networks: `requests` picks
# the IPv6 address and hangs into a ConnectionError after 20–40s, while `curl`
# silently falls back via Happy Eyeballs and succeeds. This is the real cause of
# WALCL/TGA/RRP (net liquidity) and DFII10 (real yield) failing from Python while
# curl works. Pinning the resolver to AF_INET makes requests behave like curl.
# Global to the fetch process — safe, no source we use is IPv6-only.
urllib3.util.connection.allowed_gai_family = lambda: socket.AF_INET

CURRENCIES = ["USD", "EUR", "GBP", "JPY", "AUD", "NZD", "CAD", "CHF"]

HTTP_TIMEOUT = 15
RETRIES = 2
BACKOFF_S = 1.0
UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
)
# FRED (fredgraph.csv) TARPITS browser User-Agents — a Chrome UA hangs/ReadTimeouts
# while a simple non-browser UA is served instantly. Used ONLY for the FRED CSV
# sources below (FredSource, FredSeriesSource); non-FRED sources keep `UA` (some,
# e.g. Stooq/RBA/BoE, need a browser UA to clear their own bot-walls).
FRED_UA = "macro-data-analysis/1.0 (+https://github.com/Sstrulea/macro-data-analysis)"

# Qualification thresholds.
FRESH_LAG_BD = 5     # latest within ~5 business days = fresh
STALE_LAG_BD = 10    # > 10 business days = STALE / disqualified for the pillar
MIN_HISTORY_YEARS = 2.0


class TenorMismatch(ValueError):
    """The source's OWN metadata (series id, maturity header, label) does not
    say 2 years. Never swallowed by BaseSource.fetch and never masked by a
    fallback source: rate_fetch logs it and leaves the currency unwritten, so
    freshness.rates turns it stale (audit 2026-09-23, 1.5)."""


def require_2y(source: str, ok: bool, what: str) -> None:
    if not ok:
        raise TenorMismatch(f"{source}: source metadata is not 2y ({what})")


# ---------------------------------------------------------------------------
# YieldSeries
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class YieldSeries:
    currency: str
    source: str
    tenor: str            # "2y" | "1y" | "3m" | "10y" | "n/a"
    frequency: str        # "daily" | "weekly" | "monthly"
    points: list[tuple[date, float]]  # ascending by date, yield in percent
    note: str = ""

    @property
    def n_points(self) -> int:
        return len(self.points)

    @property
    def latest_date(self) -> Optional[date]:
        return self.points[-1][0] if self.points else None

    @property
    def latest_value(self) -> Optional[float]:
        return self.points[-1][1] if self.points else None

    @property
    def history_start(self) -> Optional[date]:
        return self.points[0][0] if self.points else None

    @property
    def history_years(self) -> float:
        if not self.points:
            return 0.0
        return (self.points[-1][0] - self.points[0][0]).days / 365.25

    def business_days_lag(self, today: date) -> Optional[int]:
        if not self.latest_date:
            return None
        if self.latest_date >= today:
            return 0
        return int(np.busday_count(self.latest_date, today))


# ---------------------------------------------------------------------------
# Source protocol + base
# ---------------------------------------------------------------------------

@runtime_checkable
class RateSource(Protocol):
    name: str
    def supports(self, currency: str) -> bool: ...
    def fetch(self, currency: str) -> Optional[YieldSeries]: ...


class BaseSource:
    """Shared HTTP + bot-wall handling. Subclasses set `name` and `_fetch`."""
    name = "base"

    def __init__(self) -> None:
        self.last_status = ""   # "" | BOT-WALL | UNREACHABLE | PARSE-FAIL
        self.last_note = ""

    def supports(self, currency: str) -> bool:  # overridden
        return False

    def _get(self, url: str, extra_headers: Optional[dict] = None,
             retries: int = RETRIES) -> Optional[requests.Response]:
        headers = {"User-Agent": UA, "Accept": "*/*"}
        if extra_headers:
            headers.update(extra_headers)
        last = ""
        for attempt in range(retries + 1):
            try:
                r = requests.get(url, headers=headers, timeout=HTTP_TIMEOUT)
            except requests.RequestException as e:
                last = f"network:{type(e).__name__}"
            else:
                if r.status_code == 200:
                    return r
                last = f"HTTP {r.status_code}"
            if attempt < retries:
                time.sleep(BACKOFF_S)
        self.last_status = "UNREACHABLE"
        self.last_note = last
        return None

    @staticmethod
    def _is_botwall(text: str) -> bool:
        head = text.lstrip()[:600].lower()
        return head.startswith("<") or "requires javascript" in head or "<html" in head

    def _check_botwall(self, r: requests.Response, expect: str = "csv") -> bool:
        if expect == "csv" and self._is_botwall(r.text):
            self.last_status = "BOT-WALL"
            self.last_note = "HTML/JS challenge instead of data"
            return True
        return False

    def _parse_fail(self, why: str) -> None:
        self.last_status = "PARSE-FAIL"
        self.last_note = why

    def fetch(self, currency: str) -> Optional[YieldSeries]:
        self.last_status = ""
        self.last_note = ""
        try:
            return self._fetch(currency)
        except TenorMismatch as e:
            self.last_status, self.last_note = "TENOR-MISMATCH", str(e)
            raise
        except Exception as e:  # never propagate — graceful per source
            self._parse_fail(f"{type(e).__name__}: {e}")
            return None

    def _fetch(self, currency: str) -> Optional[YieldSeries]:  # overridden
        return None


def _safe_float(v) -> Optional[float]:
    """Coerce to float, tolerating 'NA'/'.'/''/None and numpy/NaN."""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        if s in ("", ".", "NA", "N/A", "n/a", "-", "ND"):
            return None
        try:
            return float(s)
        except ValueError:
            return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return None if f != f else f  # drop NaN


def _mk(currency, source, tenor, freq, dates, vals, note="") -> Optional[YieldSeries]:
    pts = []
    for d, v in zip(dates, vals):
        f = _safe_float(v)
        if f is None or d is None:
            continue
        dd = d.date() if isinstance(d, (pd.Timestamp, datetime)) else d
        if dd is None or (isinstance(d, float) and pd.isna(d)):
            continue
        try:
            if pd.isna(dd):
                continue
        except (TypeError, ValueError):
            pass
        pts.append((dd, f))
    pts.sort(key=lambda x: x[0])
    if not pts:
        return None
    return YieldSeries(currency, source, tenor, freq, pts, note)


def _csv_reader(text: str):
    import csv as _csv
    return _csv.reader(io.StringIO(text))


# ---------------------------------------------------------------------------
# (a) Stooq — one-stop candidate
# ---------------------------------------------------------------------------

class StooqSource(BaseSource):
    name = "stooq"
    SYMBOLS = {
        "USD": ["2usy.b"], "EUR": ["2dey.b"], "GBP": ["2uky.b", "2gby.b"],
        "JPY": ["2jpy.b"], "AUD": ["2auy.b"], "NZD": ["2nzy.b"],
        "CAD": ["2cay.b"], "CHF": ["2chy.b"],
    }
    HOSTS = ["https://stooq.com/q/d/l/?s={s}&i=d", "https://stooq.pl/q/d/l/?s={s}&i=d"]

    def supports(self, currency: str) -> bool:
        return currency in self.SYMBOLS

    def _fetch(self, currency: str) -> Optional[YieldSeries]:
        tried = []
        for sym in self.SYMBOLS[currency]:
            for tmpl in self.HOSTS:
                r = self._get(tmpl.format(s=sym))
                if r is None:
                    tried.append(f"{sym}:{self.last_note}")
                    continue
                if self._check_botwall(r):
                    tried.append(f"{sym}:bot-wall")
                    continue
                if "No data" in r.text or len(r.text.strip()) < 20:
                    tried.append(f"{sym}:no-data")
                    continue
                try:
                    df = pd.read_csv(io.StringIO(r.text))
                except Exception as e:
                    tried.append(f"{sym}:parse({e})")
                    continue
                if "Date" not in df.columns or "Close" not in df.columns:
                    tried.append(f"{sym}:cols{list(df.columns)}")
                    continue
                s = _mk(currency, self.name, "2y", "daily",
                        pd.to_datetime(df["Date"], errors="coerce"),
                        pd.to_numeric(df["Close"], errors="coerce"), note=f"symbol={sym}")
                if s:
                    return s
                tried.append(f"{sym}:empty")
        if any("bot-wall" in t for t in tried):
            self.last_status = "BOT-WALL"
        elif not self.last_status:
            self.last_status = "UNREACHABLE"
        self.last_note = "; ".join(tried)
        return None


# ---------------------------------------------------------------------------
# (b) FRED CSV — US 2y (DGS2)
# ---------------------------------------------------------------------------

class FredSource(BaseSource):
    name = "fred"
    SERIES = {"USD": ("DGS2", "2y")}

    def supports(self, currency: str) -> bool:
        return currency in self.SERIES

    def _fetch(self, currency: str) -> Optional[YieldSeries]:
        sid, tenor = self.SERIES[currency]
        r = self._get(f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={sid}", retries=3,
                      extra_headers={"User-Agent": FRED_UA})
        if r is None:
            return None
        if self._check_botwall(r):
            return None
        df = pd.read_csv(io.StringIO(r.text))
        date_col = df.columns[0]
        val_cols = [c for c in df.columns if c != date_col]
        if not val_cols:
            self._parse_fail(f"cols {list(df.columns)}")
            return None
        # fredgraph names the value column after the series id (DGS2 = 2-year)
        require_2y(self.name, str(val_cols[0]).strip() == sid,
                   f"value column {val_cols[0]!r}, expected {sid!r}")
        vals = pd.to_numeric(df[val_cols[0]].replace(".", np.nan), errors="coerce")
        s = _mk(currency, self.name, tenor, "daily",
                pd.to_datetime(df[date_col], errors="coerce"), vals, note=f"id={sid}")
        if not s:
            self._parse_fail("no numeric observations")
        return s


# ---------------------------------------------------------------------------
# Generic FRED single-series reader (NOT currency-keyed)
# ---------------------------------------------------------------------------

class FredSeriesSource(BaseSource):
    """Keyless reader for any single FRED series (e.g. DFII10), via the same
    fredgraph CSV mechanism used by FredSource. Returns a tidy DataFrame
    (date, value, source) — NOT a per-currency YieldSeries — so it stays out of
    the currency rate registry. `.` is treated as missing (FRED's NA marker).
    """
    name = "fred_series"

    def __init__(self, series_id: str) -> None:
        super().__init__()
        self.series_id = series_id

    def fetch_series(self) -> Optional[pd.DataFrame]:
        """Read-only fetch; returns DataFrame(date, value, source) or None.
        Never raises — failures are captured in last_status/last_note."""
        self.last_status = ""
        self.last_note = ""
        try:
            return self._fetch_series()
        except Exception as e:
            self._parse_fail(f"{type(e).__name__}: {e}")
            return None

    def _fetch_series(self) -> Optional[pd.DataFrame]:
        url = f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={self.series_id}"
        r = self._get(url, retries=3, extra_headers={"User-Agent": FRED_UA})
        text = None if (r is None or self._check_botwall(r)) else r.text
        if not text or not text.strip():
            # requests timed out / empty / tarpitted — retry with curl, which uses a
            # different network stack (Happy Eyeballs) and sometimes succeeds when
            # requests hangs on the same FRED series. Absolute path for launchd PATH.
            text = self._curl_csv(url)
        if not text or not text.strip():
            return None
        return self._parse_fred_csv(text)

    def _curl_csv(self, url: str) -> Optional[str]:
        """Second-attempt fetch via /usr/bin/curl. Returns CSV text or None."""
        try:
            p = subprocess.run(
                ["/usr/bin/curl", "-s", "--max-time", "20", "--compressed",
                 "-H", f"User-Agent: {FRED_UA}", url],
                capture_output=True, text=True, timeout=25,
            )
        except (OSError, subprocess.SubprocessError) as e:
            self._parse_fail(f"curl fallback {type(e).__name__}: {e}")
            return None
        if p.returncode != 0 or not p.stdout.strip():
            self._parse_fail(f"curl fallback rc={p.returncode}")
            return None
        return p.stdout

    def _parse_fred_csv(self, text: str) -> Optional[pd.DataFrame]:
        df = pd.read_csv(io.StringIO(text))
        date_col = df.columns[0]
        val_cols = [c for c in df.columns if c != date_col]
        if not val_cols:
            self._parse_fail(f"cols {list(df.columns)}")
            return None
        out = pd.DataFrame({
            "date": pd.to_datetime(df[date_col], errors="coerce"),
            "value": pd.to_numeric(df[val_cols[0]].replace(".", np.nan), errors="coerce"),
            "source": self.name,
        })
        out = out.dropna(subset=["date", "value"]).sort_values("date").reset_index(drop=True)
        if out.empty:
            self._parse_fail("no numeric observations")
            return None
        return out


class TreasuryRealYieldSource(BaseSource):
    """Keyless reader for the US Treasury 10y REAL yield (TIPS real yield curve).

    Backup/alternate for FRED DFII10 — works when fredgraph is down. The daily
    real yield curve is published per-year as CSV with header
    `Date,"5 YR","7 YR","10 YR","20 YR","30 YR"`; we take the "10 YR" column.
    Fetches the requested years and concatenates (≥3y gives enough history for
    a 252-obs vol window + 21-row change). Returns DataFrame(date, value,
    source="treasury")."""
    name = "treasury"
    URL = ("https://home.treasury.gov/resource-center/data-chart-center/"
           "interest-rates/daily-treasury-rates.csv/{year}/all"
           "?type=daily_treasury_real_yield_curve"
           "&field_tdr_date_value={year}&page&_format=csv")
    COLUMN = "10 YR"

    def fetch_real_10y(self, years: list[int]) -> Optional[pd.DataFrame]:
        """Fetch + concat the 10y real yield for the given years. Never raises."""
        self.last_status = ""
        self.last_note = ""
        try:
            return self._fetch_years(years)
        except Exception as e:
            self._parse_fail(f"{type(e).__name__}: {e}")
            return None

    def _parse_year(self, text: str) -> Optional[pd.DataFrame]:
        df = pd.read_csv(io.StringIO(text))
        cols = {str(c).strip(): c for c in df.columns}
        date_col = cols.get("Date") or df.columns[0]
        ten = cols.get(self.COLUMN)
        if ten is None:
            return None
        out = pd.DataFrame({
            "date": pd.to_datetime(df[date_col], format="%m/%d/%Y", errors="coerce"),
            "value": pd.to_numeric(df[ten], errors="coerce"),
            "source": self.name,
        })
        return out.dropna(subset=["date", "value"])

    def _fetch_years(self, years: list[int]) -> Optional[pd.DataFrame]:
        frames, got = [], []
        for y in years:
            r = self._get(self.URL.format(year=y), retries=2)
            if r is None:
                got.append(f"{y}:{self.last_note}")
                continue
            if self._check_botwall(r):
                got.append(f"{y}:bot-wall")
                continue
            part = self._parse_year(r.text)
            if part is None or part.empty:
                got.append(f"{y}:no-10YR")
                continue
            frames.append(part)
        if not frames:
            if not self.last_status:
                self._parse_fail("; ".join(got) or "no data")
            else:
                self.last_note = "; ".join(got)
            return None
        out = (pd.concat(frames, ignore_index=True)
               .drop_duplicates(subset=["date"], keep="last")
               .sort_values("date").reset_index(drop=True))
        return out


# ---------------------------------------------------------------------------
# (c) ECB Data Portal — EUR 2y
# ---------------------------------------------------------------------------

class EcbSource(BaseSource):
    name = "ecb"
    KEY = "B.U2.EUR.4F.G_N_A.SV_C_YM.SR_2Y"

    def supports(self, currency: str) -> bool:
        return currency == "EUR"

    def _fetch(self, currency: str) -> Optional[YieldSeries]:
        url = (f"https://data-api.ecb.europa.eu/service/data/YC/{self.KEY}"
               f"?format=csvdata&lastNObservations=1200")
        r = self._get(url)
        if r is None:
            return None
        if self._check_botwall(r):
            return None
        df = pd.read_csv(io.StringIO(r.text))
        if "TIME_PERIOD" not in df.columns or "OBS_VALUE" not in df.columns:
            self._parse_fail(f"cols {list(df.columns)[:8]}")
            return None
        dtypes = set(df.get("DATA_TYPE_FM", pd.Series(dtype=str)).dropna().astype(str))
        require_2y(self.name, dtypes == {"SR_2Y"}, f"DATA_TYPE_FM={sorted(dtypes)}")
        s = _mk(currency, self.name, "2y", "daily",
                pd.to_datetime(df["TIME_PERIOD"], errors="coerce"),
                pd.to_numeric(df["OBS_VALUE"], errors="coerce"), note=f"key={self.KEY}")
        if not s:
            self._parse_fail("no observations")
        return s


# ---------------------------------------------------------------------------
# (d) DBnomics FED fitted curve — USD 2y fallback
# ---------------------------------------------------------------------------

class DbnomicsFedSource(BaseSource):
    name = "dbnomics_fed"
    PATH = "FED/NOMINAL_YIELD_CURVE/SVENY02"

    def supports(self, currency: str) -> bool:
        return currency == "USD"

    def _fetch(self, currency: str) -> Optional[YieldSeries]:
        r = self._get(f"https://api.db.nomics.world/v22/series/{self.PATH}?observations=1",
                      retries=3)
        if r is None:
            return None
        payload = r.json()
        docs = (payload.get("series") or {}).get("docs") or []
        if not docs:
            self._parse_fail("no series docs")
            return None
        d = docs[0]
        code = str(d.get("series_code", ""))
        require_2y(self.name, code == self.PATH.rsplit("/", 1)[-1], f"series_code={code!r}")
        s = _mk(currency, self.name, "2y", "daily",
                pd.to_datetime(pd.Series(d.get("period", [])), errors="coerce"),
                d.get("value", []), note=self.PATH)
        if not s:
            self._parse_fail("no observations")
        return s


# ---------------------------------------------------------------------------
# (e) Bank of Canada Valet — CAD 2y
# ---------------------------------------------------------------------------

class BocValetSource(BaseSource):
    name = "boc_valet"
    SERIES = "BD.CDN.2YR.DQ.YLD"

    def supports(self, currency: str) -> bool:
        return currency == "CAD"

    def _fetch(self, currency: str) -> Optional[YieldSeries]:
        url = f"https://www.bankofcanada.ca/valet/observations/{self.SERIES}/json?recent=1500"
        r = self._get(url)
        if r is None:
            return None
        payload = r.json()
        label = str(((payload.get("seriesDetail") or {}).get(self.SERIES) or {}).get("label", ""))
        require_2y(self.name, re.search(r"\b2[- ]year\b", label, re.I) is not None,
                   f"seriesDetail label {label!r}")
        obs = payload.get("observations") or []
        dates, vals = [], []
        for o in obs:
            dates.append(o.get("d"))
            vals.append((o.get(self.SERIES) or {}).get("v"))
        s = _mk(currency, self.name, "2y", "daily",
                pd.to_datetime(pd.Series(dates), errors="coerce"),
                pd.to_numeric(pd.Series(vals), errors="coerce"), note=f"series={self.SERIES}")
        if not s:
            self._parse_fail("no observations")
        return s


# ---------------------------------------------------------------------------
# (f) MoF Japan — JPY 2y (shift_jis, Japanese-era dates, col 2年)
# ---------------------------------------------------------------------------

class MofJgbSource(BaseSource):
    """MoF JGB constant-maturity yields, column 2年 (2-year).

    Two files, both needed (audit 2026-09-23, 1.3): jgbcm_all.csv is the history
    and stops at the end of the PREVIOUS month; jgbcm.csv (not under /data/) is
    the current month. Both parse the same way (shift_jis, header=1, era dates
    like R8.9.17). They are merged and deduped on date, the current month wins.
    The historical file is required; the current-month file is best-effort (a
    failure there leaves the series ending at month end, which the freshness
    check then reports)."""
    name = "mof_jgb"
    URL_HISTORY = "https://www.mof.go.jp/jgbs/reference/interest_rate/data/jgbcm_all.csv"
    URL_CURRENT = "https://www.mof.go.jp/jgbs/reference/interest_rate/jgbcm.csv"
    COLUMN = "2年"
    ERA_BASE = {"M": 1867, "T": 1911, "S": 1925, "H": 1988, "R": 2018}

    def supports(self, currency: str) -> bool:
        return currency == "JPY"

    def _parse_date(self, raw: str) -> Optional[date]:
        raw = str(raw).strip()
        if raw and raw[0] in self.ERA_BASE:
            try:
                yr, mo, dy = raw[1:].replace("/", ".").split(".")
                return date(self.ERA_BASE[raw[0]] + int(yr), int(mo), int(dy))
            except Exception:
                return None
        try:
            return pd.to_datetime(raw).date()
        except Exception:
            return None

    def parse_csv(self, content: bytes) -> list[tuple[date, float]]:
        """One MoF file -> [(date, 2y yield)]. Raises ValueError without a 2年
        column (the maturity is read from the file's own header)."""
        raw = content.decode("shift_jis", errors="replace")
        df = pd.read_csv(io.StringIO(raw), header=1, dtype=str)
        cols = {str(c).strip(): c for c in df.columns}
        require_2y("mof_jgb", self.COLUMN in cols, f"no {self.COLUMN} column in {list(cols)[:6]}")
        vals = pd.to_numeric(df[cols[self.COLUMN]].replace("-", np.nan), errors="coerce")
        out = []
        for d, v in zip(df[df.columns[0]], vals):
            dd = self._parse_date(d)
            if dd is not None and not pd.isna(v):
                out.append((dd, float(v)))
        return out

    @staticmethod
    def merge(history: list[tuple[date, float]],
              current: list[tuple[date, float]]) -> list[tuple[date, float]]:
        """Union on date; the current-month file wins a shared date."""
        by_date = dict(history)
        by_date.update(dict(current))
        return sorted(by_date.items())

    def _fetch(self, currency: str) -> Optional[YieldSeries]:
        r = self._get(self.URL_HISTORY)
        if r is None or self._check_botwall(r):
            return None
        history = self.parse_csv(r.content)
        current: list[tuple[date, float]] = []
        note = "jgbcm_all.csv"
        rc = self._get(self.URL_CURRENT)
        if rc is not None and not self._check_botwall(rc):
            current = self.parse_csv(rc.content)
            note += f"+jgbcm.csv({len(current)})"
        else:
            log.warning("MoF current-month file unavailable (%s); series ends at "
                        "the last month end.", self.last_note)
            self.last_status, self.last_note = "", ""
        pts = self.merge(history, current)
        s = _mk(currency, self.name, "2y", "daily", [d for d, _ in pts],
                [v for _, v in pts], note=note)
        if not s:
            self._parse_fail("no parsable 2Y series")
        return s


# ---------------------------------------------------------------------------
# (g) Bank of England IADB — GBP 2y
# ---------------------------------------------------------------------------

class BoeSource(BaseSource):
    """GBP 2y = Bank of England GLC nominal SPOT curve, maturity 2.0 years
    (audit 2026-09-23, 1.4). BoE has no daily 2y series in the IADB — the old
    code IUDSNPY is the 5-year nominal PAR yield.

    Daily fetch: latest-yield-curve-data.zip -> "GLC Nominal daily data current
    month.xlsx", sheet "4. spot curve"; the maturity column is the one whose
    header row ("years:") reads 2.0. Only the current month is served, so the
    source is INCREMENTAL: rate_fetch assesses it together with the rows it
    already holds for this source (history backfilled once from
    glcnominalddata.zip by src.rate_migrations)."""
    name = "boe_glc"
    incremental = True
    URL_LATEST = ("https://www.bankofengland.co.uk/-/media/boe/files/statistics/"
                  "yield-curves/latest-yield-curve-data.zip")
    URL_HISTORY = ("https://www.bankofengland.co.uk/-/media/boe/files/statistics/"
                   "yield-curves/glcnominalddata.zip")
    MEMBER_LATEST = "GLC Nominal daily data current month.xlsx"
    SHEET = "4. spot curve"
    MATURITY_YEARS = 2.0

    def supports(self, currency: str) -> bool:
        return currency == "GBP"

    @classmethod
    def parse_spot_xlsx(cls, content: bytes) -> list[tuple[date, float]]:
        """One GLC nominal workbook -> [(date, 2.0y spot)]. Raises ValueError if
        the sheet or a 2.0-year maturity column is missing."""
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        if cls.SHEET not in wb.sheetnames:
            raise ValueError(f"no sheet {cls.SHEET!r} in {wb.sheetnames}")
        col = None
        out: list[tuple[date, float]] = []
        for row in wb[cls.SHEET].iter_rows(values_only=True):
            if not row:
                continue
            if col is None:
                if str(row[0]).strip().lower() == "years:":
                    hits = [i for i, c in enumerate(row[1:], start=1)
                            if isinstance(c, (int, float)) and float(c) == cls.MATURITY_YEARS]
                    require_2y("boe_glc", bool(hits), f"no {cls.MATURITY_YEARS}-year "
                               f"maturity in the 'years:' header row")
                    col = hits[0]
                continue
            d, v = row[0], row[col] if col < len(row) else None
            if isinstance(d, datetime) and isinstance(v, (int, float)):
                out.append((d.date(), float(v)))
        if col is None:
            raise ValueError("no 'years:' maturity header row")
        return out

    def _fetch(self, currency: str) -> Optional[YieldSeries]:
        import zipfile
        r = self._get(self.URL_LATEST)
        if r is None:
            return None
        if r.content[:2] != b"PK":
            self.last_status, self.last_note = "BOT-WALL", "not a zip"
            return None
        z = zipfile.ZipFile(io.BytesIO(r.content))
        if self.MEMBER_LATEST not in z.namelist():
            self._parse_fail(f"no {self.MEMBER_LATEST!r} in {z.namelist()}")
            return None
        pts = self.parse_spot_xlsx(z.read(self.MEMBER_LATEST))
        s = _mk(currency, self.name, "2y", "daily", [d for d, _ in pts],
                [v for _, v in pts], note=f"GLC nominal spot {self.MATURITY_YEARS}y (current month)")
        if not s:
            self._parse_fail("no 2.0y observations in the current-month file")
        return s

    def fetch_history(self, since: date) -> list[tuple[date, float]]:
        """One-off backfill (src.rate_migrations): every workbook in
        glcnominalddata.zip (~39 MB) covering >= `since`. Raises on failure."""
        return self.fetch_history_with_coverage(since)[0]

    def fetch_history_with_coverage(self, since: date) -> tuple[list[tuple[date, float]], date]:
        """(points >= since, covered_until). covered_until = the day before the
        newest workbook in the archive was written (zip entry time): every
        business day up to it without a value is a non-trading day."""
        import zipfile
        r = self._get(self.URL_HISTORY)
        if r is None or r.content[:2] != b"PK":
            raise RuntimeError(f"BoE GLC history download failed: {self.last_note or 'not a zip'}")
        z = zipfile.ZipFile(io.BytesIO(r.content))
        pts: dict[date, float] = {}
        built = max(date(*i.date_time[:3]) for i in z.infolist())
        for name in z.namelist():
            yrs = [int(t) for t in name.replace(".xlsx", "").replace("_", " ").split()
                   if t.isdigit() and len(t) == 4]
            if yrs and max(yrs) < since.year and "present" not in name:
                continue
            for d, v in self.parse_spot_xlsx(z.read(name)):
                if d >= since:
                    pts[d] = v
        return sorted(pts.items()), built - pd.Timedelta(days=1).to_pytimedelta()


# ---------------------------------------------------------------------------
# (h) RBA F2 — AUD 2y (multi-row header)
# ---------------------------------------------------------------------------

class RbaSource(BaseSource):
    name = "rba"
    URL = "https://www.rba.gov.au/statistics/tables/csv/f2-data.csv"
    SERIES_ID = "FCMYGBAG2D"   # Australian Government 2 year bond, daily

    def supports(self, currency: str) -> bool:
        return currency == "AUD"

    @classmethod
    def find_2y_column(cls, rows: list[list[str]]) -> int:
        """The F2 column whose 'Series ID' row is FCMYGBAG2D AND whose 'Title'
        row says 2 year (both are the table's own metadata). Raises
        TenorMismatch otherwise."""
        head = rows[:15]
        meta = {str(r[0]).strip().lower(): r for r in head if r}
        ids, titles = meta.get("series id"), meta.get("title")
        require_2y("rba", ids is not None and titles is not None,
                   "no 'Series ID'/'Title' header rows")
        cols = [i for i, c in enumerate(ids) if str(c).strip().upper() == cls.SERIES_ID]
        require_2y("rba", len(cols) == 1, f"series id {cls.SERIES_ID} found {len(cols)}x")
        ci = cols[0]
        title = str(titles[ci]) if ci < len(titles) else ""
        require_2y("rba", re.search(r"\b2[- ]year\b", title, re.I) is not None,
                   f"title {title!r}")
        return ci

    def _fetch(self, currency: str) -> Optional[YieldSeries]:
        r = self._get(self.URL, extra_headers={
            "Referer": "https://www.rba.gov.au/statistics/tables/",
            "Accept": "text/csv,application/csv,*/*",
        })
        if r is None:
            return None
        if self._check_botwall(r):
            return None
        rows = list(_csv_reader(r.text))
        if not rows:
            self._parse_fail("empty")
            return None
        target_col = self.find_2y_column(rows)
        dates, vals = [], []
        for row in rows:
            if len(row) <= target_col:
                continue
            d = pd.to_datetime(str(row[0]).strip(), errors="coerce", dayfirst=True)
            if pd.isna(d):
                continue
            try:
                vals.append(float(row[target_col]))
                dates.append(d)
            except (ValueError, TypeError):
                continue
        s = _mk(currency, self.name, "2y", "daily", dates, vals, note="F2 col=%d" % target_col)
        if not s:
            self._parse_fail("reachable / could not extract 2y column")
        return s


# ---------------------------------------------------------------------------
# (i) RBNZ B2 — NZD 2y (XLSX via openpyxl)
# ---------------------------------------------------------------------------

class RbnzSource(BaseSource):
    name = "rbnz"
    URLS = [
        "https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b2/hb2-daily.xlsx",
        "https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/tables/b2/hb2-daily.xlsx",
    ]

    def supports(self, currency: str) -> bool:
        return currency == "NZD"

    def _fetch(self, currency: str) -> Optional[YieldSeries]:
        for url in self.URLS:
            r = self._get(url, extra_headers={"Referer": "https://www.rbnz.govt.nz/statistics"})
            if r is None:
                continue
            if r.content[:2] != b"PK":
                self.last_status = "UNREACHABLE"
                self.last_note = "WAF/HTML, not XLSX (validate from a residential IP)"
                continue
            try:
                from openpyxl import load_workbook
            except Exception:
                self._parse_fail("openpyxl not installed (pip install openpyxl)")
                return None
            try:
                wb = load_workbook(io.BytesIO(r.content), data_only=True, read_only=True)
            except Exception as e:
                self._parse_fail(f"xlsx open {e}")
                continue
            ws = wb["Data"] if "Data" in wb.sheetnames else wb[wb.sheetnames[0]]
            rows = list(ws.iter_rows(values_only=True))
            col = None
            for hdr in rows[:12]:
                for ci, cell in enumerate(hdr):
                    c = str(cell).lower() if cell is not None else ""
                    if ("2 year" in c or "2-year" in c or "2 yr" in c) and "bond" in c:
                        col = ci
                        break
                if col is not None:
                    break
            require_2y(self.name, col is not None,
                       "no '2 year ... bond' column in the header band")
            dates, vals = [], []
            for row in rows:
                if col >= len(row):
                    continue
                d = pd.to_datetime(row[0], errors="coerce")
                if pd.isna(d):
                    continue
                try:
                    vals.append(float(row[col])); dates.append(d)
                except (TypeError, ValueError):
                    continue
            s = _mk(currency, self.name, "2y", "daily", dates, vals, note=url.split("/")[-1])
            if s:
                return s
            self._parse_fail("reachable / 2y column found but no numeric rows")
        if not self.last_status:
            self._parse_fail("unreachable / needs local validation")
        return None


# ---------------------------------------------------------------------------
# (j) SNB data portal — CHF (cube rendoblid, D0=2J; often stale)
# ---------------------------------------------------------------------------

class SnbSource(BaseSource):
    name = "snb"
    URL = "https://data.snb.ch/api/cube/rendoblid/data/csv/en"
    MATURITY = "2J"

    def supports(self, currency: str) -> bool:
        return currency == "CHF"

    def _fetch(self, currency: str) -> Optional[YieldSeries]:
        r = self._get(self.URL)
        if r is None:
            return None
        if self._check_botwall(r):
            return None
        lines = [ln for ln in r.text.splitlines() if ln.strip()]
        hdr_idx = next((i for i, ln in enumerate(lines)
                        if ln.lower().startswith('"date"') or ln.lower().startswith("date;")), None)
        if hdr_idx is None:
            self._parse_fail("no Date header; verify cube id")
            return None
        try:
            df = pd.read_csv(io.StringIO("\n".join(lines[hdr_idx:])), sep=";")
        except Exception as e:
            self._parse_fail(f"csv {e}; verify cube")
            return None
        df.columns = [str(c).strip().strip('"') for c in df.columns]
        if "D0" not in df.columns or "Value" not in df.columns:
            self._parse_fail(f"unexpected cols {list(df.columns)}; verify cube")
            return None
        sub = df[df["D0"].astype(str).str.strip() == self.MATURITY]
        avail = sorted(df["D0"].dropna().astype(str).unique())[:12]
        require_2y(self.name, not sub.empty, f"no {self.MATURITY} maturity (have {avail})")
        s = _mk(currency, self.name, "2y", "daily",
                pd.to_datetime(sub["Date"], errors="coerce"),
                pd.to_numeric(sub["Value"], errors="coerce"),
                note=f"cube=rendoblid D0={self.MATURITY}")
        if not s:
            self._parse_fail("reachable / no 2J observations")
        return s


# ---------------------------------------------------------------------------
# Registry + preference
# ---------------------------------------------------------------------------

ALL_SOURCES: dict[str, BaseSource] = {
    s.name: s for s in [
        StooqSource(), FredSource(), EcbSource(), DbnomicsFedSource(),
        BocValetSource(), MofJgbSource(), BoeSource(), RbaSource(),
        RbnzSource(), SnbSource(),
    ]
}

SOURCE_PREFERENCE: dict[str, list[str]] = {
    "USD": ["fred", "dbnomics_fed", "stooq"],
    "EUR": ["ecb", "stooq"],
    "GBP": ["boe_glc", "stooq"],
    "JPY": ["mof_jgb", "stooq"],
    "AUD": ["rba", "stooq"],
    "NZD": ["rbnz", "stooq"],
    "CAD": ["boc_valet", "stooq"],
    "CHF": ["snb", "stooq"],
}


# ---------------------------------------------------------------------------
# Qualification
# ---------------------------------------------------------------------------

@dataclass
class Assessment:
    status: str          # OK | OK* | WRONG-TENOR | STALE | MONTHLY | SHORT-HIST | EMPTY
    qualifies: bool
    flags: list[str] = field(default_factory=list)


def assess(s: YieldSeries, today: date) -> Assessment:
    flags: list[str] = []
    if s.frequency == "monthly":
        return Assessment("MONTHLY", False, ["monthly→disqualified"])
    lag = s.business_days_lag(today)
    hist = s.history_years
    if s.frequency == "weekly":
        flags.append("weekly")
    if s.tenor != "2y":
        flags.append(f"tenor={s.tenor}")
    if hist < MIN_HISTORY_YEARS:
        flags.append(f"hist={hist:.1f}y")
    if lag is None:
        return Assessment("EMPTY", False, flags)
    if lag > STALE_LAG_BD:
        return Assessment("STALE", False, flags + [f"lag={lag}bd"])
    if hist < MIN_HISTORY_YEARS:
        return Assessment("SHORT-HIST", False, flags)
    if s.tenor == "10y":
        return Assessment("WRONG-TENOR", True, flags)
    status = "OK*" if flags else "OK"
    if lag > FRESH_LAG_BD:
        flags.append(f"lag={lag}bd")
        status = "OK*"
    return Assessment(status, True, flags)
