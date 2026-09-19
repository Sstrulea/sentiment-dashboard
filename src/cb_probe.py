"""Read-only data spike for the Central Banks section — NUMERIC sources + calendar.

Diagnostic only. Probes official / candidate sources for eight banks (Fed/USD,
ECB/EUR, BoE/GBP, BoJ/JPY, BoC/CAD, RBA/AUD, RBNZ/NZD, SNB/CHF) and prints a
coverage matrix. Writes nothing (the optional --json-out flag is the only file
write, to a path the caller chooses). Every probe follows the `RateSource`
contract of `src.rate_sources` (`name` / `supports(bank)` / `fetch(bank)` /
`last_status`) and reuses its HTTP layer (`BaseSource._get`: IPv4 pin, retries,
bot-wall sniffing) so it can be lifted into `src/cb_sources/` unchanged.

    A1  policy rate + last 4 decisions (official machine-readable series)
    A2  the bank's own rate trajectory (Fed SEP dots, RBNZ OCR track)
    A3  market-implied path (futures / OIS / probability feeds)  <- critical
    A4  Forex Factory consensus for decisions (LOCAL parquet + archive only)
    A5  official 2026-2027 meeting calendar + quiet-period rule

    0B (text + latency + trigger):
    E   errata for the 0A implied-path numbers      B1  decision statement (last 4)     B2  votes / dissents
    B3  press conference (video + transcript)       B4  minutes / accounts / summaries  B5  speeches, BIS feed, roster
    B6  Forex Factory central-bank event names      B7  <=15 min trigger (GH cron)      B8  PDF extractor check
    B9  private-repo scenario (GH minutes, Vercel)  a3bis  targeted A3 retry (Eurex, proxies, surveys)

    python -m src.cb_probe
    python -m src.cb_probe --bank USD --section a1
    python -m src.cb_probe --section a3
    python -m src.cb_probe --section a5 --bank CHF --json-out /tmp/cb.json
    python -m src.cb_probe --section b1 --section b2          # 0B; PDF text needs a PDF lib on PYTHONPATH (none in requirements)

Failure codes: BOTWALL / STALE / JS_ONLY / NO_HISTORY / SHORT_HORIZON / LICENSE / NONE.
Every run is from THIS machine's IP; `cdn=` in the per-source log marks hosts
behind a bot-management CDN (Cloudflare/Akamai/...) that may block GitHub Actions
even when they answer here.
"""
from __future__ import annotations

import argparse
import calendar
import csv
import hashlib
import io
import json
import logging
import math
import os
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from datetime import time as dtime
from difflib import SequenceMatcher
from email.utils import parsedate_to_datetime
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
from typing import Optional
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd

from .rate_sources import FRED_UA, BaseSource, _safe_float

log = logging.getLogger(__name__)

BANKS = {"USD": "Fed", "EUR": "ECB", "GBP": "BoE", "JPY": "BoJ",
         "CAD": "BoC", "AUD": "RBA", "NZD": "RBNZ", "CHF": "SNB"}
CORE_SECTIONS = ("a1", "a2", "a3", "a4", "a5")
SECTIONS = CORE_SECTIONS + ("e", "b1", "b2", "b3", "b4", "b5", "b6", "b7", "b8", "b9", "a3bis")
DATA_DIR = Path(__file__).resolve().parents[1] / "data"

# Pre-registered thresholds (spike brief).
MAX_LAG_BD = 1          # EOD, lag <= 1 business day
MIN_HIST_BD = 30        # minimum history; ideal >= 250 (12 months)
NB_UA = FRED_UA         # honest non-browser UA: RBA/FRED serve it, 403 the Chrome UA

MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_name) if m}
MONTHS.update({m.lower(): i for i, m in enumerate(calendar.month_abbr) if m})
MONTHS["sept"] = 9


# ---------------------------------------------------------------------------
# Data shapes
# ---------------------------------------------------------------------------

@dataclass
class CbSeries:
    bank: str
    source: str
    kind: str                                   # policy | curve | futures | sep | calendar | ff
    points: list = field(default_factory=list)  # [(date, float)] ascending
    extra: dict = field(default_factory=dict)
    note: str = ""

    @property
    def n_points(self) -> int:
        return len(self.points)

    @property
    def latest_date(self) -> Optional[date]:
        return self.points[-1][0] if self.points else self.extra.get("asof")

    @property
    def latest_value(self) -> Optional[float]:
        return self.points[-1][1] if self.points else None

    @property
    def history_start(self) -> Optional[date]:
        return self.points[0][0] if self.points else None

    def lag_bd(self, today: date) -> Optional[int]:
        """Business days between the latest observation and the most recent business day
        (weekend-safe: Friday data is lag 0 on a Saturday, Thursday data lag 1)."""
        d = self.latest_date
        if d is None:
            return None
        ref = today - timedelta(days=max(0, today.weekday() - 4))    # Sat/Sun -> Friday
        return 0 if d >= ref else int(np.busday_count(d, ref))


@dataclass
class Probe:
    bank: str
    section: str
    source: str
    status: str                  # OK | PARTIAL | FAIL
    code: str                    # failure code or ""
    series: Optional[CbSeries]
    note: str = ""
    flags: list = field(default_factory=list)
    cdn: str = ""


def _cdn_hint(h) -> str:
    keys = " ".join(k.lower() for k in h.keys())
    s = f"{h.get('Server', '')} {h.get('Via', '')}".lower()
    if "cf-ray" in keys or "cloudflare" in s:
        return "cloudflare"
    if "akamai" in s or "x-akamai" in keys or "akamai" in keys:
        return "akamai"
    if "x-iinfo" in keys or "incap" in keys:
        return "imperva"
    if "x-amz-cf-id" in keys or "cloudfront" in s:
        return "cloudfront"
    return ""


def _text(s: str) -> str:
    s = re.sub(r"<script.*?</script>|<style.*?</style>", " ", s, flags=re.S)
    return re.sub(r"\s+", " ", unescape(re.sub(r"<[^>]+>", " ", s)))


def _tables(s: str) -> list[list[list[str]]]:
    out = []
    for tb in re.findall(r"<table.*?</table>", s, flags=re.S):
        rows = [[_text(c).strip() for c in re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", tr, flags=re.S)]
                for tr in re.findall(r"<tr.*?</tr>", tb, flags=re.S)]
        out.append(rows)
    return out


def _pts(dates, vals) -> list[tuple[date, float]]:
    out = []
    for d, v in zip(dates, vals):
        f = _safe_float(v)
        if f is None or d is None or (isinstance(d, float) and pd.isna(d)):
            continue
        d = pd.Timestamp(d)
        if pd.isna(d):
            continue
        out.append((d.date(), f))
    out.sort(key=lambda x: x[0])
    return out


def _at(points, d: date) -> Optional[float]:
    """Last value on or before d (step function — policy rates are piecewise constant)."""
    v = None
    for pd_, pv in points:
        if pd_ > d:
            break
        v = pv
    return v


def _hist_bd(points) -> int:
    return int(np.busday_count(points[0][0], points[-1][0])) if len(points) > 1 else 0


# ---------------------------------------------------------------------------
# Source base (RateSource contract)
# ---------------------------------------------------------------------------

class CbSource(BaseSource):
    """RateSource-contract probe: `supports(bank)`, `fetch(bank) -> CbSeries|None`,
    `last_status` / `last_note`. `fetch` never raises (BaseSource wraps `_fetch`)."""
    name = "cb_base"
    section = ""
    banks: tuple = ()
    official = True          # published by the central bank / statistical authority itself
    policy_level = True      # series IS the policy-rate level (not a market proxy)

    def __init__(self) -> None:
        super().__init__()
        self.cdn = ""
        self._cache: dict = {}

    def supports(self, bank: str) -> bool:
        return bank in self.banks

    def get(self, url: str, ua: Optional[str] = NB_UA, retries: int = 1, **hdr):
        """`BaseSource._get` + CDN fingerprint. UA defaults to the non-browser one."""
        extra = dict(hdr)
        if ua:
            extra["User-Agent"] = ua
        r = self._get(url, extra_headers=extra, retries=retries)
        if r is not None:
            self.cdn = _cdn_hint(r.headers) or self.cdn
        return r

    def failure_code(self) -> str:
        if self.last_status == "BOT-WALL":
            return "BOTWALL"
        m = re.search(r"HTTP (\d+)", self.last_note or "")
        if m and m.group(1) in {"401", "403", "429", "451", "503"}:
            return "BOTWALL"
        return "NONE"


# ---------------------------------------------------------------------------
# A1 — policy rate (official machine-readable series)
# ---------------------------------------------------------------------------

def _fred_csv(src: CbSource, sid: str) -> Optional[list]:
    r = src.get(f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={sid}", ua=FRED_UA, retries=3)
    if r is None or src._check_botwall(r):
        return None
    df = pd.read_csv(io.StringIO(r.text))
    d = pd.to_datetime(df.iloc[:, 0], errors="coerce")
    v = pd.to_numeric(df.iloc[:, 1].replace(".", np.nan), errors="coerce")
    return _pts(d, v) or None


class FredTarget(CbSource):
    name = "fred_target"
    section = "a1"
    banks = ("USD",)

    def _fetch(self, bank):
        up, lo = _fred_csv(self, "DFEDTARU"), _fred_csv(self, "DFEDTARL")
        if not up or not lo:
            self.last_status = self.last_status or "UNREACHABLE"
            return None
        eff, sofr = _fred_csv(self, "DFF"), _fred_csv(self, "SOFR")
        return CbSeries(bank, self.name, "policy", up,
                        {"lower": lo, "bench": {"EFFR": eff, "SOFR": sofr}},
                        "DFEDTARU (upper) + DFEDTARL (lower); DFF/SOFR overnight")


class EcbPolicy(CbSource):
    name = "ecb_data_portal"
    section = "a1"
    banks = ("EUR",)
    BASE = "https://data-api.ecb.europa.eu/service/data/"

    def _obs(self, key: str, start="2024-01-01"):
        r = self.get(f"{self.BASE}{key}?format=csvdata&startPeriod={start}")
        if r is None or self._check_botwall(r):
            return None
        df = pd.read_csv(io.StringIO(r.text))
        if "TIME_PERIOD" not in df.columns:
            self._parse_fail(f"cols {list(df.columns)[:6]}")
            return None
        return _pts(pd.to_datetime(df["TIME_PERIOD"], errors="coerce"),
                    pd.to_numeric(df["OBS_VALUE"], errors="coerce")) or None

    def _fetch(self, bank):
        dfr = self._obs("FM/D.U2.EUR.4F.KR.DFR.LEV")
        if not dfr:
            return None
        mro = self._obs("FM/D.U2.EUR.4F.KR.MRR_FR.LEV")
        estr = self._obs("EST/B.EU000A2X2A25.WT")
        return CbSeries(bank, self.name, "policy", dfr,
                        {"mro": mro, "bench": {"EUR_STR": estr}},
                        "FM.D.U2.EUR.4F.KR.DFR.LEV (deposit facility, effective-dated, daily ffill); "
                        "MRR_FR = main refinancing; EST = euro short-term rate")


class BoePolicy(CbSource):
    name = "boe_iadb"
    section = "a1"
    banks = ("GBP",)

    def _obs(self, code: str):
        url = ("https://www.bankofengland.co.uk/boeapps/iadb/fromshowcolumns.asp?csv.x=yes"
               f"&Datefrom=01/Jan/2024&Dateto=now&SeriesCodes={code}&CSVF=TN&UsingCodes=Y")
        r = self.get(url, ua=None)      # BoE serves the browser UA (existing BoeSource does)
        if r is None or self._check_botwall(r):
            return None
        df = pd.read_csv(io.StringIO(r.text))
        return _pts(pd.to_datetime(df.iloc[:, 0], format="%d %b %Y", errors="coerce"),
                    pd.to_numeric(df.iloc[:, 1], errors="coerce")) or None

    def _fetch(self, bank):
        br = self._obs("IUDBEDR")
        if not br:
            return None
        return CbSeries(bank, self.name, "policy", br, {"bench": {"SONIA": self._obs("IUDSOIA")}},
                        "IUDBEDR official Bank Rate (daily); IUDSOIA SONIA")


class BojPolicy(CbSource):
    name = "boj_timeseries_api"
    section = "a1"
    banks = ("JPY",)
    policy_level = False    # API has the market call rate + Basic Loan Rate, NOT the guideline itself

    def _obs(self, db: str, code: str, start="202401"):
        url = ("https://www.stat-search.boj.or.jp/api/v1/getDataCode?format=json&lang=en"
               f"&db={db}&code={code}&startDate={start}")
        r = self.get(url, ua=None)
        if r is None:
            return None
        rs = (r.json().get("RESULTSET") or [])
        if not rs:
            return None
        v = rs[0]["VALUES"]
        return _pts(pd.to_datetime([str(x) for x in v["SURVEY_DATES"]], format="%Y%m%d"), v["VALUES"]) or None

    def _fetch(self, bank):
        call = self._obs("FM01", "STRDCLUCON")
        if not call:
            return None
        blr = self._obs("IR01", "MADR1Z@D")
        return CbSeries(bank, self.name, "policy", call, {"blr": blr, "bench": {"TONA~call": call}},
                        "STRDCLUCON = uncollateralized O/N call rate avg (market rate ~ policy guideline); "
                        "IR01 MADR1Z@D = Basic Loan Rate (policy + 25bp). Guideline itself: PDF only")


class BocPolicy(CbSource):
    name = "boc_valet"
    section = "a1"
    banks = ("CAD",)

    def _obs(self, sid: str):
        r = self.get(f"https://www.bankofcanada.ca/valet/observations/{sid}/json?start_date=2024-01-01", ua=None)
        if r is None:
            return None
        obs = r.json().get("observations") or []
        return _pts(pd.to_datetime([o["d"] for o in obs], errors="coerce"),
                    [(o.get(sid) or {}).get("v") for o in obs]) or None

    def _fetch(self, bank):
        tgt = self._obs("V39079")
        if not tgt:
            return None
        return CbSeries(bank, self.name, "policy", tgt, {"bench": {"CORRA": self._obs("AVG.INTWO")}},
                        "V39079 target for the overnight rate (business daily); AVG.INTWO = CORRA")


class RbaPolicy(CbSource):
    name = "rba_f1"
    section = "a1"
    banks = ("AUD",)
    URL = "https://www.rba.gov.au/statistics/tables/csv/f1-data.csv"

    def _fetch(self, bank):
        r = self.get(self.URL)          # NB_UA: the Chrome UA is what earns the 403 here
        if r is None:
            return None
        if self._check_botwall(r):
            return None
        rows = list(csv.reader(io.StringIO(r.text.lstrip("﻿"))))
        sid_i = next((i for i, x in enumerate(rows) if x and x[0] == "Series ID"), None)
        if sid_i is None:
            self._parse_fail("no 'Series ID' row in F1")
            return None
        col = {sid: j for j, sid in enumerate(rows[sid_i]) if sid}
        need = ("FIRMMCRTD", "FIRMMCRID")
        if any(n not in col for n in need):
            self._parse_fail(f"F1 ids missing: {[n for n in need if n not in col]}")
            return None
        data = [x for x in rows[sid_i + 1:] if x and re.match(r"\d{2}-[A-Za-z]{3}-\d{4}$", x[0])]
        dts = [datetime.strptime(x[0], "%d-%b-%Y") for x in data]

        def series(sid):
            return _pts(dts, [x[col[sid]] if len(x) > col[sid] else None for x in data])
        ois = {k: series(sid)[-1:] for k, sid in (("1M", "FIRMMOIS1D"), ("3M", "FIRMMOIS3D"), ("6M", "FIRMMOIS6D"))
               if sid in col}
        return CbSeries(bank, self.name, "policy", series("FIRMMCRTD"),
                        {"bench": {"AONIA/interbank O/N": series("FIRMMCRID")}, "ois_short": ois},
                        "F1 FIRMMCRTD cash rate target (daily); FIRMMCRID interbank overnight; "
                        "FIRMMOIS{1,3,6}D = 1/3/6M OIS (FENICS)")


class RbnzPolicy(CbSource):
    name = "rbnz_site"
    section = "a1"
    banks = ("NZD",)
    URLS = ["https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b2/hb2-daily.xlsx",
            "https://www.rbnz.govt.nz/monetary-policy/about-monetary-policy/the-official-cash-rate"]

    def _fetch(self, bank):
        notes = []
        for u in self.URLS:
            r = self.get(u, ua=None)
            if r is not None and r.content[:2] == b"PK":
                self._parse_fail("XLSX reachable but parser intentionally not built in this spike")
                return None
            notes.append(f"{u.rsplit('/', 1)[-1]}:{self.last_note or 'non-XLSX'}")
        self.last_note = "; ".join(notes)
        return None


class SnbPolicy(CbSource):
    name = "snb_data_portal"
    section = "a1"
    banks = ("CHF",)
    URL = "https://data.snb.ch/api/cube/snbgwdzid/data/csv/en?fromDate=2024-01-01"

    def _fetch(self, bank):
        r = self.get(self.URL, ua=None)
        if r is None or self._check_botwall(r):
            return None
        lines = [ln for ln in r.text.lstrip("﻿").splitlines() if ln.strip()]
        hi = next((i for i, ln in enumerate(lines) if ln.startswith('"Date"')), None)
        if hi is None:
            self._parse_fail("no Date header")
            return None
        df = pd.read_csv(io.StringIO("\n".join(lines[hi:])), sep=";")
        df["Value"] = pd.to_numeric(df["Value"], errors="coerce")
        lz, sa = df[df.D0 == "LZ"], df[df.D0 == "SARON"]
        return CbSeries(bank, self.name, "policy",
                        _pts(pd.to_datetime(lz.Date), lz.Value),
                        {"bench": {"SARON": _pts(pd.to_datetime(sa.Date), sa.Value)}},
                        f"cube snbgwdzid D0=LZ (SNB policy rate) / SARON; publishing "
                        f"{next((l for l in lines if 'PublishingDate' in l), '')[:40]}")


class BisPolicy(CbSource):
    """BIS WS_CBPOL — all 8 banks, keyless, but published with a multi-day lag: fallback only."""
    name = "bis_cbpol"
    section = "a1"
    banks = tuple(BANKS)
    official = False
    AREA = {"USD": "US", "EUR": "XM", "GBP": "GB", "JPY": "JP", "CAD": "CA", "AUD": "AU", "NZD": "NZ", "CHF": "CH"}
    URL = ("https://stats.bis.org/api/v2/data/dataflow/BIS/WS_CBPOL/1.0/D."
           + "+".join(AREA.values()) + "?startPeriod=2024-01-01&format=csv")

    def _fetch(self, bank):
        if "df" not in self._cache:
            r = self.get(self.URL, ua=None, retries=1)
            if r is None:
                return None
            self._cache["df"] = pd.read_csv(io.StringIO(r.text))[["REF_AREA", "TIME_PERIOD", "OBS_VALUE"]]
        df = self._cache["df"]
        g = df[df.REF_AREA == self.AREA[bank]]
        pts = _pts(pd.to_datetime(g.TIME_PERIOD), g.OBS_VALUE)
        if not pts:
            self._parse_fail("no rows for area")
            return None
        note = {"USD": "US = midpoint of target range", "EUR": "XM = deposit facility rate"}.get(bank, "")
        return CbSeries(bank, self.name, "policy", pts, {}, f"WS_CBPOL daily; {note}".strip("; "))


A1_PRIMARY = {"USD": "fred_target", "EUR": "ecb_data_portal", "GBP": "boe_iadb", "JPY": "bis_cbpol",
              "CAD": "boc_valet", "AUD": "rba_f1", "NZD": "bis_cbpol", "CHF": "snb_data_portal"}


# ---------------------------------------------------------------------------
# A2 — the bank's own rate trajectory
# ---------------------------------------------------------------------------

class FedSep(CbSource):
    name = "fed_sep_html"
    section = "a2"
    banks = ("USD",)

    def _fetch(self, bank):
        r = self.get("https://www.federalreserve.gov/monetarypolicy/fomccalendars.htm", ua=None)
        if r is None:
            return None
        ds = sorted(set(re.findall(r"fomcprojtabl(\d{8})\.htm", r.text)))
        ds = [d for d in ds if d <= date.today().strftime("%Y%m%d")]
        if not ds:
            self._parse_fail("no SEP links on calendar page")
            return None
        sep = ds[-1]
        r2 = self.get(f"https://www.federalreserve.gov/monetarypolicy/fomcprojtabl{sep}.htm", ua=None)
        if r2 is None:
            return None
        tabs = _tables(r2.text)
        med = next((row for tb in tabs for row in tb if row and row[0].lower() == "federal funds rate"), None)
        dots = next((tb for tb in tabs if tb and tb[0] and tb[0][0].startswith("Midpoint of target range")), None)
        if med is None or dots is None:
            self._parse_fail("median row or dot table not found")
            return None
        years = dots[0][1:]
        counts = {y: {} for y in years}
        for row in dots[1:]:
            for y, c in zip(years, row[1:]):
                if c.strip().isdigit():
                    counts[y][float(row[0])] = int(c)
        calc = {}
        for y, cs in counts.items():
            flat = sorted(v for v, n in cs.items() for _ in range(n))
            calc[y] = (float(np.median(flat)) if flat else None, len(flat))
        sd = datetime.strptime(sep, "%Y%m%d").date()
        return CbSeries(bank, self.name, "sep", [(sd, _safe_float(med[1]) or 0.0)],
                        {"asof": sd, "median_published": dict(zip(years, med[1:6])), "dots": counts,
                         "median_from_dots": calc},
                        f"SEP {sd} — accessible-version HTML tables (Table 1 median, Figure 2 dot counts)")


class RbnzOcrTrack(CbSource):
    name = "rbnz_mps"
    section = "a2"
    banks = ("NZD",)
    URL = "https://www.rbnz.govt.nz/monetary-policy/monetary-policy-statement/monetary-policy-statement-filtered-listing-page"

    def _fetch(self, bank):
        r = self.get(self.URL, ua=None)
        if r is None:
            return None
        self._parse_fail("reachable — MPS OCR-track parser not built in this spike")
        return None


class PathAbsence(CbSource):
    """Evidence that a bank does NOT publish its own rate path (official statement text)."""
    name = "no_own_path_evidence"
    section = "a2"
    banks = ("EUR", "GBP", "JPY", "CAD", "AUD", "CHF")
    official = True
    CHECKS = {
        "CHF": ("https://www.snb.ch/en/publications/communication/press-releases-restricted/pre_20260618",
                r"conditional inflation forecast.{0,160}", "conditional forecast on constant policy rate"),
        "EUR": ("https://www.ecb.europa.eu/press/pr/date/2026/html/ecb.mp260723~29f24d99bc.en.html",
                r"data-dependent and meeting-by-meeting approach", "meeting-by-meeting, no path"),
        "GBP": ("https://www.bankofengland.co.uk/monetary-policy-summary-and-minutes/2026/september-2026",
                r"market curve", "MPC conditions on the market curve, no own path"),
        "AUD": ("https://www.rba.gov.au/publications/smp/2026/aug/overview.html",
                r"cash rate is assumed to move in line with expectations derived from financial market pricing",
                "SMP assumes the market-implied cash-rate path; no own path"),
    }

    def _fetch(self, bank):
        chk = self.CHECKS.get(bank)
        if chk is None:
            self.last_status = "UNVERIFIED"
            self.last_note = "no page fetched in this spike (documentary: BoJ Outlook Report / BoC MPR carry no rate path)"
            return None
        url, rx, why = chk
        r = self.get(url, ua=NB_UA)
        if r is None:
            return None
        m = re.search(rx, _text(r.text))
        if not m:
            self.last_status = "UNVERIFIED"
            self.last_note = f"phrase not found on {url.rsplit('/', 1)[-1]}"
            return None
        return CbSeries(bank, self.name, "no_path", [], {"evidence": m.group(0)[:200], "url": url}, why)


# ---------------------------------------------------------------------------
# A3 — market-implied path (CRITICAL)
# ---------------------------------------------------------------------------

def _xlsx_rows(content: bytes, sheet: str):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    return list(wb[sheet].iter_rows(values_only=True))


class BoeOis(CbSource):
    """UK OIS instantaneous forward curve, monthly grid 1..60m, daily; published by noon t+1."""
    name = "boe_ois_curve"
    section = "a3"
    banks = ("GBP",)
    BASE = "https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/"
    SHEET = "1. fwds, short end"

    def _parse(self, blob: bytes, member_rx: str) -> dict:
        z = zipfile.ZipFile(io.BytesIO(blob))
        name = next(n for n in z.namelist() if re.search(member_rx, n))
        rows = _xlsx_rows(z.read(name), self.SHEET)
        mat_row = next(r for r in rows if r and str(r[0]).strip().lower().startswith("months"))
        months = [float(x) for x in mat_row[1:] if x is not None]
        out = {}
        for r in rows:
            if isinstance(r[0], (datetime, date)):
                vals = [v for v in r[1:1 + len(months)]]
                if all(isinstance(v, (int, float)) for v in vals) and vals:
                    out[pd.Timestamp(r[0]).date()] = vals
        return {"months": months, "curves": out}

    def _fetch(self, bank):
        r1 = self.get(self.BASE + "latest-yield-curve-data.zip", ua=None)
        if r1 is None:
            return None
        cur = self._parse(r1.content, r"OIS daily data current month")
        r2 = self.get(self.BASE + "oisddata.zip", ua=None)
        hist = self._parse(r2.content, r"2025 to present") if r2 is not None else {"curves": {}}
        curves = {**hist["curves"], **cur["curves"]}
        months = cur["months"]
        i12 = min(range(len(months)), key=lambda i: abs(months[i] - 12))
        pts = [(d, v[i12]) for d, v in sorted(curves.items())]
        last = max(curves)
        return CbSeries(bank, self.name, "curve", pts,
                        {"asof": last, "months": months, "curve": curves[last], "n_dates": len(curves),
                         "horizon_months": max(months)},
                        "points = 12m-ahead instantaneous OIS forward; current-month + '2025 to present' workbooks")


class AsxFutures(CbSource):
    """ASX futures via the JSON API behind asx.com.au (undocumented, used by the site itself)."""
    name = "asx_markit_json"
    section = "a3"
    banks = ("AUD", "NZD")
    CODE = {"AUD": "IB", "NZD": "BB"}
    URL = "https://asx.api.markitdigital.com/asx-research/1.0/derivatives/interest-rate/{c}/futures?days=1&height=179&width=179"

    def _fetch(self, bank):
        r = self.get(self.URL.format(c=self.CODE[bank]), ua=None)
        if r is None:
            return None
        items = (r.json().get("data") or {}).get("items") or []
        cons = []
        for it in items:
            px = it.get("pricePreviousSettlement")
            if px is None or not it.get("dateExpiry"):
                continue
            cons.append({"symbol": it["symbol"], "expiry": date.fromisoformat(it["dateExpiry"]),
                         "settle": float(px), "implied": round(100 - float(px), 4),
                         "settle_date": date.fromisoformat(it["datePreviousSettlement"])})
        if not cons:
            self._parse_fail("no settled contracts")
            return None
        return CbSeries(bank, self.name, "futures", [],
                        {"contracts": cons, "asof": max(c["settle_date"] for c in cons),
                         "instrument": "30-day interbank cash rate (1M)" if bank == "AUD" else "90-day NZ bank bill (3M BKBM)"},
                        f"code={self.CODE[bank]} snapshot only (no history endpoint)")


class AsxHistoryMirror(CbSource):
    """Community daily capture of the ASX RBA rate tracker (github.com/bpalmer4/ASX)."""
    name = "asx_history_mirror"
    section = "a3"
    banks = ("AUD",)
    official = False
    URL = "https://raw.githubusercontent.com/bpalmer4/ASX/main/ASX-COMBINED/ASX-COMBINED.csv"

    def _fetch(self, bank):
        r = self.get(self.URL, ua=None)
        if r is None:
            return None
        df = pd.read_csv(io.StringIO(r.text), index_col=0)
        d = pd.to_datetime(df.index, errors="coerce")
        mask = ~pd.isna(d)
        front = df[mask].apply(lambda r: r.dropna().iloc[0] if r.notna().any() else np.nan, axis=1)
        pts = _pts(d[mask], front.values)
        lastrow = df[mask].iloc[-1].dropna()
        hz = None
        if len(lastrow):
            yy, mm_ = str(lastrow.index[-1]).split("-")
            hz = date(int(yy), int(mm_), 28)
        return CbSeries(bank, self.name, "futures", pts,
                        {"n_scrapes": int(mask.sum()), "n_expiries": df.shape[1], "horizon": hz},
                        "unofficial third-party repo; ASX-COMBINED.csv = implied cash rate per expiry month per scrape day")


class MxCorra(CbSource):
    """Montréal Exchange 'Canadian Interest Rate Expectations' — server-rendered HTML tables."""
    name = "mx_corra_expectations"
    section = "a3"
    banks = ("CAD",)
    URL = "https://www.m-x.ca/en/trading/tools/canadian-interest-rate-expectations"

    def _fetch(self, bank):
        r = self.get(self.URL, ua=None)
        if r is None:
            return None
        html = r.text
        rx = (r'<td class="text-left">([A-Za-z]+ \d{4}) \((CRA|COA)([A-Z])(\d{2})\)</td>\s*<td>([\d.]+)</td>')
        cons = {"CRA": [], "COA": []}
        for label, fam, mc, yy, px in re.findall(rx, html):
            mon, yr = label.split()
            cons[fam].append({"month": date(int(yr), MONTHS[mon.lower()], 1), "settle": float(px),
                              "implied": round(100 - float(px), 4)})
        txt = _text(html)
        m = re.search(r"Actual CORRA Rate - ([\d.]+)%", txt)
        i0 = txt.find("BoC Meeting Date")
        region = txt[i0:txt.find("Bank of Canada Target Rates", i0)] if i0 >= 0 else ""
        meets = sorted({date(int(y), MONTHS[mo.lower()], int(d)) for mo, d, y in
                        re.findall(r"([A-Z][a-z]+) (\d{1,2}), (\d{4})", region) if mo.lower() in MONTHS})
        if not cons["CRA"]:
            self._parse_fail("no CRA rows — layout changed?")
            return None
        return CbSeries(bank, self.name, "futures", [],
                        {"contracts": cons, "corra_spot": float(m.group(1)) if m else None,
                         "boc_meetings": meets, "asof": None},
                        "HTML snapshot: CRA (3M) + COA (1M) prices; page carries NO as-of stamp")


class JpxTona(CbSource):
    """JPX daily settlement-price CSV (all derivatives, cp932); 3M TONA futures rows only."""
    name = "jpx_settlement_csv"
    section = "a3"
    banks = ("JPY",)
    PAGE = "https://www.jpx.co.jp/english/markets/derivatives/settlement-price/index.html"

    def _fetch(self, bank):
        r = self.get(self.PAGE, ua=None)
        if r is None:
            return None
        m = re.search(r'href="([^"]*/rb_e(\d{8})\.csv)"', r.text)
        if not m:
            self._parse_fail("no rb_e*.csv link (JS-only page?)")
            return None
        r2 = self.get("https://www.jpx.co.jp" + m.group(1), ua=None)
        if r2 is None:
            return None
        cons = []
        for ln in r2.content.decode("cp932", errors="replace").splitlines():
            mm = re.match(r"\d+,FUT_TOA3M_(\d{6}),,(\d{6}),,([\d.]+)", ln)
            if mm:
                exp = datetime.strptime(mm.group(1), "%y%m%d").date()
                cons.append({"expiry": exp, "month": mm.group(2), "settle": float(mm.group(3)),
                             "implied": round(100 - float(mm.group(3)), 4)})
        if not cons:
            self._parse_fail("no FUT_TOA3M rows")
            return None
        return CbSeries(bank, self.name, "futures", [],
                        {"contracts": cons, "asof": datetime.strptime(m.group(2), "%Y%m%d").date()},
                        "3M TONA futures; only the latest business day is exposed (older rb_e dates 404)")


class AtlantaMpt(CbSource):
    """Atlanta Fed Market Probability Tracker — SOFR-option-implied 3M SOFR distribution per reference quarter."""
    name = "atlantafed_mpt"
    section = "a3"
    banks = ("USD",)
    URL = "https://www.atlantafed.org/-/media/Project/Atlanta/FRBA/Documents/cenfis/market-probability-tracker/mpt_histdata.xlsx"

    def _fetch(self, bank):
        r = self.get(self.URL, ua=None, retries=1)
        if r is None:
            return None
        df = pd.read_excel(io.BytesIO(r.content), sheet_name="DATA")
        df["date"] = pd.to_datetime(df["date"])
        df["reference_start"] = pd.to_datetime(df["reference_start"])
        mean = df[df.field == "Rate: mean"]
        last = mean.date.max()
        cur = mean[mean.date == last].sort_values("reference_start")
        pts = _pts(mean[mean.reference_start == mean[mean.date == last].reference_start.min()].date,
                   mean[mean.reference_start == mean[mean.date == last].reference_start.min()].value)
        return CbSeries(bank, self.name, "curve", pts,
                        {"asof": last.date(), "n_dates": int(mean.date.nunique()),
                         "first_date": mean.date.min().date(),
                         "windows": [(a.date(), float(b)) for a, b in zip(cur.reference_start, cur.value)],
                         "target_range": str(cur.target_range.iloc[0]) if len(cur) else "",
                         "horizon": cur.reference_start.max().date() if len(cur) else None},
                        "xlsx sheet DATA: 'Rate: mean' (bps) per reference quarter; embedded licence text: "
                        "'personal and educational purposes only'")


class TreasuryBills(CbSource):
    """US Treasury daily par yield curve (1M..30Y) — public domain; coarse T-bill proxy."""
    name = "treasury_par_curve"
    section = "a3"
    banks = ("USD",)

    def _fetch(self, bank):
        y = date.today().year
        frames = []
        for yr in (y - 1, y):
            url = ("https://home.treasury.gov/resource-center/data-chart-center/interest-rates/"
                   f"daily-treasury-rates.csv/{yr}/all?type=daily_treasury_yield_curve&field_tdr_date_value={yr}&page&_format=csv")
            r = self.get(url, ua=None)
            if r is None or self._check_botwall(r):
                if yr == y:
                    return None
                continue
            frames.append(pd.read_csv(io.StringIO(r.text)))
        df = pd.concat(frames, ignore_index=True)
        d = pd.to_datetime(df["Date"], format="%m/%d/%Y")
        last = df.iloc[int(np.argmax(d.values))]
        cols = ["1 Mo", "2 Mo", "3 Mo", "4 Mo", "6 Mo", "1 Yr", "2 Yr"]
        return CbSeries(bank, self.name, "curve", _pts(d, df["3 Mo"]),
                        {"asof": d.max().date(), "curve": {c: float(last[c]) for c in cols}},
                        "par yield curve CSV (current year); bills+notes, NOT OIS")


class EcbYcForward(CbSource):
    """ECB euro-area AAA govt instantaneous forwards (IF_*) — coarse, government (not OIS) proxy."""
    name = "ecb_yc_forward"
    section = "a3"
    banks = ("EUR",)

    def _fetch(self, bank):
        start = (date.today() - timedelta(days=12)).isoformat()
        r = self.get("https://data-api.ecb.europa.eu/service/data/YC/B.U2.EUR.4F.G_N_A.SV_C_YM."
                     f"?startPeriod={start}&format=csvdata")
        if r is None or self._check_botwall(r):
            return None
        df = pd.read_csv(io.StringIO(r.text))
        df["tenor"] = df["KEY"].str.split(".").str[-1]
        f = df[df.tenor.str.startswith("IF_")].copy()
        if f.empty:
            self._parse_fail("no IF_ tenors")
            return None

        def months(t):
            m = re.fullmatch(r"IF_(?:(\d+)Y)?(?:(\d+)M)?", t)
            return int(m.group(1) or 0) * 12 + int(m.group(2) or 0) if m else None
        f["m"] = f.tenor.map(months)
        last = f.TIME_PERIOD.max()
        cur = f[f.TIME_PERIOD == last].dropna(subset=["m"]).sort_values("m")
        rh = self.get("https://data-api.ecb.europa.eu/service/data/YC/B.U2.EUR.4F.G_N_A.SV_C_YM.IF_1Y"
                      "?startPeriod=2004-01-01&format=csvdata")
        dh = pd.read_csv(io.StringIO(rh.text)) if rh is not None else f[f.tenor == "IF_1Y"]
        pts = _pts(pd.to_datetime(dh.TIME_PERIOD), dh.OBS_VALUE)
        return CbSeries(bank, self.name, "curve", pts,
                        {"asof": date.fromisoformat(last), "curve": {int(m): float(v) for m, v in zip(cur.m, cur.OBS_VALUE)}},
                        "YC AAA govt instantaneous forward IF_*; series history since 2004; AAA govt != €STR OIS")


# Status-only probes of sources we expect to be blocked / JS-only / paid (negative evidence).
BLOCKED_TARGETS = [
    ("USD", "cme_fedwatch", "https://www.cmegroup.com/markets/interest-rates/cme-fedwatch-tool.html", "http"),
    ("USD", "yahoo_zq_chart", "https://query1.finance.yahoo.com/v8/finance/chart/ZQZ26.CBT?range=5d&interval=1d", "http"),
    ("EUR", "eurex_fesr_page", "https://www.eurex.com/ex-en/markets/int/mon/euro-short-term-rate-futures/euro-short-term-rate-futures-1404746", "js"),
    ("CHF", "eurex_fsr3_page", "https://www.eurex.com/ex-en/markets/int/mon/saron-futures/saron/3M-SARON-Futures-1405958", "js"),
    ("EUR", "ice_euribor_page", "https://www.ice.com/products/27994948/Three-Month-Euribor-Futures/data?marketId=5920503", "js"),
    ("CAD", "mx_quotes_page", "https://www.m-x.ca/en/trading/data/quotes?symbol=CRA", "js"),
    ("NZD", "rbnz_b2_xlsx", RbnzPolicy.URLS[0], "http"),
]


class BlockedProbe(CbSource):
    name = "blocked_evidence"
    section = "a3"
    banks = tuple(BANKS)
    official = False

    def probe_target(self, bank, name, url, kind) -> Probe:
        self.last_status = self.last_note = ""
        r = self.get(url, ua=None, retries=0)
        if r is None:
            return Probe(bank, "a3", name, "FAIL", self.failure_code(), None, self.last_note, cdn=self.cdn)
        txt = _text(r.text)
        hint = re.search(r"(?i)(statistics loading|tradingview loading|loading\.\.\.|enable javascript|qmod)", r.text)
        priced = len(re.findall(r"\b9\d\.\d{2,3}\b", txt))
        if kind == "js" and (hint or priced == 0):
            return Probe(bank, "a3", name, "FAIL", "JS_ONLY", None,
                         f"HTTP 200 but no prices in server HTML ({priced} price-like tokens; hint={hint.group(1) if hint else '-'})",
                         cdn=self.cdn)
        return Probe(bank, "a3", name, "PARTIAL", "", None, f"HTTP 200, {priced} price-like tokens", cdn=self.cdn)


# ---------------------------------------------------------------------------
# A3 — diagnostics on the fetched curves (no model: arithmetic only, assumptions stated)
# ---------------------------------------------------------------------------

def implied_1m_chain(months: dict, meetings: list, r0: float) -> list:
    """Bootstrap post-meeting overnight rates from 1M-average contracts.
    `months`: {(y,m): avg implied rate}; `meetings`: effective dates (first day of new rate).
    Assumes one meeting per month, constant overnight-target spread (folded into r0), and that
    the rate only changes on effective dates. Noise scales with N/(N-d_pre): late-month
    meetings are poorly identified."""
    out, prev = [], r0
    for eff in sorted(meetings):
        ym = (eff.year, eff.month)
        if ym not in months:
            continue
        n = calendar.monthrange(*ym)[1]
        d_pre = eff.day - 1
        if n - d_pre <= 3:
            out.append((eff, None, f"only {n - d_pre}d after effective date"))
            continue
        post = (n * months[ym] - d_pre * prev) / (n - d_pre)
        out.append((eff, post, ""))
        prev = post
    return out


def fwd_at(months: list, curve: list, t_months: float) -> float:
    return float(np.interp(t_months, months, curve))


# ---------------------------------------------------------------------------
# A4 — Forex Factory consensus (LOCAL data only)
# ---------------------------------------------------------------------------

FF_ID = {"USD": "usd_fed", "EUR": "eur_ecb", "GBP": "gbp_boe", "JPY": "jpy_boj",
         "CAD": "cad_boc", "AUD": "aud_rba", "NZD": "nzd_rbnz", "CHF": "chf_snb"}


class FfConsensus(CbSource):
    """Rows are joined to the OFFICIAL meeting calendar (decision day +/- 1d), not to raw FF days:
    FF also emits tentative/placeholder rows (BoJ 21:00 UTC rows, actual=0.0) around a meeting."""
    name = "ff_local_parquet"
    section = "a4"
    banks = tuple(BANKS)
    official = False
    RAW_NAME = {"USD": "Federal Funds Rate", "EUR": "Main Refinancing Rate", "GBP": "Official Bank Rate",
                "JPY": "BOJ Policy Rate", "CAD": "Overnight Rate", "AUD": "Cash Rate",
                "NZD": "Official Cash Rate", "CHF": "SNB Policy Rate"}

    def _load(self):
        if "pq" not in self._cache:
            self._cache["pq"] = pd.read_parquet(DATA_DIR / "economic_calendar_ff.parquet")
            rng = json.loads((DATA_DIR / "archive" / "ff_calendar_range.json").read_text())
            self._cache["rng"] = pd.DataFrame(rng)
        return self._cache["pq"], self._cache["rng"]

    def _fetch(self, bank):
        pq, rng = self._load()
        cid = f"{FF_ID[bank]}_interest_rate_decision"
        rows = pq[pq.canonical_id == cid].sort_values("datetime_utc").copy()
        rows["day"] = rows.datetime_utc.dt.date
        today = date.today()
        meetings = [m for m in official_meetings(bank) if m.decision <= today][-4:]
        dec = []
        for m in meetings:
            win = rows[(rows.day >= (m.first_day or m.decision) - timedelta(days=1)) & (rows.day <= m.decision + timedelta(days=1))]
            good = win[win.actual.notna() & win.released]
            pick = good.iloc[-1] if len(good) else (win.iloc[-1] if len(win) else None)
            zero = win[(win.actual == 0.0)]
            dec.append({"meeting": m.decision, "basis": m.basis, "n_rows": len(win),
                        "day": None if pick is None else pick.day,
                        "actual": None if pick is None else pick.actual,
                        "forecast": None if pick is None else pick.forecast,
                        "previous": None if pick is None else pick.previous,
                        "released": None if pick is None else bool(pick.released),
                        "zero_rows": [str(t) for t in zero.datetime_utc]})
        # marker events (0/0/0 rows for statements / press conferences) prove the meeting exists in FF
        markers = {}
        if bank == "CHF":
            mk = rng[(rng.Currency == "CHF") & rng.Name.isin(["SNB Monetary Policy Assessment", "SNB Press Conference"])]
            markers = {str(pd.to_datetime(d, format="%Y.%m.%d %H:%M:%S").date()): n
                       for d, n in zip(mk.Date, mk.Name)}
        rr = rng[(rng.Currency == bank) & (rng.Name == self.RAW_NAME[bank])]
        return CbSeries(bank, self.name, "ff",
                        [(d["day"], float(d["actual"])) for d in dec if d["actual"] is not None and pd.notna(d["actual"])],
                        {"decisions": dec, "n_rows": len(rows), "raw_names": sorted(set(rows.name_raw)),
                         "canonical_names": sorted(set(rows.name_canonical)),
                         "zero_rows_all": [str(t) for t in rows[rows.actual == 0.0].datetime_utc],
                         "range_zero_prev_pos": int(((rr.Actual == 0.0) & (rr.Previous > 0)).sum()) if len(rr) else 0,
                         "range_json_rows": int(len(rr)), "range_json_last": rr.Date.max() if len(rr) else None,
                         "last_row_day": rows.day.max() if len(rows) else None, "markers": markers,
                         "can_be_zero": _can_be_zero("interest_rate_decision")},
                        f"canonical_id={cid}")


def _can_be_zero(key: str):
    try:
        import yaml
        return bool(yaml.safe_load((DATA_DIR / "economic_indicators.yaml").read_text())["indicators"][key].get("can_be_zero"))
    except Exception:
        return None


# ---------------------------------------------------------------------------
# A5 — official calendar (2026-2027)
# ---------------------------------------------------------------------------

@dataclass
class Meeting:
    decision: date
    first_day: Optional[date] = None
    projections: Optional[bool] = None    # SEP / staff projections / MPR / SMP / MPS / Outlook Report
    presser: Optional[bool] = None
    basis: str = "page"                   # page | derived


def _dm(day: int, mon: str, year: int) -> date:
    return date(year, MONTHS[mon.lower().rstrip(".")], day)


class CalendarSource(CbSource):
    section = "a5"
    kind = "calendar"
    URL = ""

    def _page(self):
        r = self.get(self.URL, ua=None)
        return r.text if r is not None else None

    def _fetch(self, bank):
        html = self._page()
        if html is None:
            return None
        ms = [m for m in self.parse(html) if 2025 <= m.decision.year <= 2027]
        if not ms:
            self._parse_fail("no 2026-2027 meetings parsed")
            return None
        ms.sort(key=lambda m: m.decision)
        return CbSeries(bank, self.name, "calendar", [(m.decision, 1.0) for m in ms],
                        {"meetings": ms, "url": self.URL}, self.URL)

    def parse(self, html: str) -> list[Meeting]:  # overridden
        return []


class FedCalendar(CalendarSource):
    name = "fed_fomccalendars"
    banks = ("USD",)
    URL = "https://www.federalreserve.gov/monetarypolicy/fomccalendars.htm"

    def parse(self, html):
        txt = _text(html)
        out = []
        for m in re.finditer(r"(\d{4}) FOMC Meetings(.*?)(?=\d{4} FOMC Meetings|\* Meeting associated|\Z)", txt):
            year, blk = int(m.group(1)), m.group(2)
            for mm in re.finditer(r"(January|February|March|April|May|June|July|August|September|October|November|December)"
                                  r" (\d{1,2})(?:-(\d{1,2}))?(\*?)(?![\d,])", blk):
                mon, d1, d2, star = mm.groups()
                last = int(d2 or d1)
                seg = blk[mm.end():mm.end() + 260]
                nxt = re.search(r"(January|February|March|April|May|June|July|August|September|October|November|December) \d", seg)
                seg = seg[:nxt.start()] if nxt else seg
                dec = _dm(last, mon, year)
                out.append(Meeting(dec, _dm(int(d1), mon, year), bool(star),
                                   True if "Press Conference" in seg else (None if dec > date.today() else False),
                                   "page" if dec <= date.today() else "page(projections)/presser unannounced"))
        return out


class EcbCalendar(CalendarSource):
    name = "ecb_gc_calendar"
    banks = ("EUR",)
    URL = "https://www.ecb.europa.eu/press/calendars/mgcgc/html/index.en.html"

    def parse(self, html):
        txt = _text(html)
        out = []
        for d, mth in re.findall(r"(\d{2}/\d{2}/\d{4}) Governing Council of the ECB: monetary policy meeting[^0-9]{0,60}\(Day 2\)"
                                 r"(, followed by press conference)?", txt):
            dd = datetime.strptime(d, "%d/%m/%Y").date()
            # staff projections accompany the Mar/Jun/Sep/Dec meetings (ECB practice, not on this page)
            out.append(Meeting(dd, dd - timedelta(days=1), dd.month in (3, 6, 9, 12), bool(mth), "page(presser)/derived(proj)"))
        return out


class BoeCalendar(CalendarSource):
    name = "boe_mpc_dates"
    banks = ("GBP",)
    URL = "https://www.bankofengland.co.uk/monetary-policy/upcoming-mpc-dates"

    def parse(self, html):
        txt = _text(html)
        out = []
        for year, blk in (("2026", txt[txt.find("2026 confirmed dates"):txt.find("2027 provisional dates")]),
                          ("2027", txt[txt.find("2027 provisional dates"):])):
            for mm in re.finditer(r"(?:Monday|Tuesday|Wednesday|Thursday|Friday) (\d{1,2}) ([A-Z][a-z]+) (.*?)"
                                  r"(?=(?:Monday|Tuesday|Wednesday|Thursday|Friday) \d{1,2} [A-Z][a-z]+ |Current Bank Rate|"
                                  r"Monetary Policy Committee voting|2027 provisional|$)", blk):
                d, mon, desc = mm.groups()
                mpr = "Monetary Policy Report" in desc
                out.append(Meeting(_dm(int(d), mon, int(year)), None, mpr, mpr, "page(mpr)/derived(presser=MPR days)"))
        return out


class BojCalendar(CalendarSource):
    name = "boj_mpm_schedule"
    banks = ("JPY",)
    URL = "https://www.boj.or.jp/en/mopo/mpmsche_minu/index.htm"

    def parse(self, html):
        out = []
        for year, tb in zip((2026, 2027), [t for t in _tables(html) if t and t[0] and "MPM" in " ".join(t[0])][:2]):
            for row in tb[1:]:
                if not row or not re.match(r"[A-Z][a-z]{2,4}\.? \d", row[0]):
                    continue
                mm = re.match(r"([A-Z][a-z]{2,4})\.? (\d{1,2}) \([^)]*\)(?:, (\d{1,2}))?", row[0])
                if not mm:
                    continue
                mon, d1, d2 = mm.groups()
                outlook = len(row) > 1 and bool(re.match(r"[A-Z][a-z]{2,4}\.? \d", row[1]))
                out.append(Meeting(_dm(int(d2 or d1), mon, year), _dm(int(d1), mon, year), outlook, None,
                                   "page(outlook)/presser not on page"))
        return out


class BocCalendar(CalendarSource):
    name = "boc_schedule"
    banks = ("CAD",)
    URL = "https://www.bankofcanada.ca/core-functions/monetary-policy/key-interest-rate/"

    def parse(self, html):
        txt = _text(html)
        out = []
        for m in re.finditer(r"Schedule for (\d{4}) Dates Publications(.*?)(?=Schedule for \d{4}|\Z)", txt):
            year = int(m.group(1))
            for mm in re.finditer(r"([A-Z][a-z]+) (\d{1,2}) Interest rate announcement( and Monetary Policy Report)?", m.group(2)):
                mon, d, mpr = mm.groups()
                out.append(Meeting(_dm(int(d), mon, year), None, bool(mpr), True, "page(mpr) + blackout page: press conference at every decision"))
        return out


class RbaCalendar(CalendarSource):
    name = "rba_board_schedule"
    banks = ("AUD",)
    URL = "https://www.rba.gov.au/schedules-events/board-meeting-schedules.html"

    def _page(self):
        r = self.get(self.URL)          # NB_UA
        return r.text if r is not None else None

    def parse(self, html):
        txt = _text(html)
        out = []
        for m in re.finditer(r"Board meeting schedules (\d{4}) Month Monetary Policy Board(.*?)(?=Board meeting schedules \d{4}|Related Information|\Z)", txt):
            year = int(m.group(1))
            for mm in re.finditer(r"(\d{1,2})[–-](\d{1,2}) ([A-Z][a-z]+)", m.group(2)):
                d1, d2, mon = mm.groups()
                mo = MONTHS[mon.lower()]
                # Statement on Monetary Policy accompanies Feb/May/Aug/Nov (RBA practice, not on this page)
                out.append(Meeting(date(year, mo, int(d2)), date(year, mo, int(d1)), mo in (2, 5, 8, 11), None,
                                   "page(dates)/derived(SMP months)"))
        return out


class SnbCalendar(CalendarSource):
    name = "snb_event_schedule"
    banks = ("CHF",)
    URL = "https://www.snb.ch/en/services-events/digital-services/event-schedule"
    PAST = "https://www.snb.ch/en/the-snb/mandates-goals/monetary-policy/decisions"

    def parse(self, html):
        txt = _text(html)
        out = []
        past = self.get(self.PAST, ua=None)          # the schedule lists only upcoming events
        if past is not None:
            for d, mon, y in re.findall(r"Monetary policy assessment of (\d{1,2}) ([A-Z][a-z]+) (\d{4})", _text(past.text)):
                out.append(Meeting(_dm(int(d), mon, int(y)), None, True, None, "page(decisions archive)/derived(proj)"))
        for d in re.findall(r"(\d{2}\.\d{2}\.\d{4}) \d{2}:\d{2} Monetary policy assessment of [^()]*\(press release\)", txt):
            dd = datetime.strptime(d, "%d.%m.%Y").date()
            pc = bool(re.search(rf"{re.escape(d)} \d{{2}}:\d{{2}} Monetary policy assessment of [^()]*\(introductory remarks, news conference\)", txt))
            out.append(Meeting(dd, None, True, pc, "page(presser)/derived(conditional forecast every quarter)"))
        return out


class RbnzCalendar(CalendarSource):
    """RBNZ site is behind a Cloudflare challenge (HTTP 403 for every UA, also via WebFetch) —
    dates are MANUAL, transcribed from RBNZ's own release 'Monetary policy and OCR decision dates
    until February 2027' as surfaced by web search; verify by hand."""
    name = "rbnz_manual"
    banks = ("NZD",)
    URL = "https://www.rbnz.govt.nz/news-and-events/how-we-release-information/monetary-policy-and-ocr-decision-dates-until-february-2027"
    MANUAL = [(date(2026, 2, 18), True), (date(2026, 4, 8), False), (date(2026, 5, 27), True),
              (date(2026, 7, 8), False), (date(2026, 9, 2), True), (date(2026, 10, 28), False),
              (date(2026, 12, 9), True), (date(2027, 2, 17), True)]

    def _fetch(self, bank):
        r = self.get(self.URL, ua=None, retries=0)       # attempt the official page first
        blocked = r is None
        ms = [Meeting(d, None, mps, True if mps else None, "manual") for d, mps in self.MANUAL]
        note = ("official page BOTWALL (" + self.last_note + ") — manual list; 2027 dates after 17 Feb not yet published"
                if blocked else "official page reachable — parser not built; manual list shown")
        return CbSeries(bank, self.name, "calendar", [(m.decision, 1.0) for m in ms], {"meetings": ms, "url": self.URL}, note)


CALENDARS = [FedCalendar, EcbCalendar, BoeCalendar, BojCalendar, BocCalendar, RbaCalendar, SnbCalendar, RbnzCalendar]


QUIET = {   # bank: (rule text, official URL, regex proving the wording on the fetched page | None)
    "USD": ("blackout: 00:00 ET on the second Saturday before the meeting -> 23:59 ET the day after it ends",
            "https://www.federalreserve.gov/monetarypolicy/files/FOMC_ExtCommunicationParticipants.pdf", None),
    "EUR": ("quiet period: 7 days before each monetary-policy meeting",
            "https://www.ecb.europa.eu/ecb-and-you/explainers/tell-me/html/what-is-the_quiet_period.en.html", r"seven days"),
    "GBP": ("MPC quiet period: from start of deliberations (usually 8-9 days before) until announcement",
            "https://www.bankofengland.co.uk/about/governance-and-funding/mpc-external-communications-code", r"(?i)quiet period"),
    "JPY": ("blackout: 2 business days before day 1 of the MPM -> end of last MPM day (wording on an older BoJ notice; verify current)",
            "https://www.boj.or.jp/en/mopo/mpmsche_minu/m_ref/mpm0104a.htm", r"starts two business days before"),
    "CAD": ("blackout: Tuesday 8 days before in MPR months (Jan/Apr/Jul/Oct), else Wednesday 7 days before; ends 10:30 ET when the press conference starts",
            "https://www.bankofcanada.ca/core-functions/monetary-policy/key-interest-rate/blackout-guidelines/",
            r"blackout begins on Wednesday . seven days prior"),
    "AUD": ("external MPB members: from 14:00 Sydney Wednesday before the meeting -> minutes; trading blackout to 17:00 on decision day",
            "https://www.rba.gov.au/about-rba/our-policies/monetary-policy-board-external-members-comms-and-public-engagement-policy.html",
            r"2\.00 pm on the Wednesday prior"),
    "NZD": ("NOT FOUND — RBNZ site bot-walled; no quiet-period document located via search", "", None),
    "CHF": ("NOT FOUND — no quiet-period document located on snb.ch via search", "", None),
}


class QuietPeriod(CbSource):
    name = "quiet_period_rules"
    section = "a5"
    banks = tuple(BANKS)

    def _fetch(self, bank):
        rule, url, rx = QUIET[bank]
        if not url:
            self.last_status, self.last_note = "UNVERIFIED", rule
            return None
        checked = "official PDF (wording from search excerpt, not parsed)"
        if rx:
            r = self.get(url, ua=NB_UA)
            if r is None:
                return None
            if not re.search(rx, _text(r.text)):
                self.last_status, self.last_note = "UNVERIFIED", f"phrase /{rx}/ not found on {url}"
                return None
            checked = f"phrase /{rx}/ present on fetched page"
        return CbSeries(bank, self.name, "quiet", [], {"rule": rule, "url": url}, checked)


# ---------------------------------------------------------------------------
# 0B — TEXT sources (B1-B5). Stdlib-only mini-DOM (html.parser) for deterministic extraction:
# fixed per-bank container selectors, no heuristics, no new dependency.
# ---------------------------------------------------------------------------

_VOID = {"meta", "br", "img", "input", "link", "hr", "area", "base", "col", "embed", "source", "track", "wbr"}
_SKIP = {"script", "style", "noscript", "svg", "template", "iframe", "head"}


class _Node:
    __slots__ = ("tag", "attrs", "kids", "parent")

    def __init__(self, tag, attrs, parent):
        self.tag, self.attrs, self.kids, self.parent = tag, dict(attrs), [], parent

    @property
    def cls(self) -> list[str]:
        return (self.attrs.get("class") or "").split()

    def walk(self):
        yield self
        for k in self.kids:
            if isinstance(k, _Node):
                yield from k.walk()


class _DomBuilder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = _Node("root", [], None)
        self.cur = self.root

    def handle_starttag(self, tag, attrs):
        n = _Node(tag, attrs, self.cur)
        self.cur.kids.append(n)
        if tag not in _VOID:
            self.cur = n

    def handle_endtag(self, tag):
        c = self.cur
        while c is not None and c.tag != tag:
            c = c.parent
        if c is not None and c.parent is not None:
            self.cur = c.parent

    def handle_data(self, data):
        if data:
            self.cur.kids.append(data)


def _dom(html: str) -> _Node:
    b = _DomBuilder()
    b.feed(html)
    return b.root


def _node_text(n) -> str:
    if isinstance(n, str):
        return n
    if n.tag in _SKIP:
        return ""
    return "".join(_node_text(k) for k in n.kids)


def _norm(t: str) -> str:
    return re.sub(r"\s+", " ", t).strip()


def _paras(n, tags=("p", "li", "h1", "h2", "h3", "h4", "blockquote"), loose: bool = False) -> list[str]:
    """Text blocks in document order. loose=True also emits runs of bare text sitting directly in a
    container (e.g. the Fed's 'approved the following statement ... by a 12-0 vote:' line)."""
    out: list[str] = []

    def walk(x):
        if isinstance(x, str) or x.tag in _SKIP:
            return
        if x.tag in tags:
            t = _norm(_node_text(x))
            if t:
                out.append(t)
            return
        buf: list[str] = []
        for k in x.kids:
            if isinstance(k, str):
                if loose:
                    buf.append(k)
                continue
            if loose and buf:
                t = _norm("".join(buf))
                if t:
                    out.append(t)
                buf = []
            walk(k)
        if loose and buf:
            t = _norm("".join(buf))
            if t:
                out.append(t)
    walk(n)
    return out


def _html_text(r) -> str:
    """Decode an HTTP response as UTF-8 unless the server declared another charset
    (requests would otherwise guess ISO-8859-1 for text/html and mangle en-dashes)."""
    if "charset" not in r.headers.get("Content-Type", "").lower():
        try:
            return r.content.decode("utf-8")
        except UnicodeDecodeError:
            pass
    return r.text


def _pick(root: _Node, tag: str, id_: Optional[str] = None, cls: Optional[str] = None) -> Optional[_Node]:
    """Largest-text node matching the selector (tag, #id, .class)."""
    best, size = None, -1
    for n in root.walk():
        if n.tag != tag or (id_ and n.attrs.get("id") != id_) or (cls and cls not in n.cls):
            continue
        s = len(_node_text(n))
        if s > size:
            best, size = n, s
    return best


BOILERPLATE = ("Skip to", "Cookie", "cookies", "Subscribe", "Share this", "Privacy", "Sign up", "Follow us", "Back to top",
               "Contact us", "Site map", "Search")


def _feed_items(content: bytes) -> list[dict]:
    """RSS 2.0 / Atom / RDF -> [{title, link, pub(datetime|None), raw, author, desc}]."""
    try:
        root = ET.fromstring(content)
    except ET.ParseError:
        return []
    out = []
    for it in root.iter():
        if it.tag.split("}")[-1] not in ("item", "entry"):
            continue

        def g(tag):
            return next((_norm(e.text) for e in it if e.tag.split("}")[-1] == tag and (e.text or "").strip()), "")
        link = g("link") or next((e.get("href") for e in it if e.tag.split("}")[-1] == "link" and e.get("href")), "")
        raw = g("pubDate") or g("published") or g("updated") or g("date")
        out.append({"title": re.sub(r"<!\[CDATA\[|\]\]>", "", g("title")), "link": link, "raw": raw,
                    "pub": _parse_dt(raw), "author": g("creator") or g("author"), "desc": g("description")})
    return out


def _parse_dt(s: str) -> Optional[datetime]:
    if not s:
        return None
    try:
        d = parsedate_to_datetime(s)
    except (TypeError, ValueError):
        try:
            d = datetime.fromisoformat(s.replace("Z", "+00:00"))
        except ValueError:
            return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


# Official release time of the decision (local clock, DST-aware) — from bank pages / feeds, see report.
OFFICIAL_TIME = {"USD": (dtime(14, 0), "America/New_York"), "EUR": (dtime(14, 15), "Europe/Berlin"),
                 "GBP": (dtime(12, 0), "Europe/London"), "JPY": (None, "Asia/Tokyo"),
                 "CAD": (dtime(9, 45), "America/Toronto"), "AUD": (dtime(14, 30), "Australia/Sydney"),
                 "NZD": (dtime(14, 0), "Pacific/Auckland"), "CHF": (dtime(9, 30), "Europe/Zurich")}


def official_dt(bank: str, day: date) -> Optional[datetime]:
    t, tz = OFFICIAL_TIME[bank]
    return None if t is None else datetime.combine(day, t, tzinfo=ZoneInfo(tz))


def _mins(a: Optional[datetime], b: Optional[datetime]) -> Optional[float]:
    return None if a is None or b is None else round((a - b).total_seconds() / 60, 1)


ECB_SEEDS = {   # hashed ECB URLs cannot be built from a date — seeds found via web search; new ones come from the press feed
    date(2026, 4, 30): "https://www.ecb.europa.eu/press/pr/date/2026/html/ecb.mp260430~81b7179e6f.en.html",
    date(2026, 6, 11): "https://www.ecb.europa.eu/press/pr/date/2026/html/ecb.mp260611~4d41bd5e83.en.html",
    date(2026, 7, 23): "https://www.ecb.europa.eu/press/pr/date/2026/html/ecb.mp260723~29f24d99bc.en.html",
    date(2026, 9, 10): "https://www.ecb.europa.eu/press/pr/date/2026/html/ecb.mp260910~314e508016.en.html",
}
ECB_MPS = {     # monetary policy statement with Q&A (`ecb.is…`) — same limitation
    date(2026, 6, 11): "https://www.ecb.europa.eu/press/press_conference/monetary-policy-statement/2026/html/ecb.is260611~372040d313.en.html",
    date(2026, 9, 10): "https://www.ecb.europa.eu/press/press_conference/monetary-policy-statement/2026/html/ecb.is260910~6a45359cfc.en.html",
}
ECB_ACCOUNTS = {date(2026, 7, 23): "https://www.ecb.europa.eu/press/accounts/2026/html/ecb.mg260827~f06c21fd54.en.html"}

# Fixed extraction rules per bank: container selector (tag, #id, .class) + optional paragraph cut markers.
STATEMENT_RULES = {
    "USD": {"sel": ("div", "article", None), "stop": r"^(For media inquiries|Implementation Note)", "loose": True},
    "EUR": {"sel": ("div", None, "section"), "stop": None},
    "GBP": {"sel": ("div", "output", None), "stop": r"^Minutes of the Monetary Policy Committee meeting"},
    "CAD": {"sel": ("div", None, "post-content"), "stop": None},
    "AUD": {"sel": ("div", None, "rss-mr-content"), "stop": None},
    "CHF": {"sel": ("div", None, "a-text"), "stop": None},
}


def _statement_url(bank: str, d: date, feeds: dict) -> Optional[tuple[str, Optional[str]]]:
    """(url, ua) for the decision-day statement of `bank`; ua=None -> default browser UA."""
    if bank == "USD":
        return f"https://www.federalreserve.gov/newsevents/pressreleases/monetary{d:%Y%m%d}a.htm", None
    if bank == "EUR":
        u = ECB_SEEDS.get(d) or next((i["link"] for i in feeds.get("ecb", [])
                                      if f"ecb.mp{d:%y%m%d}" in i["link"]), None)
        return (u, None) if u else None
    if bank == "GBP":
        return f"https://www.bankofengland.co.uk/monetary-policy-summary-and-minutes/{d.year}/{d:%B}-{d.year}".lower(), None
    if bank == "JPY":
        return f"https://www.boj.or.jp/en/mopo/mpmdeci/mpr_{d.year}/k{d:%y%m%d}a.pdf", None
    if bank == "CAD":
        return f"https://www.bankofcanada.ca/{d.year}/{d:%m}/fad-press-release-{d:%Y-%m-%d}/", None
    if bank == "AUD":
        u = feeds.get("rba_decisions", {}).get(d)
        return (u, NB_UA) if u else None
    if bank == "NZD":
        return "https://www.rbnz.govt.nz/monetary-policy/monetary-policy-decisions", None
    if bank == "CHF":
        return f"https://www.snb.ch/en/publications/communication/press-releases-restricted/pre_{d:%Y%m%d}", None
    return None


class TextSource(CbSource):
    """Base for 0B probes: fetch a document and describe it (HTTP, headers, deterministic extraction)."""
    section = "b1"

    def describe(self, bank: str, kind: str, meeting: Optional[date], url: str, ua: Optional[str],
                 rule: Optional[dict] = None, feed_pub: Optional[datetime] = None) -> dict:
        self.last_status = self.last_note = ""
        rec = {"bank": bank, "kind": kind, "meeting": meeting, "url": url, "http": None, "fmt": "", "size": 0,
               "lm": None, "feed_pub": feed_pub, "paras": 0, "chars": 0, "sha": "", "first": "", "last": "",
               "boiler": None, "text": "", "code": "", "note": ""}
        r = self.get(url, ua=ua, retries=0)
        if r is None:
            rec["code"], rec["note"] = self.failure_code(), self.last_note
            rec["http"] = re.sub(r"\D", "", self.last_note or "") or None
            return rec
        rec["http"], rec["size"] = r.status_code, len(r.content)
        ctype = r.headers.get("Content-Type", "").lower()
        rec["lm"] = _parse_dt(r.headers.get("Last-Modified", ""))
        if "pdf" in ctype or url.lower().endswith(".pdf"):
            rec["fmt"] = "pdf"
            text = _pdf_text(r.content)
            if text is None:
                rec["code"], rec["note"] = "PDF_ONLY", "no PDF extractor in requirements.txt"
                return rec
            paras = [p for p in (_norm(x) for x in re.split(r"\n\s*\n", text)) if p]
        else:
            rec["fmt"] = "html"
            if rule is None:
                return rec
            root = _dom(_html_text(r))
            node = _pick(root, *rule["sel"])
            if node is None:
                rec["code"], rec["note"] = "NONE", f"selector {rule['sel']} not found"
                return rec
            paras = _paras(node, loose=bool(rule.get("loose")))
            if rule.get("stop"):
                cut = next((i for i, p in enumerate(paras) if re.match(rule["stop"], p)), None)
                paras = paras[:cut] if cut is not None else paras
        text = "\n".join(paras)
        rec.update(paras=len(paras), chars=len(text), sha=hashlib.sha1(text.encode()).hexdigest()[:8],
                   first=paras[0][:90] if paras else "", last=paras[-1][:90] if paras else "",
                   boiler=sum(1 for p in paras for b in BOILERPLATE if p.startswith(b)), text=text)
        return rec


def _pdf_text(content: bytes) -> Optional[str]:
    """Only if a PDF library happens to be importable; the project has none (see B8)."""
    try:
        from pypdf import PdfReader        # noqa: not in requirements.txt — proposal only
    except ImportError:
        return None
    return "\n\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(content)).pages)


VOTE_RX = {
    "USD": [r"approved the following statement for release by an? \d+\s*[–-]\s*\d+ vote", r"(?m)^Voting for the monetary policy action were.*$", r"(?m)^Voting against .*$"],
    "GBP": [r"voted (?:by a majority of \d+\s*[–-]\s*\d+|unanimously) to [^%]*%"],
    "JPY": [r"by an? \d+\s*[–-]\s*\d+ majority vote", r"by a unanimous vote"],
    "AUD": [r"[A-Za-z]+ members? voted in favour;.*?(?=\.(?:\s|$))", r"The Board decided unanimously.*?(?=\.(?:\s|$))"],
}


class StatementSource(TextSource):
    """B1 — decision statement of the last 4 meetings."""
    name = "statement_last4"
    section = "b1"
    banks = tuple(BANKS)

    def _feeds(self) -> dict:
        if "feeds" not in self._cache:
            f: dict = {}
            r = self.get("https://www.ecb.europa.eu/rss/press.html", ua=None)
            f["ecb"] = _feed_items(r.content) if r is not None else []
            r = self.get("https://www.federalreserve.gov/feeds/press_monetary.xml", ua=None)
            f["fed"] = _feed_items(r.content) if r is not None else []
            r = self.get("https://www.boj.or.jp/en/rss/whatsnew.xml", ua=None)
            f["boj"] = _feed_items(r.content) if r is not None else []
            r = self.get("https://www.snb.ch/public/rss/en/mopo", ua=None)
            f["snb"] = _feed_items(r.content) if r is not None else []
            r = self.get("https://www.rba.gov.au/monetary-policy/int-rate-decisions/2026/")
            f["rba_decisions"] = {}
            if r is not None:
                for li in re.findall(r"<li[^>]*>(.*?)</li>", r.text, flags=re.S):
                    m = re.search(r"href=\"(/media-releases/[^\"]+)\"[^>]*>\s*(\d{1,2} [A-Z][a-z]+ \d{4})", li, flags=re.S)
                    if m:
                        f["rba_decisions"][datetime.strptime(m.group(2), "%d %B %Y").date()] = "https://www.rba.gov.au" + m.group(1)
            self._cache["feeds"] = f
        return self._cache["feeds"]

    def _feed_pub(self, bank: str, url: str, d: date, feeds: dict) -> Optional[datetime]:
        pool = {"USD": feeds["fed"], "EUR": feeds["ecb"], "JPY": feeds["boj"], "CHF": feeds["snb"]}.get(bank, [])
        for it in pool:
            if url and (url.rsplit("/", 1)[-1].split("~")[0] in it["link"] or url.rsplit("/", 1)[-1] in it["link"]):
                return it["pub"]
        return None

    def _fetch(self, bank):
        feeds = self._feeds()
        ms = [m.decision for m in official_meetings(bank) if m.decision <= date.today()][-4:]
        docs = []
        for d in ms:
            tgt = _statement_url(bank, d, feeds)
            if tgt is None:
                docs.append({"bank": bank, "kind": "statement", "meeting": d, "url": "", "http": None, "code": "NONE",
                             "note": "URL not derivable from the date (hash) and not in the feed", "text": "", "fmt": ""})
                continue
            url, ua = tgt
            rec = self.describe(bank, "statement", d, url, ua, STATEMENT_RULES.get(bank), self._feed_pub(bank, url, d, feeds))
            rec["official"] = official_dt(bank, d)
            docs.append(rec)
        # stability across meetings: similarity of consecutive statements (redline feasibility)
        prev = None
        for rec in docs:
            rec["sim_prev"] = None
            if prev is not None and rec.get("text") and prev.get("text"):
                rec["sim_prev"] = round(SequenceMatcher(None, prev["text"][:9000], rec["text"][:9000]).ratio(), 2)
            prev = rec if rec.get("text") else prev
        return CbSeries(bank, self.name, "text", [], {"docs": docs}, "decision statement, last 4 meetings")


def _fmt_t(x: Optional[float]) -> str:
    return "—" if x is None else f"{x:+.1f}m"


def report_b1(probes: list[Probe], today: date) -> None:
    print("\n[B1] DECISION STATEMENT — last 4 meetings (deterministic extraction; latency = evidence − official release time)")
    for p in probes:
        if p.series is None:
            print(f"  {p.bank} FAIL {p.code} {p.note}")
            continue
        for d in p.series.extra["docs"]:
            off = d.get("official")
            lat_lm = _fmt_t(_mins(d.get("lm"), off))
            lat_feed = _fmt_t(_mins(d.get("feed_pub"), off))
            tail = d["url"].rsplit("/", 1)[-1][:40] if d["url"] else "(no url)"
            print(f"  {p.bank} {d['meeting']} http={d['http']} {d['fmt'] or '-':4} {d['code'] or 'ok':8} paras={d.get('paras', 0):3} "
                  f"chars={d.get('chars', 0):6} boiler={d.get('boiler')} sim_prev={d.get('sim_prev')} "
                  f"LM={d['lm'].strftime('%m-%d %H:%M') if d.get('lm') else '—'}Z LM−off={lat_lm} feed−off={lat_feed} {tail}")
            if d.get("first"):
                print(f"        first: {d['first']!r}")
                print(f"        last : {d['last']!r}")


def report_b2(probes: list[Probe], today: date) -> None:
    print("\n[B2] VOTES / DISSENTS — parsed from the statement text of the last 4 meetings")
    for p in probes:
        if p.series is None:
            print(f"  {p.bank} FAIL {p.code} {p.note}")
            continue
        for v in p.series.extra["votes"]:
            print(f"  {p.bank} {v['meeting']} {v['where']:<26} {v['vote'][:170]}")


class VoteSource(TextSource):
    """B2 — where votes/dissents appear and their format, parsed off the B1 documents (+ minutes for RBA)."""
    name = "votes_last4"
    section = "b2"
    banks = tuple(BANKS)
    WHERE = {"USD": "statement (Voting for/against)", "GBP": "MPC summary (+ minutes: per member)",
             "JPY": "statement PDF (majority + dissenters)", "AUD": "MPB minutes (2 weeks later)",
             "EUR": "not published (consensus; account)", "CAD": "not published (consensus; summary of deliberations)",
             "CHF": "not published", "NZD": "MPC record of meeting (site bot-walled)"}

    def _fetch(self, bank):
        if bank in ("EUR", "CAD", "CHF"):
            return CbSeries(bank, self.name, "votes", [], {"votes": [
                {"meeting": None, "where": self.WHERE[bank], "vote": "NOT PUBLISHED (no vote count in statement or minutes-equivalent)"}]},
                "no votes published")
        if bank == "NZD":
            self.last_status, self.last_note = "BOT-WALL", "HTTP 403"
            return None
        st = StatementSource()
        s = st.fetch(bank)
        if s is None:
            return None
        votes = []
        for d in s.extra["docs"]:
            txt = d.get("text") or ""
            if bank == "AUD" and d.get("meeting"):      # votes live in the minutes, not the release
                mu = f"https://www.rba.gov.au/monetary-policy/rba-board-minutes/{d['meeting'].year}/{d['meeting']:%Y-%m-%d}.html"
                r = self.get(mu, ua=NB_UA, retries=0)
                txt = _text(_html_text(r)) if r is not None else ""
            hits = [_norm(m.group(0)) for rx in VOTE_RX[bank] for m in [re.search(rx, txt)] if m] if txt else []
            hits = [(f"Voting for ({h.split('Voting against')[0].count(';') + 1} names)"
                     + (f"; against: {h.split('Voting against', 1)[1][:130]}" if "Voting against" in h else "")) if h.startswith("Voting for") else h[:170]
                    for h in hits]
            votes.append({"meeting": d["meeting"], "where": self.WHERE[bank],
                          "vote": " | ".join(hits) if hits else ("PDF not parsed (no extractor)" if d.get("code") == "PDF_ONLY" else "not found in text")})
        return CbSeries(bank, self.name, "votes", [], {"votes": votes}, self.WHERE[bank])



# ---------------------------------------------------------------------------
# B3 — press conference (video + transcript), B4 — minutes/accounts/summaries
# ---------------------------------------------------------------------------

YT_CHANNEL = {   # official channels; ids resolved from the bank's own site link (RBA: youtube.com/RBAInfo, RBNZ: /user/reservebankofnz)
    "USD": "UCAzhpt9DmG6PnHXjmJTvRGQ", "EUR": "UCXB8fM4VyQubRu3UVGhd3wA", "GBP": "UCY70MMJ8Rj4wtwt7bVN6hIA",
    "JPY": "UC32Yu7NyStgmKYsXvYofPvQ", "CAD": "UCY4EvEbIox0M4JuEsu5OKnQ", "AUD": "UCaLmgAMEglL-yGvuGzTx6ig",
    "NZD": "UC1v0CeB83SX6Px9r5KT0NuQ", "CHF": "UC4vQTVEqtj2orppzBkdGmyg"}
YT_FEED = "https://www.youtube.com/feeds/videos.xml?channel_id={cid}"
PRESS_RX = re.compile(r"press conference|media conference|news conference|conférence de presse|mediengespräch|記者会見|"
                      r"introductory statement|pooled broadcast", re.I)
# Does the bank hold a press conference at every rate meeting? (official wording / practice, with the source)
PRESSER_HELD = {
    "USD": "every meeting (8/8) — transcript PDF exists for each of the last 4",
    "EUR": "every meeting (8/8) — statement + Q&A on the ECB site",
    "GBP": "MPR meetings only (Feb/Apr/Jul/Nov); other meetings = Governor's pooled broadcast interview + transcript (BoE news feed)",
    "JPY": "every MPM, ~15:30 JST on the last day (BoJ site); English transcript NOT published (JP summary next business day)",
    "CAD": "every decision (BoC blackout page: 'blackout ends … when the press conference begins')",
    "AUD": "every meeting, 15:30 AEST (RBA media-conference page); audio + transcript page mc-gov-YYYY-MM-DD",
    "NZD": "every OCR announcement (MPS and MPR), 15:00 NZ (RBNZ, via search — site bot-walled)",
    "CHF": "quarterly assessments only (4/yr); introductory remarks published, no Q&A transcript",
}


def _transcript_target(bank: str, d: date) -> Optional[tuple[str, Optional[str], dict]]:
    body = {"sel": ("body", None, None)}
    if bank == "USD":
        return f"https://www.federalreserve.gov/mediacenter/files/FOMCpresconf{d:%Y%m%d}.pdf", None, body
    if bank == "EUR":
        u = ECB_MPS.get(d)
        return (u, None, {"sel": ("div", None, "section")}) if u else None
    if bank == "CAD":
        return f"https://www.bankofcanada.ca/{d.year}/{d:%m}/opening-statement-{d:%Y-%m-%d}/", None, {"sel": ("div", None, "post-content")}
    if bank == "AUD":
        return f"https://www.rba.gov.au/speeches/{d.year}/mc-gov-{d:%Y-%m-%d}.html", NB_UA, {"sel": ("div", "content", None)}
    if bank == "CHF":
        return (f"https://www.snb.ch/en/publications/communication/speeches-restricted/ref_{d:%Y%m%d}_mslanmargpe", None,
                {"sel": ("div", None, "a-text")})
    return None


class PressConfSource(TextSource):
    """B3 — is there a press conference, where is the official video (YouTube feed, keyless) and the transcript."""
    name = "presser_last4"
    section = "b3"
    banks = tuple(BANKS)

    def _fetch(self, bank):
        ms = [m.decision for m in official_meetings(bank) if m.decision <= date.today()][-4:]
        cid = YT_CHANNEL[bank]
        r = self.get(YT_FEED.format(cid=cid), ua=None, retries=0)
        yt = _feed_items(r.content) if r is not None else []
        rows = []
        for d in ms:
            off = official_dt(bank, d)
            vids = [v for v in yt if v["pub"] and PRESS_RX.search(v["title"]) and 0 <= (v["pub"].date() - d).days <= 3]
            v = min(vids, key=lambda x: x["pub"]) if vids else None
            tr = _transcript_target(bank, d)
            trec = self.describe(bank, "transcript", d, tr[0], tr[1], tr[2]) if tr else None
            rows.append({"meeting": d, "video": v, "video_lag_h": None if v is None or off is None else round((v["pub"] - off).total_seconds() / 3600, 1),
                         "transcript": trec})
        oldest = min((v["pub"] for v in yt if v["pub"]), default=None)
        return CbSeries(bank, self.name, "presser", [], {"rows": rows, "yt_items": len(yt), "yt_oldest": oldest, "channel": cid,
                                                        "held": PRESSER_HELD[bank]}, PRESSER_HELD[bank])


def report_b3(probes: list[Probe], today: date) -> None:
    print("\n[B3] PRESS CONFERENCE — official video (YouTube feed, keyless) + transcript, last 4 meetings")
    for p in probes:
        if p.series is None:
            print(f"  {p.bank} FAIL {p.code} {p.note}")
            continue
        x = p.series.extra
        print(f"  {p.bank} channel_id={x['channel']} feed={x['yt_items']} entries (oldest {x['yt_oldest']:%Y-%m-%d}) | {x['held']}")
        for r in x["rows"]:
            v, t = r["video"], r["transcript"]
            vs = (f"video {v['pub']:%m-%d %H:%MZ}" + (f" (+{r['video_lag_h']}h)" if r['video_lag_h'] is not None else " (official time variable)") + f" '{v['title'][:44]}'") if v else "video: none in feed window"
            ts = "transcript: none derived" if t is None else f"transcript http={t['http']} {t['fmt'] or '-'} {t['code'] or 'ok'} LM={t['lm']:%m-%d %H:%MZ}" if t.get("lm") else \
                f"transcript http={t['http']} {t['fmt'] or '-'} {t['code'] or 'ok'} chars={t.get('chars', 0)}"
            print(f"      {r['meeting']}  {vs} | {ts}")


class MinutesSource(TextSource):
    """B4 — minutes / accounts / summary of opinions / summary of deliberations: source and lag, last 4 meetings."""
    name = "minutes_last4"
    section = "b4"
    banks = tuple(BANKS)
    KIND = {"USD": "minutes (HTML/PDF)", "EUR": "account of the meeting", "GBP": "minutes (same page as summary)",
            "JPY": "Summary of Opinions (+ Minutes ~2 months later)", "CAD": "Summary of Governing Council deliberations",
            "AUD": "MPB minutes", "NZD": "MPC record of meeting", "CHF": "Summary of monetary policy discussion"}

    def _fetch(self, bank):
        ms = [m.decision for m in official_meetings(bank) if m.decision <= date.today()][-4:]
        rows = []
        if bank == "USD":
            r = self.get("https://www.federalreserve.gov/monetarypolicy/fomccalendars.htm", ua=None)
            txt = _text(_html_text(r)) if r is not None else ""
            rel = {}
            for m in re.finditer(r"([A-Z][a-z]+) (\d{1,2})-(\d{1,2})\*?[^()]{0,260}?\(Released ([A-Z][a-z]+ \d{2}, \d{4})\)", txt):
                mo, d1, d2, rd = m.groups()
                rel[date(2026, MONTHS[mo.lower()], int(d2))] = datetime.strptime(rd, "%B %d, %Y").date()
            for d in ms:
                u = f"https://www.federalreserve.gov/monetarypolicy/fomcminutes{d:%Y%m%d}.htm"
                rec = self.describe(bank, "minutes", d, u, None, {"sel": ("div", "article", None)}) if rel.get(d) else None
                rows.append({"meeting": d, "released": rel.get(d), "url": u, "rec": rec})
        elif bank == "EUR":
            for d in ms:
                u = ECB_ACCOUNTS.get(d)
                rel = None
                if u:
                    mm = re.search(r"ecb\.mg(\d{6})", u)
                    rel = datetime.strptime(mm.group(1), "%y%m%d").date() if mm else None
                rows.append({"meeting": d, "released": rel, "url": u or "", "rec": self.describe(bank, "account", d, u, None, {"sel": ("div", None, "section")}) if u else None,
                             "note": "" if u else "URL has a hash — not derivable from the date"})
        elif bank == "GBP":
            for d in ms:
                u = _statement_url("GBP", d, {})[0]
                r = self.get(u, ua=None, retries=0)
                has = r is not None and "Minutes of the Monetary Policy Committee meeting" in _html_text(r)
                rows.append({"meeting": d, "released": d if has else None, "url": u, "rec": None})
        elif bank == "JPY":
            r = self.get(BojCalendar.URL, ua=None)
            tabs = _tables(_html_text(r)) if r is not None else []
            tb = next((t for t in tabs if t and t[0] and "MPM" in " ".join(t[0])), [])
            byd = {}
            for row in tb[1:]:
                mm = re.match(r"([A-Z][a-z]{2,4})\.? (\d{1,2}) \([^)]*\)(?:, (\d{1,2}))?", row[0]) if row else None
                if mm and len(row) >= 4:
                    dd = _dm(int(mm.group(3) or mm.group(2)), mm.group(1), 2026)
                    dm = [re.match(r"([A-Z][a-z]{2,4})\.? (\d{1,2})", c) for c in row[-2:]]
                    rd = [_dm(int(x.group(2)), x.group(1), 2026 if MONTHS[x.group(1).lower().rstrip('.')] >= dd.month else 2027) if x else None for x in dm]
                    byd[dd] = rd
            for d in ms:
                rd = byd.get(d, [None, None])
                u = f"https://www.boj.or.jp/en/mopo/mpmsche_minu/opinion_{d.year}/opi{d:%y%m%d}.pdf"
                rows.append({"meeting": d, "released": rd[0], "url": u, "rec": self.describe(bank, "opinions", d, u, None), "minutes_released": rd[1]})
        elif bank == "CAD":
            r = self.get("https://www.bankofcanada.ca/content_type/summary-of-deliberations/feed/", ua=None)
            items = _feed_items(r.content) if r is not None else []
            for d in ms:
                it = next((i for i in items if i["pub"] and f"{d:%B}".lower() in i["link"].lower() and f"-{d.day}-{d.year}" in i["link"]), None)
                rows.append({"meeting": d, "released": it["pub"].date() if it else None, "url": it["link"] if it else "", "rec": None,
                             "note": "" if it else "not in feed window"})
        elif bank == "AUD":
            for d in ms:
                u = f"https://www.rba.gov.au/monetary-policy/rba-board-minutes/{d.year}/{d:%Y-%m-%d}.html"
                rec = self.describe(bank, "minutes", d, u, NB_UA, {"sel": ("div", "content", None)})
                r = self.get(u, ua=NB_UA, retries=0)
                m = re.search(r'name="dc.date" content="(\d{4}-\d{2}-\d{2})"', r.text) if r is not None else None
                rows.append({"meeting": d, "released": date.fromisoformat(m.group(1)) if m else None, "url": u, "rec": rec})
        elif bank == "CHF":
            for d in ms:
                rd = d + timedelta(days=28)
                u = f"https://www.snb.ch/en/publications/communication/summaries/zus_{rd:%Y%m%d}"
                rows.append({"meeting": d, "released": rd, "url": u, "rec": self.describe(bank, "summary", d, u, None, {"sel": ("div", None, "a-text")})})
        else:
            self.last_status, self.last_note = "BOT-WALL", "HTTP 403"
            return None
        return CbSeries(bank, self.name, "minutes", [], {"rows": rows}, self.KIND[bank])


def report_b4(probes: list[Probe], today: date) -> None:
    print("\n[B4] MINUTES / ACCOUNTS / SUMMARIES — last 4 meetings (lag = release date − meeting decision day)")
    for p in probes:
        if p.series is None:
            print(f"  {p.bank} FAIL {p.code} {p.note}")
            continue
        print(f"  {p.bank} {p.series.note}")
        for r in p.series.extra["rows"]:
            lag = (r["released"] - r["meeting"]).days if r.get("released") else None
            rec = r.get("rec") or {}
            ex = f" minutes→{r['minutes_released']} (+{(r['minutes_released'] - r['meeting']).days}d)" if r.get("minutes_released") else ""
            print(f"      {r['meeting']}  released={r.get('released') or '—'} lag={lag if lag is not None else '—'}d  "
                  f"http={rec.get('http', '-')} {rec.get('fmt', '') or '-'} {rec.get('code', '') or ''} LM={rec['lm']:%m-%d %H:%MZ}{ex}  "
                  f"{(r.get('url') or '')[-58:]} {r.get('note', '')}" if rec.get("lm") else
                  f"      {r['meeting']}  released={r.get('released') or '—'} lag={lag if lag is not None else '—'}d  "
                  f"http={rec.get('http', '-')} {rec.get('fmt', '') or '-'} {rec.get('code', '') or ''}{ex}  {(r.get('url') or '')[-58:]} {r.get('note', '')}")


# ---------------------------------------------------------------------------
# B5 — speeches / hearings / BIS central bankers' speeches / committee roster
# ---------------------------------------------------------------------------

SPEECH_FEEDS = {
    "USD": [("https://www.federalreserve.gov/feeds/speeches.xml", None, "speeches"),
            ("https://www.federalreserve.gov/feeds/testimony.xml", None, "testimony")],
    "EUR": [("https://www.ecb.europa.eu/rss/press.html", None, "press feed (filter sp/in)")],
    "GBP": [("https://www.bankofengland.co.uk/rss/speeches", None, "speeches")],
    "CAD": [("https://www.bankofcanada.ca/content_type/speeches/feed/", None, "speeches")],
    "AUD": [("https://www.rba.gov.au/rss/rss-cb-speeches.xml", NB_UA, "speeches (1 item retained)")],
    "CHF": [("https://www.snb.ch/public/rss/en/speeches", None, "speeches"),
            ("https://www.snb.ch/public/rss/en/interviews", None, "interviews")],
}
HEARING_RX = re.compile(r"hearing|testimon|committee|parliament|senate|house of|congress|treasury select", re.I)


def _speaker(bank: str, it: dict) -> str:
    t = it["title"]
    if bank == "USD":
        return t.split(",")[0]
    if bank == "EUR":
        return t.split(":")[0] if ":" in t else ""
    if bank == "CHF":
        m = re.match(r"\d{4}-\d{2}-\d{2} - ([^:]+):", t)
        return m.group(1) if m else ""
    return ""


class SpeechSource(TextSource):
    """B5 — official speech / testimony feeds: fields, lag, 30-day volume, full text availability."""
    name = "speech_feeds"
    section = "b5"
    banks = tuple(BANKS)

    def _items(self, bank: str) -> tuple[list[dict], list[str]]:
        notes, out = [], []
        if bank == "JPY":
            r = self.get("https://www.boj.or.jp/en/about/press/koen_2026/index.htm", ua=None)
            if r is not None:
                for m in re.finditer(r'<a[^>]+href="(/en/about/press/koen_\d{4}/ko(\d{6})a\.htm)"[^>]*>(.*?)</a>', _html_text(r), flags=re.S):
                    d = datetime.strptime(m.group(2), "%y%m%d").replace(tzinfo=timezone.utc)
                    out.append({"title": _norm(re.sub(r"<[^>]+>", " ", m.group(3))), "link": "https://www.boj.or.jp" + m.group(1),
                                "raw": d.date().isoformat(), "pub": d, "author": "", "desc": "", "feed": "koen list (HTML)"})
            notes.append("no RSS for speeches only (whatsnew.xml mixes all); HTML list koen_YYYY/index.htm")
            return out, notes
        if bank == "AUD":
            r = self.get("https://www.rba.gov.au/speeches/")
            if r is not None:
                for m in re.finditer(r'<a[^>]+href="(/speeches/\d{4}/([a-z\-]+)-(\d{4}-\d{2}-\d{2})[^"]*\.html)"[^>]*>(.*?)</a>', _html_text(r), flags=re.S):
                    d = datetime.strptime(m.group(3), "%Y-%m-%d").replace(tzinfo=timezone.utc)
                    out.append({"title": _norm(re.sub(r"<[^>]+>", " ", m.group(4))), "link": "https://www.rba.gov.au" + m.group(1),
                                "raw": m.group(3), "pub": d, "author": m.group(2), "desc": "", "feed": "speeches list (HTML)"})
            notes.append("RSS keeps only the latest item; HTML list used for volume")
            return out, notes
        for url, ua, label in SPEECH_FEEDS.get(bank, []):
            r = self.get(url, ua=ua, retries=0)
            if r is None:
                notes.append(f"{label}: {self.last_note}")
                continue
            its = _feed_items(r.content)
            if bank == "EUR":
                its = [i for i in its if re.search(r"ecb\.(sp|in)\d{6}", i["link"])]
            for i in its:
                i["feed"] = label
            out += its
        return out, notes

    def _fetch(self, bank):
        if bank == "NZD":
            self.last_status, self.last_note = "BOT-WALL", "HTTP 403 (rbnz.govt.nz Cloudflare challenge)"
            return None
        items, notes = self._items(bank)
        if not items:
            self.last_note = "; ".join(notes) or "no items"
            return None
        now = datetime.now(timezone.utc)
        last30 = [i for i in items if i["pub"] and now - timedelta(days=30) <= i["pub"] <= now]
        future = [i for i in items if i["pub"] and i["pub"] > now]
        items.sort(key=lambda i: i["pub"] or datetime.min.replace(tzinfo=timezone.utc), reverse=True)
        sample = [i for i in items if i["pub"] and i["pub"] <= now][:3]
        for i in sample:                       # LM of the item page vs feed time (lag evidence), format
            r = self.get(i["link"], ua=NB_UA if bank == "AUD" else None, retries=0) if i["link"] and not i["link"].lower().endswith(".pdf") else None
            i["lm"] = _parse_dt(r.headers.get("Last-Modified", "")) if r is not None else None
            i["fmt"] = "pdf" if i["link"].lower().endswith(".pdf") else "html"
            i["http"] = r.status_code if r is not None else None
            i["speaker"] = _speaker(bank, i)
        pdf_share = sum(1 for i in items if i["link"].lower().endswith(".pdf"))
        hearings = [i["title"][:70] for i in items if HEARING_RX.search(i["title"])][:3]
        return CbSeries(bank, self.name, "speeches", [], {"n_items": len(items), "n30": len(last30), "future": len(future), "sample": sample,
                                                       "pdf": pdf_share, "hearings": hearings, "notes": notes,
                                                       "speaker_in_feed": bool(_speaker(bank, items[0]) if bank in ("USD", "EUR", "CHF") else False)},
                        "; ".join(sorted({i.get("feed", "") for i in items})))


BIS_BANK = {"Federal Reserve": "USD", "European Central Bank": "EUR", "Bank of England": "GBP", "Bank of Japan": "JPY",
            "Bank of Canada": "CAD", "Reserve Bank of Australia": "AUD", "Reserve Bank of New Zealand": "NZD", "Swiss National Bank": "CHF"}


class BisSpeeches(TextSource):
    """B5 — BIS 'central bankers' speeches' feed: coverage per bank, dedup key vs the bank's own feed."""
    name = "bis_cbspeeches"
    section = "b5"
    banks = ("ALL",)

    def _fetch(self, bank):
        r = self.get("https://www.bis.org/doclist/cbspeeches.rss", ua=None)
        if r is None:
            return None
        items = _feed_items(r.content)
        now = datetime.now(timezone.utc)
        per = {b: [] for b in BANKS}
        for i in items:
            for k, cur in BIS_BANK.items():
                if k in i["desc"] and i["pub"] and now - timedelta(days=30) <= i["pub"] <= now:
                    per[cur].append(i)
        # dedup key: speaker surname + title similarity (BIS 'date' is the BIS posting date, not the speech date)
        match = {}
        ss = SpeechSource()
        for cur in ("USD", "EUR", "CHF"):
            own, _ = ss._items(cur)
            own = [o for o in own if o["pub"] and now - timedelta(days=45) <= o["pub"] <= now]
            lags, hit = [], 0
            for o in own:
                last = _speaker(cur, o).split()[-1].lower() if _speaker(cur, o) else ""
                tt = re.sub(r"^.*?[,:]\s*", "", re.sub(r"^\d{4}-\d{2}-\d{2} - ", "", o["title"])).lower()
                for b in per[cur] + [i for i in items if last and last in (i["author"] + i["desc"]).lower()]:
                    if last and last in (b["author"] + b["desc"]).lower() and SequenceMatcher(None, tt, b["title"].lower()).ratio() >= 0.6:
                        hit += 1
                        lags.append((b["pub"].date() - o["pub"].date()).days)
                        break
            match[cur] = (hit, len(own), sorted(lags))
        lm = _parse_dt(r.headers.get("Last-Modified", ""))
        newest = max((i["pub"] for i in items if i["pub"]), default=None)
        return CbSeries("ALL", self.name, "bis", [], {"n": len(items), "per_bank": {b: len(v) for b, v in per.items()},
                                                     "sample": items[:2], "match": match, "lm": lm, "newest": newest,
                                                     "range": (min(i["pub"] for i in items if i["pub"]), newest)},
                        "BIS feed: title, dc:creator, day-only date, description 'Speech by <role>, <bank>'")


ROSTER_URLS = {
    "USD": "https://www.federalreserve.gov/monetarypolicy/fomc.htm",
    "EUR": "https://www.ecb.europa.eu/ecb/orga/decisions/govc/html/index.en.html",
    "GBP": "https://www.bankofengland.co.uk/about/people/monetary-policy-committee",
    "JPY": "https://www.boj.or.jp/en/about/organization/policyboard/index.htm",
    "CAD": "https://www.bankofcanada.ca/about/people/governing-council/",
    "AUD": "https://www.rba.gov.au/about-rba/boards/monetary-policy-board.html",
    "NZD": "https://www.rbnz.govt.nz/about-us/who-we-are/monetary-policy-committee",
    "CHF": "https://www.snb.ch/en/the-snb/organisation/supervisory-management-boards",
}
ROLE_RX = (r"([A-Z][\w'’\-\.]+(?: [A-Z][\w'’\-\.]+){1,3}),? (?:Governor|Deputy Governor|Senior Deputy Governor|Chair(?:man)?|Vice[- ]Chair(?:man)?|"
           r"President|Vice-President|Member|Chief Economist|Executive Director|Director)\b")


class RosterSource(TextSource):
    """B5 — official committee roster: source, format, machine-parseability."""
    name = "committee_roster"
    section = "b5"
    banks = tuple(BANKS)

    def _fetch(self, bank):
        r = self.get(ROSTER_URLS[bank], ua=NB_UA if bank == "AUD" else None, retries=0)
        if r is None:
            return None
        h = _html_text(r)
        t = _norm(_text(h))
        names = list(dict.fromkeys(m.group(1) for m in re.finditer(ROLE_RX, t)))
        extra = {"http": r.status_code, "size": len(r.content), "names": names[:8], "n_names": len(names)}
        if bank == "USD":
            extra["years_present"] = [y for y in ("2026", "2027", "2028") if y in t]
            extra["voting_terms"] = bool(re.search(r"(?i)voting", t))
        return CbSeries(bank, self.name, "roster", [], extra, ROSTER_URLS[bank])


def report_b5(probes: list[Probe], today: date) -> None:
    print("\n[B5] SPEECHES / HEARINGS — official feeds (lag-of-evidence = item-page Last-Modified − feed time)")
    for p in probes:
        if p.series is None:
            print(f"  {p.bank} {p.source} FAIL {p.code} {p.note}")
            continue
        s, x = p.series, p.series.extra
        if s.kind == "speeches":
            print(f"  {p.bank} feed: {s.note} | items={x['n_items']} last30d={x['n30']} future-dated={x['future']} pdf-links={x['pdf']} "
                  f"speaker_in_feed={x['speaker_in_feed']} hearings={x['hearings'][:2]} {' | '.join(x['notes'])}")
            for i in x["sample"]:
                lag = _fmt_t(_mins(i.get("lm"), i["pub"]))
                print(f"      {i['pub']:%Y-%m-%d %H:%M}Z {i['fmt']:4} http={i.get('http')} LM−feed={lag} speaker={i.get('speaker') or '—'!r} {i['title'][:60]!r} {i['link'][-52:]}")
        elif s.kind == "bis":
            print(f"  BIS cbspeeches: {x['n']} items {x['range'][0]:%Y-%m-%d}..{x['range'][1]:%Y-%m-%d}, feed LM {x['lm']:%m-%d %H:%MZ} (newest item {x['newest']:%m-%d} → publication lag ≥ "
                  f"{(x['lm'] - x['newest']).days}d); last-30d per bank {x['per_bank']}; dedup (surname + title≥0.6) matched/own/lag-days: {x['match']}")
            for i in x["sample"]:
                print(f"      {i['pub']:%Y-%m-%d} creator={i['author']!r} {i['desc'][:80]!r} {i['link'][-60:]}")
        elif s.kind == "roster":
            print(f"  {p.bank} roster {x['http']} {x['size']}B names={x['n_names']} {x['names'][:5]} " + (f"years={x.get('years_present')} voting={x.get('voting_terms')}" if p.bank == "USD" else ""))


# ---------------------------------------------------------------------------
# B6 — Forex Factory event names for central-bank events, last 12 weeks (parquet + raw)
# ---------------------------------------------------------------------------

CB_EVENT_RX = re.compile(r"(?i)speaks|testif|press conference|statement|minutes|accounts|summary of|policy|rate\b|meeting|"
                         r"FOMC|MPC|assessment|hearing|report|OCR|cash rate|beige")
NOT_CB = re.compile(r"(?i)unemployment|cpi|ppi|gdp|pmi|retail|trade balance|employment|earnings|sales|production|confidence|"
                    r"claims|housing|building|current account|manufacturing|services|inflation|bond auction|budget|"
                    r"ism|jolts|payroll|survey|index|orders|inventor|import|export|permits|starts|sentiment|tankan|sight deposits|"
                    r"m3|m2|money|credit|lending|reserves|assets|deposits|foreign|net |leading|coincident|cash earnings")


class FfEventNames(TextSource):
    """B6 — distinct central-bank-type event names per currency, last 12 weeks (ff_raw weekly snapshots + range archive)."""
    name = "ff_cb_event_names"
    section = "b6"
    banks = ("ALL",)

    def _fetch(self, bank):
        since = date.today() - timedelta(weeks=12)
        rows = []
        for f in sorted((DATA_DIR / "ff_raw").glob("ff_weekly_*.json")):
            for e in json.loads(f.read_text()):
                d = str(e.get("date", ""))[:10]
                if d and date.fromisoformat(d) >= since and e.get("country") in BANKS:
                    rows.append((e["country"], e["title"], d, "ff_raw"))
        for e in json.loads((DATA_DIR / "archive" / "ff_calendar_range.json").read_text()):
            d = e["Date"][:10].replace(".", "-")
            if e["Currency"] in BANKS and date.fromisoformat(d) >= since:
                rows.append((e["Currency"], e["Name"], d, "range"))
        df = pd.DataFrame(rows, columns=["ccy", "name", "day", "src"]).drop_duplicates(["ccy", "name", "day"])
        df = df[df.name.map(lambda n: bool(CB_EVENT_RX.search(n)) and not NOT_CB.search(n))]
        first_raw = min((f.stem.split("_")[-1] for f in (DATA_DIR / "ff_raw").glob("ff_weekly_*.json")), default="")
        return CbSeries("ALL", self.name, "ff_names", [], {"df": df, "since": since, "ff_raw_first": first_raw,
                                                        "range_last": "2026-07-03"}, "ff_raw + ff_calendar_range.json")


def report_b6(probes: list[Probe], today: date) -> None:
    for p in probes:
        if p.series is None:
            print(f"  B6 FAIL {p.note}")
            continue
        x = p.series.extra
        df = x["df"]
        print(f"\n[B6] FF CENTRAL-BANK EVENT NAMES since {x['since']} (ff_raw snapshots from {x['ff_raw_first']}; range archive to {x['range_last']}; "
              f"gap between them is NOT covered by either)")
        for ccy in BANKS:
            g = df[df.ccy == ccy].groupby("name").agg(n=("day", "count"), last=("day", "max")).sort_values("n", ascending=False)
            print(f"  {ccy}: " + " | ".join(f"{n}×{r.n} (last {r['last'][5:]})" for n, r in g.iterrows())[:900])


# ---------------------------------------------------------------------------
# B7 — ≤15 min trigger: official release times, real GitHub cron delay, external workflow_dispatch
# ---------------------------------------------------------------------------

GH_REPO = "Sstrulea/sentiment-dashboard"
ROOT = Path(__file__).resolve().parents[1]


def _gh_headers() -> dict:
    tok = os.environ.get("GITHUB_TOKEN") or os.environ.get("GH_TOKEN")
    if not tok:
        try:
            tok = subprocess.run(["gh", "auth", "token"], capture_output=True, text=True, timeout=15).stdout.strip()
        except (OSError, subprocess.SubprocessError):
            tok = ""
    return {"Authorization": f"Bearer {tok}", "Accept": "application/vnd.github+json"} if tok else {}


def _gh_json(src: CbSource, path: str):
    r = src.get(f"https://api.github.com/{path}", ua="cb-probe/0.1", retries=0, **_gh_headers())
    return (r.json(), r) if r is not None else (None, None)


class TriggerLatency(TextSource):
    name = "trigger_latency"
    section = "b7"
    banks = ("ALL",)

    def _cron_stats(self) -> dict:
        runs: list[dict] = []
        for pg in (1, 2, 3):
            d, _ = _gh_json(self, f"repos/{GH_REPO}/actions/workflows/econ-refresh.yml/runs?per_page=100&page={pg}&event=schedule")
            got = (d or {}).get("workflow_runs", [])
            runs += got
            if len(got) < 100:
                break
        runs = runs[:200]
        if not runs:
            return {}
        cr = sorted(_parse_dt(r["created_at"]) for r in runs)
        t0, t1 = cr[0], cr[-1]
        ticks, c = [], (t0 - timedelta(hours=4)).replace(minute=5, second=0, microsecond=0)
        while c <= t1:                                    # econ-refresh.yml: '5 * * * 1-5' and '5 */4 * * 0,6' (UTC)
            if c.weekday() < 5 or c.hour % 4 == 0:
                ticks.append(c)
            c += timedelta(hours=1)
        ticks = [t for t in ticks if t >= t0 - timedelta(minutes=1)]
        used, delays, missed = set(), [], 0
        for i, t in enumerate(ticks):
            nxt = ticks[i + 1] if i + 1 < len(ticks) else t1 + timedelta(hours=1)
            k = next((k for k, x in enumerate(cr) if t <= x < nxt and k not in used), None)
            if k is None:
                missed += 1
            else:
                used.add(k)
                delays.append((cr[k] - t).total_seconds() / 60)
        delays.sort()
        gaps = sorted((cr[i + 1] - cr[i]).total_seconds() / 60 for i in range(len(cr) - 1))
        q = lambda a, p: a[min(len(a) - 1, int(round(p * (len(a) - 1))))]
        return {"runs": len(cr), "ticks": len(ticks), "missed": missed, "days": round((t1 - t0).total_seconds() / 86400, 1),
                "delay_p50": q(delays, .5), "delay_p90": q(delays, .9), "delay_max": delays[-1],
                "gap_p50": q(gaps, .5), "gap_p90": q(gaps, .9), "gap_max": gaps[-1],
                "within15": sum(d <= 15 for d in delays)}

    def _fetch(self, bank):
        cron = self._cron_stats()
        jpn = StatementSource().fetch("JPY")
        boj = [(x["meeting"], x["lm"]) for x in (jpn.extra["docs"] if jpn else []) if x.get("lm")]
        wf, resp = _gh_json(self, f"repos/{GH_REPO}/actions/workflows/econ-refresh.yml")
        yml = (ROOT / ".github/workflows/econ-refresh.yml").read_text()
        api = (ROOT / "api/manual-actual.py").read_text()
        return CbSeries("ALL", self.name, "trigger", [], {
            "official": {b: (t.strftime("%H:%M") if t else "variable", tz) for b, (t, tz) in OFFICIAL_TIME.items()},
            "boj_lm": boj, "cron": cron,
            "dispatch": {"workflow_state": (wf or {}).get("state"), "has_workflow_dispatch": "workflow_dispatch" in yml,
                         "concurrency": re.search(r"concurrency:\s*\n\s+group: (\S+)\n\s+cancel-in-progress: (\S+)", yml).groups() if re.search(r"concurrency:", yml) else None,
                         "token_scopes": resp.headers.get("X-OAuth-Scopes") if resp is not None else None,
                         "repo_already_dispatches": "actions/workflows/" in api and "dispatches" in api,
                         "endpoint": f"POST https://api.github.com/repos/{GH_REPO}/actions/workflows/econ-refresh.yml/dispatches  {{\"ref\":\"main\"}} -> 204"}},
                        "GitHub API (read-only)")


def report_b7(probes: list[Probe], today: date) -> None:
    for p in probes:
        if p.series is None:
            print(f"\n[B7] FAIL {p.note}")
            continue
        x = p.series.extra
        print("\n[B7] TRIGGER ≤15 min")
        print("  (a) official release time of the decision (local): " + "; ".join(f"{b} {t} {tz}" for b, (t, tz) in x["official"].items()))
        if x["boj_lm"]:
            print("      BoJ statement PDF Last-Modified (JST), last 4: " + ", ".join(f"{m} {lm.astimezone(ZoneInfo('Asia/Tokyo')):%H:%M}" for m, lm in x["boj_lm"]))
        c = x["cron"]
        if c:
            print(f"  (b) econ-refresh.yml scheduled runs: {c['runs']} runs over {c['days']}d vs {c['ticks']} cron ticks -> {c['missed']} ticks never ran "
                  f"({100 * c['missed'] / c['ticks']:.0f}%); delay of matched runs p50={c['delay_p50']:.0f}m p90={c['delay_p90']:.0f}m max={c['delay_max']:.0f}m; "
                  f"only {c['within15']} runs started ≤15m after their tick; gap between consecutive runs p50={c['gap_p50']:.0f}m p90={c['gap_p90']:.0f}m max={c['gap_max']:.0f}m")
        d = x["dispatch"]
        print(f"  (c) workflow_dispatch: state={d['workflow_state']} declared={d['has_workflow_dispatch']} concurrency={d['concurrency']} "
              f"current token scopes={d['token_scopes']} | repo already dispatches from api/manual-actual.py={d['repo_already_dispatches']}\n      {d['endpoint']}")


# ---------------------------------------------------------------------------
# B8 — PDF-only texts + extractor check
# ---------------------------------------------------------------------------

class PdfCheck(TextSource):
    name = "pdf_check"
    section = "b8"
    banks = ("ALL",)

    def _fetch(self, bank):
        import importlib.util
        req = (ROOT / "requirements.txt").read_text().lower()
        libs = {m: bool(importlib.util.find_spec(m)) for m in ("pypdf", "PyPDF2", "pdfminer", "fitz", "pdfplumber")}
        in_req = [m for m in ("pypdf", "pdfminer", "pymupdf", "pdfplumber") if m in req]
        quality = None
        if libs["pypdf"]:
            boj = StatementSource().fetch("JPY")
            d = boj.extra["docs"][-1] if boj else None
            if d and d.get("chars"):
                quality = {"chars": d["chars"], "has_1_25": "1.25 percent" in d["text"], "has_footnote_date": "September 24, 2026" in d["text"],
                           "sha": d["sha"], "paras": d["paras"], "first": d["first"]}
        ecb_pdf = sum(1 for i in _feed_items(self.get("https://www.ecb.europa.eu/rss/press.html", ua=None).content)
                      if i["link"].endswith(".pdf")) if self.get("https://www.ecb.europa.eu/rss/press.html", ua=None) else None
        return CbSeries("ALL", self.name, "pdf", [], {"importable": libs, "in_requirements": in_req, "quality": quality, "ecb_pdf_items": ecb_pdf}, "requirements/import check")


def report_b8(probes: list[Probe], today: date) -> None:
    for p in probes:
        if p.series is None:
            print(f"\n[B8] FAIL {p.note}")
            continue
        x = p.series.extra
        print(f"\n[B8] PDF — pdf libs in requirements.txt: {x['in_requirements'] or 'NONE'}; importable here: "
              f"{ {k: v for k, v in x['importable'].items() if v} or 'none' }; ECB press-feed items that are PDF-only: {x['ecb_pdf_items']}")
        if x["quality"]:
            print(f"     extraction test (BoJ 2026-09-18 statement via pypdf): {x['quality']}")


# ---------------------------------------------------------------------------
# B9 — private-repo scenario (read-only): visibility, Actions minutes (30d), Vercel gate
# ---------------------------------------------------------------------------

LINUX_RATE_USD_MIN = 0.006          # docs.github.com "Actions runner pricing" (fetched 2026-09-19): Linux 2-core x64
FREE_MINUTES = {"GitHub Free": 2000, "GitHub Pro": 3000, "GitHub Team": 3000}


class PrivateScenario(TextSource):
    name = "private_scenario"
    section = "b9"
    banks = ("ALL",)

    def _fetch(self, bank):
        repo, _ = _gh_json(self, f"repos/{GH_REPO}")
        wfs, _ = _gh_json(self, f"repos/{GH_REPO}/actions/workflows?per_page=50")
        since = (datetime.now(timezone.utc) - timedelta(days=30)).strftime("%Y-%m-%d")
        runs = []
        for w in (wfs or {}).get("workflows", []):
            pg = 1
            while True:
                d, _ = _gh_json(self, f"repos/{GH_REPO}/actions/workflows/{w['id']}/runs?per_page=100&page={pg}&created=%3E%3D{since}")
                rs = (d or {}).get("workflow_runs", [])
                for r in rs:
                    r["_wf"] = w["path"].rsplit("/", 1)[-1]
                runs += rs
                if len(rs) < 100:
                    break
                pg += 1
        from concurrent.futures import ThreadPoolExecutor

        def jobs(r):
            d, _ = _gh_json(self, f"repos/{GH_REPO}/actions/runs/{r['id']}/jobs?per_page=100")
            secs = [(_parse_dt(j["completed_at"]) - _parse_dt(j["started_at"])).total_seconds()
                    for j in (d or {}).get("jobs", []) if j.get("started_at") and j.get("completed_at")]
            return r["_wf"], r["event"], secs
        with ThreadPoolExecutor(8) as ex:
            res = list(ex.map(jobs, runs))
        agg: dict = {}
        for wf, ev, secs in res:
            a = agg.setdefault(wf, {"runs": 0, "jobs": 0, "raw_min": 0.0, "billed_min": 0, "events": {}})
            a["runs"] += 1
            a["events"][ev] = a["events"].get(ev, 0) + 1
            for s_ in secs:
                a["jobs"] += 1
                a["raw_min"] += s_ / 60
                a["billed_min"] += max(1, math.ceil(s_ / 60))
        total = sum(a["billed_min"] for a in agg.values())
        mw = (ROOT / "middleware.js").read_text()
        vj = json.loads((ROOT / "vercel.json").read_text())
        return CbSeries("ALL", self.name, "private", [], {
            "visibility": (repo or {}).get("visibility"), "private": (repo or {}).get("private"), "since": since, "runs": len(runs),
            "agg": agg, "billed_total": total, "raw_total": sum(a["raw_min"] for a in agg.values()),
            "cost_no_free": round(total * LINUX_RATE_USD_MIN, 2),
            "cost_by_plan": {k: round(max(0, total - v) * LINUX_RATE_USD_MIN, 2) for k, v in FREE_MINUTES.items()},
            "middleware": {"exists": True, "auth": "cookie 'dash_auth' + POST /login form (not HTTP Basic)", "matcher": "config.matcher" in mw,
                           "env": "DASH_TOKENS" in mw, "covers_api": "config.matcher" not in mw, "vercel_functions": list((vj.get("functions") or {}).keys())}},
                        "GitHub API + repo files (read-only)")


def report_b9(probes: list[Probe], today: date) -> None:
    for p in probes:
        if p.series is None:
            print(f"\n[B9] FAIL {p.note}")
            continue
        x = p.series.extra
        print(f"\n[B9] PRIVATE SCENARIO — repo visibility={x['visibility']} (private={x['private']})")
        print(f"  GitHub Actions since {x['since']}: {x['runs']} runs, raw {x['raw_total']:.0f} min, billed (ceil per job) {x['billed_total']} min")
        for wf, a in sorted(x["agg"].items(), key=lambda kv: -kv[1]["billed_min"]):
            print(f"      {wf:20} runs={a['runs']:4} jobs={a['jobs']:4} raw={a['raw_min']:6.1f}m billed={a['billed_min']:4}m events={a['events']}")
        print(f"  cost if private @ ${LINUX_RATE_USD_MIN}/min Linux 2-core: no free tier ${x['cost_no_free']}; overage by plan {x['cost_by_plan']} "
              f"(included minutes {FREE_MINUTES})")
        print(f"  Vercel gate already in repo: {x['middleware']}")


# ---------------------------------------------------------------------------
# A3-bis — targeted retry: Eurex JS/XHR, official proxies, participant surveys
# ---------------------------------------------------------------------------

class A3bisProbe(TextSource):
    name = "a3bis"
    section = "a3bis"
    banks = ("ALL",)

    def _eurex(self) -> dict:
        out = {}
        for label, u in (("FESR", "https://www.eurex.com/ex-en/markets/int/mon/euro-short-term-rate-futures/euro-short-term-rate-futures-1404746"),
                         ("FSR3", "https://www.eurex.com/ex-en/markets/int/mon/saron-futures/saron/3M-SARON-Futures-1405958"),
                         ("FEU3", "https://www.eurex.com/ex-en/markets/int/mon/euribor-derivatives/euribor/Three-Month-EURIBOR-Futures-137458")):
            r = self.get(u, ua=None, retries=0)
            if r is None:
                out[label] = {"http": None, "note": self.last_note}
                continue
            h = _html_text(r)
            host = re.search(r'"factsetHostUrl":\s*"([^"]+)"', h)
            out[label] = {"http": r.status_code, "price_tokens": len(re.findall(r"\b9\d\.\d{2,3}\b", _text(h))),
                          "scripts": len(re.findall(r"<script[^>]+src=", h)), "widget_host": host.group(1) if host else None,
                          "token_in_html": bool(re.search(r'"connectionToken"', h)), "tradingview": "s3.tradingview.com/tv.js" in h}
        r = self.get("https://www.eurex.com/ex-en/data/statistics/market-statistics-online/100!onlineStats?productId=1410330&viewType=3&busDate=20260918", ua=None, retries=0)
        out["stats_online"] = {"http": r.status_code if r else None, "has_settlement_table": bool(r and "Settlement" in _text(_html_text(r))),
                               "chars": len(_text(_html_text(r))) if r else 0}
        return out

    def _proxies(self) -> list[dict]:
        rows = []
        for sid, lab in (("TB.CDN.30D.MID", "1M"), ("TB.CDN.90D.MID", "3M"), ("TB.CDN.180D.MID", "6M")):
            r = self.get(f"https://www.bankofcanada.ca/valet/observations/{sid}/json?recent=3", ua=None, retries=0)
            if r is not None:
                ob = max(r.json()["observations"], key=lambda o: o["d"])     # Valet returns newest first with recent=N
                rows.append({"ccy": "CAD", "src": "BoC Valet T-bill " + lab, "asof": ob["d"], "value": float(ob[sid]["v"]), "horizon": lab})
        for cid, lab in (("zirepo", "SNB repo cube (SARON compound 1M/3M/6M, backward-looking)"), ("rendoblid", "SNB Confederation bond yields")):
            r = self.get(f"https://data.snb.ch/api/cube/{cid}/data/csv/en?fromDate=2026-08-01", ua=None, retries=0)
            if r is not None:
                L = [ln for ln in r.text.lstrip("﻿").splitlines() if ln.strip()]
                pub = next((ln for ln in L if "PublishingDate" in ln), "")
                dts = [re.match(r'"(\d{4}-\d{2}-\d{2})', ln) for ln in L[3:]]
                rows.append({"ccy": "CHF", "src": lab, "asof": max((m.group(1) for m in dts if m), default=None), "value": None,
                             "horizon": "≤6M backward" if cid == "zirepo" else "bonds", "pub": pub[-20:-1]})
        r = self.get(RbaPolicy.URL)
        if r is not None:
            rows_ = list(csv.reader(io.StringIO(r.text.lstrip("﻿"))))
            sid_i = next(i for i, x in enumerate(rows_) if x and x[0] == "Series ID")
            col = {s_: j for j, s_ in enumerate(rows_[sid_i]) if s_}
            data = [x for x in rows_[sid_i + 1:] if x and re.match(r"\d{2}-[A-Za-z]{3}-\d{4}$", x[0]) and len(x) > col["FIRMMBAB180D"] and x[col["FIRMMBAB180D"]]]
            last = data[-1]
            for sid, lab in (("FIRMMBAB30D", "1M"), ("FIRMMBAB90D", "3M"), ("FIRMMBAB180D", "6M")):
                rows.append({"ccy": "AUD", "src": "RBA F1 EOD bank bills " + lab, "asof": datetime.strptime(last[0], "%d-%b-%Y").date().isoformat(),
                             "value": float(last[col[sid]]), "horizon": lab})
        return rows

    def _fetch(self, bank):
        return CbSeries("ALL", self.name, "a3bis", [], {"eurex": self._eurex(), "proxies": self._proxies()}, "A3-bis retry")


SURVEYS = [  # (bank, name, cadence, published, horizon, format, terms)
    ("USD", "NY Fed Survey of Market Expectations (SME)", "8/yr, before each FOMC", "results ~3 weeks AFTER the meeting (after the minutes)",
     "expected fed funds at next FOMC meetings + quarter/year ends", "PDF (…/markets/survey/2026/<mon>-2026-sme-results.pdf); page lists 445 PDF links, no XLSX seen",
     "NY Fed Terms of Use — permission sentence not extracted; verify"),
    ("EUR", "ECB Survey of Monetary Analysts (SMA)", "8/yr (every GC meeting)", "Monday of the week after the GC meeting",
     "median DFR path for next meetings and end-year horizons", "PDF aggregate results (ecb.smar<yymmdd>_<month>.en.pdf)",
     "ECB copyright page: reuse with source acknowledgement (page not fully parsed)"),
    ("GBP", "BoE Market Participants Survey", "~8/yr (Feb, Mar, Apr, Jun, Jul, Sep … in 2026)", "18 Sep 2026 for the survey run 2–4 Sep (day AFTER the 17 Sep MPC)",
     "modal Bank Rate after the next MPC meetings + probability distributions", "XLSX + HTML tables (…/survey-results/2026/market-participants-survey-results-<month>-2026.xlsx)",
     "BoE data: OGL v3.0 (bankofengland.co.uk/legal)"),
    ("CAD", "BoC Market Participants Survey", "quarterly", "~2 weeks after the Jan/Apr/Jul/Oct decision (Q2 2026 on 2026-07-27)",
     "policy-rate forecast by horizon (Canada + US)", "HTML tables on the publication page",
     "BoC Terms of Use §1: 'freely use, copy, distribute and transmit' with attribution"),
]


def report_a3bis(probes: list[Probe], today: date) -> None:
    for p in probes:
        if p.series is None:
            print(f"\n[A3-bis] FAIL {p.note}")
            continue
        x = p.series.extra
        print("\n[A3-bis] (a) EUREX quote pages — server-side HTML / XHR behind the widget")
        for k, v in x["eurex"].items():
            print(f"     {k}: {v}")
        print("  (b) OFFICIAL PROXIES (lag vs last business day; not OIS — bills / bank bills / backward compounded SARON)")
        for r in x["proxies"]:
            lag = None
            if r["asof"]:
                d = date.fromisoformat(r["asof"])
                ref = today - timedelta(days=max(0, today.weekday() - 4))
                lag = 0 if d >= ref else int(np.busday_count(d, ref))
            print(f"     {r['ccy']} {r['src']:<62} asof {r['asof']} lag={lag}bd value={r['value']} horizon={r['horizon']} {r.get('pub', '')}")
        print("  (c) PARTICIPANT SURVEYS (public-scenario fallback; publication is AFTER the meeting by design -> criterion 2 fails, so use as cross-check)")
        for b, n, cad, pub, hor, fmt, terms in SURVEYS:
            print(f"     {b} {n}: {cad} | {pub} | horizon: {hor} | {fmt} | terms: {terms}")


# ---------------------------------------------------------------------------
# E — errata for 0A (USD spread, 3M contract windows, one window-selection rule)
# ---------------------------------------------------------------------------

def pick_window(windows: list[dict], eff: date, tol: Optional[int] = None):
    """THE window rule: earliest 3M reference window whose start is on/after (effective date − tol days).
    Forward-only: a window that began more than `tol` days before the meeting is never used, because it
    only partly reflects the decision. Returns (window, gap_days = start − eff) or (None, None)."""
    tol = WINDOW_TOL_DAYS if tol is None else tol
    w = next((r for r in sorted(windows, key=lambda r: r["start"]) if r["start"] >= eff - timedelta(days=tol)), None)
    return (w, (w["start"] - eff).days) if w else (None, None)


class ErrataProbe(TextSource):
    name = "errata_0a"
    section = "e"
    banks = ("ALL",)

    def _fetch(self, bank):
        today = date.today()
        got = {}
        for src_cls, bk in ((AtlantaMpt, "USD"), (JpxTona, "JPY"), (MxCorra, "CAD"), (AsxFutures, "NZD"), (FredTarget, "USD")):
            s = src_cls().fetch(bk)
            if s is not None:
                got[(bk, s.source)] = s
        ctx = _a1_context(["USD", "JPY", "CAD", "NZD"], today)
        # USD: which spread was applied?
        usd = got[("USD", "fred_target")]
        lo = usd.extra["lower"]
        mid = (usd.latest_value + _at(lo, usd.latest_date)) / 2
        sofr = usd.extra["bench"]["SOFR"][-1]
        effr = usd.extra["bench"]["EFFR"][-1]
        mpt = got[("USD", "atlantafed_mpt")]
        w26 = pick_window(a3_windows("atlantafed_mpt", "USD", mpt.extra), max(m.decision for m in official_meetings("USD") if m.decision.year == 2026) + timedelta(days=1))[0]
        usd_err = {"mid": mid, "sofr": sofr, "effr": effr, "sofr_spread_bp": (sofr[1] - mid) * 100, "effr_spread_bp": (effr[1] - mid) * 100,
                   "window": w26, "with_sofr": w26["rate"] - (sofr[1] - mid), "with_effr": w26["rate"] - (effr[1] - mid)}
        # contract listing with reference windows + naming convention
        listing = []
        for r in a3_windows("atlantafed_mpt", "USD", mpt.extra)[:4]:
            listing.append(("USD", "Atlanta MPT", f"ref_start {r['start']}", r["start"], r["end"], "explicit reference_start (CME 3M SOFR options, quarter starts 3rd Wed)", r["rate"]))
        for c in got[("JPY", "jpx_settlement_csv")].extra["contracts"][:4]:
            y, m = int(c["month"][:4]), int(c["month"][4:])
            y2, m2 = _add_months(y, m, 3)
            listing.append(("JPY", "JPX TONA-3M", f"month {c['month']} / last-trade {c['expiry']}", _third_wed(y, m), _third_wed(y2, m2), "START month (spec: 3rd Wed of contract month → Tuesday before 3rd Wed 3 months later)", c["implied"]))
        for c in got[("CAD", "mx_corra_expectations")].extra["contracts"]["CRA"][:4]:
            y, m = c["month"].year, c["month"].month
            y2, m2 = _add_months(y, m, 3)
            listing.append(("CAD", "MX CRA (3M CORRA)", f"CRA{'HMUZ'[[3, 6, 9, 12].index(m)]}{y % 100}", _third_wed(y, m), _third_wed(y2, m2), "START month = 'Contract Reference Month' (spec)", c["implied"]))
        for r in a3_windows("asx_markit_json", "NZD", got[("NZD", "asx_markit_json")].extra)[:4]:
            listing.append(("NZD", "ASX BB (90d bank bill)", f"expiry {r['start']}", r["start"], r["end"], "SETTLEMENT month (expiry = first Wednesday after the 9th; 90d forward from expiry)", r["rate"]))
        # rule applied uniformly
        applied = []
        for bk, srcname in (("USD", "atlantafed_mpt"), ("JPY", "jpx_settlement_csv"), ("CAD", "mx_corra_expectations"), ("NZD", "asx_markit_json")):
            s = got[(bk, srcname)]
            c = ctx[bk]
            sp = (c["spread_bp"] or 0.0) / 100.0
            ms = [m for m in official_meetings(bk) if m.decision > today]
            for lbl, yr in (("last-2026", 2026), ("last-2027", 2027)):
                mm = max((m for m in ms if m.decision.year == yr), key=lambda m: m.decision, default=None)
                if mm is None:
                    continue
                eff = mm.decision + timedelta(days=EFF_LAG_DAYS[bk])
                w, gap = pick_window(a3_windows(srcname, bk, s.extra), eff)
                inside = sum(1 for m in official_meetings(bk) if w and w["start"] <= m.decision < w["end"] and m is not mm) if w else None
                applied.append({"bank": bk, "lbl": lbl, "meeting": mm.decision, "eff": eff, "window": w, "gap": gap, "inside": inside,
                                "implied": (w["rate"] - sp) if w else None, "policy": c["policy"], "spread_bp": c["spread_bp"]})
        return CbSeries("ALL", self.name, "errata", [], {"usd": usd_err, "listing": listing, "applied": applied}, "0A errata")


def report_e(probes: list[Probe], today: date) -> None:
    for p in probes:
        if p.series is None:
            print(f"\n[E] FAIL {p.note}")
            continue
        x = p.series.extra
        u = x["usd"]
        print("\n[E1] USD — which overnight–policy spread was subtracted from the MPT window?")
        print(f"   midpoint {u['mid']:.3f}; SOFR {u['sofr'][0]}={u['sofr'][1]:.2f} ({u['sofr_spread_bp']:+.1f}bp vs mid); EFFR {u['effr'][0]}={u['effr'][1]:.2f} ({u['effr_spread_bp']:+.1f}bp)")
        print(f"   MPT window {u['window']['start']}: mean {u['window']['rate']:.3f}  -> minus SOFR spread = {u['with_sofr']:.3f} ({(u['with_sofr'] - u['mid']) * 100:+.0f}bp)   "
              f"| minus EFFR spread = {u['with_effr']:.3f} ({(u['with_effr'] - u['mid']) * 100:+.0f}bp)")
        print("\n[E2] 3M CONTRACTS — reference window [start, end) and exchange naming convention")
        for r in x["listing"]:
            print(f"   {r[0]} {r[1]:<24} {r[2]:<38} window {r[3]}..{r[4]}  implied {r[6]:.3f}  | {r[5]}")
        print(f"\n[E3] WINDOW RULE: earliest window with start ≥ (effective date − {WINDOW_TOL_DAYS}d); forward-only; effective = decision + "
              f"{EFF_LAG_DAYS}; result = window mean − current overnight–policy spread; 'inside' = other meetings inside the window (upper bound).")
        for a in x["applied"]:
            w = a["window"]
            print(f"   {a['bank']} {a['lbl']} {a['meeting']} (eff {a['eff']}): window {w['start']}..{w['end']} gap={a['gap']:+d}d inside={a['inside']} "
                  f"implied {a['implied']:.3f} ({(a['implied'] - a['policy']) * 100:+.0f}bp vs {a['policy']:.3f}; spread {'n/a' if a['spread_bp'] is None else format(a['spread_bp'], '+.1f') + 'bp'})"
                  if w else f"   {a['bank']} {a['lbl']} {a['meeting']}: no window")


# ---------------------------------------------------------------------------
# Probe runners + verdicts
# ---------------------------------------------------------------------------

def run(src: CbSource, bank: str, today: date) -> Probe:
    log.info("probing %s / %s ...", bank, src.name)
    s = src.fetch(bank)
    if s is None:
        if src.last_status == "UNVERIFIED":        # nothing failed — we just could not confirm it programmatically
            return Probe(bank, src.section, src.name, "PARTIAL", "NONE", None, "UNVERIFIED: " + src.last_note, cdn=src.cdn)
        return Probe(bank, src.section, src.name, "FAIL", src.failure_code(), None, src.last_note, cdn=src.cdn)
    return Probe(bank, src.section, src.name, "OK", "", s, s.note, [], src.cdn)


A1_SOURCES = [FredTarget, EcbPolicy, BoePolicy, BojPolicy, BocPolicy, RbaPolicy, RbnzPolicy, SnbPolicy, BisPolicy]
A2_SOURCES = [FedSep, RbnzOcrTrack, PathAbsence]
A3_SOURCES = [BoeOis, AsxFutures, AsxHistoryMirror, MxCorra, JpxTona, AtlantaMpt, TreasuryBills, EcbYcForward]
A4_SOURCES = [FfConsensus]
A5_SOURCES = CALENDARS + [QuietPeriod]
SECTION_SOURCES = {"a1": A1_SOURCES, "a2": A2_SOURCES, "a3": A3_SOURCES, "a4": A4_SOURCES, "a5": A5_SOURCES,
                   "b1": [StatementSource], "b2": [VoteSource], "b3": [PressConfSource], "b4": [MinutesSource],
                   "b5": [SpeechSource, BisSpeeches, RosterSource], "b6": [FfEventNames],
                   "b7": [TriggerLatency], "b8": [PdfCheck], "b9": [PrivateScenario], "a3bis": [A3bisProbe], "e": [ErrataProbe]}


def probe_section(section: str, banks: list[str], today: date) -> list[Probe]:
    out: list[Probe] = []
    insts = [cls() for cls in SECTION_SOURCES[section]]
    for src in insts:
        if src.banks == ("ALL",):                    # not bank-specific: run once
            out.append(run(src, "ALL", today))
    for bank in banks:
        for src in insts:
            if src.banks != ("ALL",) and src.supports(bank):
                out.append(run(src, bank, today))
        if section == "a3":
            b = BlockedProbe()
            for bk, name, url, kind in BLOCKED_TARGETS:
                if bk == bank:
                    out.append(b.probe_target(bk, name, url, kind))
    return out


# ---- meetings (official calendar) helper, reused by A1 (last-4 decisions) and A3 (horizon) ----

_MEET_CACHE: dict = {}


def official_meetings(bank: str) -> list[Meeting]:
    """Official calendar; if the page lists only upcoming meetings (ECB), past 2026 decision days
    are taken from the local FF rows and tagged basis='ff' (NOT official)."""
    if bank not in _MEET_CACHE:
        src = next(c for c in CALENDARS if bank in c.banks)()
        s = src.fetch(bank)
        ms = list(s.extra["meetings"]) if s else []
        today = date.today()
        if not any(m.decision <= today for m in ms):
            ff = FfConsensus()
            pq, _ = ff._load()
            rows = pq[(pq.canonical_id == f"{FF_ID[bank]}_interest_rate_decision") & pq.actual.notna()
                      & (pq.datetime_utc.dt.year == 2026)]
            ms += [Meeting(d, None, None, None, "ff (not official)") for d in sorted(set(rows.datetime_utc.dt.date))]
        _MEET_CACHE[bank] = sorted(ms, key=lambda m: m.decision)
    return _MEET_CACHE[bank]


def last_decisions(points: list, meetings: list[Meeting], today: date, n=4) -> list[dict]:
    """Level after / change at the last n decision dates <= today, from an effective-dated policy series.
    Level = value just before the next meeting (captures ECB's +6d implementation lag)."""
    past = sorted(m.decision for m in meetings if m.decision <= today)
    nxt = sorted(m.decision for m in meetings if m.decision > today)
    out = []
    latest = points[-1][0] if points else None
    for d in past[-n:]:
        idx = past.index(d)
        upper = (past[idx + 1] if idx + 1 < len(past) else (nxt[0] if nxt else today + timedelta(days=30))) - timedelta(days=1)
        if latest is None or latest < d:
            out.append({"date": d, "level": None, "delta_bp": None, "pending": True})   # series lags the decision
            continue
        after = _at(points, min(upper, today))
        before = _at(points, d - timedelta(days=1))
        out.append({"date": d, "level": after, "pending": False,
                    "delta_bp": None if after is None or before is None else round((after - before) * 100, 1)})
    return out


# ---------------------------------------------------------------------------
# Reporting
# ---------------------------------------------------------------------------

def _fmt_lag(s: CbSeries, today: date) -> str:
    lag = s.lag_bd(today)
    return "—" if lag is None else f"{lag}bd"


def report_a1(probes: list[Probe], today: date) -> None:
    print("\n[A1] POLICY RATE — per-source log")
    by_bank: dict[str, dict[str, Probe]] = {}
    for p in probes:
        by_bank.setdefault(p.bank, {})[p.source] = p
        s = p.series
        if s is None:
            print(f"  {p.bank} {p.source:22} FAIL {p.code:8} {p.note}   cdn={p.cdn or '-'}")
        else:
            print(f"  {p.bank} {p.source:22} OK   latest {s.latest_date} = {s.latest_value:g}  lag={_fmt_lag(s, today)} "
                  f"n={s.n_points} since {s.history_start}  cdn={p.cdn or '-'}")
    print("\n[A1] LAST 4 DECISIONS (primary source x official calendar) and overnight-benchmark spread")
    yaml_rates = _load_yaml_rates()
    ff_src = FfConsensus()
    for bank, d in by_bank.items():
        prim = d.get(A1_PRIMARY[bank])
        if prim is None or prim.series is None:
            print(f"  {bank}: primary {A1_PRIMARY[bank]} unavailable")
            continue
        s = prim.series
        ms = official_meetings(bank)
        dec = last_decisions(s.points, ms, today)
        lo = s.extra.get("lower")
        rng = ""
        if lo:
            rng = f" range {_at(lo, s.latest_date):g}–{s.latest_value:g}"
        print(f"  {bank} [{prim.source}] level={s.latest_value:g}{rng}")
        for x in dec:
            print(f"      {x['date']}  " + ("PENDING — series lags this decision (latest obs "
                  f"{s.latest_date})" if x["pending"] else f"level {x['level']}  d {x['delta_bp']}bp"))
        for bname, bp in (s.extra.get("bench") or {}).items():
            if bp:
                pol = _at(s.points, bp[-1][0])
                if lo:
                    pol_mid = (pol + _at(lo, bp[-1][0])) / 2
                    print(f"      bench {bname} {bp[-1][0]}={bp[-1][1]:g}  vs upper {(bp[-1][1] - pol) * 100:+.0f}bp / midpoint {(bp[-1][1] - pol_mid) * 100:+.1f}bp")
                elif pol is not None:
                    print(f"      bench {bname} {bp[-1][0]}={bp[-1][1]:g}  vs policy {pol:g}: {(bp[-1][1] - pol) * 100:+.1f}bp")
        if bank == "EUR" and s.extra.get("mro"):
            print(f"      MRO {s.extra['mro'][-1][1]:g} -> DFR {s.latest_value:g}: MRO-DFR = {(s.extra['mro'][-1][1] - s.latest_value) * 100:.0f}bp")
        if bank == "JPY":
            blr = by_bank["JPY"]["boj_timeseries_api"].series if "boj_timeseries_api" in by_bank["JPY"] and by_bank["JPY"]["boj_timeseries_api"].series else None
            if blr and blr.extra.get("blr"):
                print(f"      BoJ API: call rate avg {blr.latest_date}={blr.latest_value:g}; Basic Loan Rate {blr.extra['blr'][-1][0]}={blr.extra['blr'][-1][1]:g}")
        if bank == "AUD" and s.extra.get("ois_short"):
            print("      OIS(F1): " + ", ".join(f"{k}={v[0][1]:g}@{v[0][0]}" for k, v in s.extra["ois_short"].items() if v))
        ffs = ff_src.fetch(bank)
        if ffs is not None:
            have = [x for x in ffs.extra["decisions"] if x["actual"] is not None and pd.notna(x["actual"])]
            if have:
                d = have[-1]
                print(f"      FF last decision row: {d['day']} actual={d['actual']} forecast={d['forecast']} previous={d['previous']}")
            else:
                print("      FF last decision row: none with a numeric actual in the last 4 meetings")
        y = yaml_rates.get(bank)
        if y:
            off = s.latest_value if not lo else (s.latest_value + _at(lo, s.latest_date)) / 2
            print(f"      yaml policy_rates.yaml: {y['rate_pct']} (eff {y['effective']}, verified {y['verified']}) "
                  f"vs official {off:g} -> {'MATCH' if abs(y['rate_pct'] - off) < 1e-9 else 'DIFF ' + format((y['rate_pct'] - off) * 100, '+.1f') + 'bp'}")
    bis = {p.bank: p for p in probes if p.source == "bis_cbpol" and p.series}
    if bis:
        print("\n[A1] BIS WS_CBPOL fallback: latest obs / lag")
        for b, p in bis.items():
            print(f"  {b} {p.series.latest_date} = {p.series.latest_value:g}  lag={_fmt_lag(p.series, today)}")


def _load_yaml_rates() -> dict:
    try:
        import yaml
        return yaml.safe_load((DATA_DIR / "policy_rates.yaml").read_text())["rates"]
    except Exception:
        return {}


def report_a2(probes: list[Probe], today: date) -> None:
    print("\n[A2] OWN TRAJECTORY")
    for p in probes:
        s = p.series
        if s is None:
            print(f"  {p.bank} {p.source:22} FAIL {p.code:8} {p.note}   cdn={p.cdn or '-'}")
        elif s.kind == "sep":
            print(f"  {p.bank} {p.source:22} OK   SEP {s.extra['asof']}  cdn={p.cdn or '-'}")
            print(f"      median (published): {s.extra['median_published']}")
            print(f"      median (from dots) : { {y: v for y, v in s.extra['median_from_dots'].items()} }")
            for y, cs in s.extra["dots"].items():
                print(f"      dots {y}: " + ", ".join(f"{k:g}x{n}" for k, n in sorted(cs.items(), reverse=True)))
        else:
            print(f"  {p.bank} {p.source:22} OK   {s.note} | evidence: {s.extra.get('evidence', '')[:120]}")


def report_a3(probes: list[Probe], today: date) -> None:
    print("\n[A3] MARKET-IMPLIED PATH — per-source log")
    got: dict[tuple, CbSeries] = {}
    for p in probes:
        s = p.series
        if s is None:
            print(f"  {p.bank} {p.source:24} FAIL {p.code:8} {p.note}   cdn={p.cdn or '-'}")
            continue
        got[(p.bank, p.source)] = s
        hist = f"n_dates={s.extra.get('n_dates')}" if s.extra.get("n_dates") else (
            f"n_scrapes={s.extra.get('n_scrapes')}" if s.extra.get("n_scrapes") else f"n={s.n_points}")
        print(f"  {p.bank} {p.source:24} OK   asof={s.extra.get('asof') or s.latest_date}  lag={_fmt_lag(s, today)}  {hist}  cdn={p.cdn or '-'}")
    print("\n[A3] DETAIL")
    for (bank, src), s in got.items():
        x = s.extra
        if src == "boe_ois_curve":
            print(f"  GBP BoE OIS fwd (asof {x['asof']}): months {[round(m) for m in x['months'][:12]]}...{round(x['months'][-1])}")
            print("      " + ", ".join(f"{round(m)}m={v:.3f}" for m, v in list(zip(x['months'], x['curve']))[:24]))
            print(f"      history {s.history_start}..{s.latest_date} ({x['n_dates']} dates), horizon {x['horizon_months']:.0f}m")
        elif src == "asx_markit_json":
            print(f"  {bank} ASX {x['instrument']} asof {x['asof']} ({len(x['contracts'])} contracts, expiries {x['contracts'][0]['expiry']}..{x['contracts'][-1]['expiry']})")
            print("      " + ", ".join(f"{c['symbol']}={c['implied']:.3f}" for c in x["contracts"]))
        elif src == "asx_history_mirror":
            print(f"  AUD ASX mirror: {x['n_scrapes']} daily scrapes {s.history_start}..{s.latest_date}, {x['n_expiries']} expiry columns")
        elif src == "mx_corra_expectations":
            print(f"  CAD MX (no as-of stamp) CORRA spot {x['corra_spot']}; BoC meetings on page: {[str(d) for d in x['boc_meetings']]}")
            for fam in ("COA", "CRA"):
                print(f"      {fam}: " + ", ".join(f"{c['month']:%b%y}={c['implied']:.3f}" for c in x["contracts"][fam]))
        elif src == "jpx_settlement_csv":
            print(f"  JPY JPX 3M TONA futures asof {x['asof']} ({len(x['contracts'])} contracts to {x['contracts'][-1]['expiry']})")
            print("      " + ", ".join(f"{c['expiry']:%y%m%d}={c['implied']:.3f}" for c in x["contracts"][:12]))
        elif src == "atlantafed_mpt":
            print(f"  USD Atlanta Fed MPT asof {x['asof']} target range {x['target_range']}: history {x['first_date']}..{x['asof']} "
                  f"({x['n_dates']} dates), farthest window {x['horizon']}")
            print("      window mean (bp): " + ", ".join(f"{a:%y-%m-%d}={b:.0f}" for a, b in x["windows"][:14]))
        elif src == "treasury_par_curve":
            print(f"  USD Treasury par curve asof {x['asof']}: " + ", ".join(f"{k}={v:.2f}" for k, v in x["curve"].items()))
        elif src == "ecb_yc_forward":
            c = x["curve"]
            print(f"  EUR ECB YC AAA-govt inst. forward asof {x['asof']}: " + ", ".join(f"{m}m={c[m]:.3f}" for m in (3, 6, 9, 12, 18, 24, 36) if m in c))


def _print_negative(probes: list[Probe]) -> None:
    print("\n[A3] NEGATIVE EVIDENCE (blocked / JS-only / paid candidates probed from this IP)")
    for p in probes:
        if p.source in {t[1] for t in BLOCKED_TARGETS}:
            print(f"  {p.bank} {p.source:20} {p.status} {p.code:8} {p.note}  cdn={p.cdn or '-'}")


def report_a4(probes: list[Probe], today: date) -> None:
    print("\n[A4] FOREX FACTORY CONSENSUS — last 4 decisions per bank, joined to the official calendar (local parquet)")
    for p in probes:
        s = p.series
        if s is None:
            print(f"  {p.bank} FAIL {p.note}")
            continue
        x = s.extra
        print(f"  {p.bank} {s.note} rows={x['n_rows']} last_row={x['last_row_day']} can_be_zero={x['can_be_zero']} "
              f"raw={x['raw_names']} canon={x['canonical_names']}")
        for d in x["decisions"]:
            if d["n_rows"] == 0:
                mk = x["markers"].get(str(d["meeting"]))
                print(f"      {d['meeting']}  NO NUMERIC ROW in FF" + (f"  (marker event present: {mk})" if mk else "") + f"  [{d['basis']}]")
                continue
            cov = ("A" if pd.notna(d["actual"]) else "-") + ("F" if pd.notna(d["forecast"]) else "-") + ("P" if pd.notna(d["previous"]) else "-")
            zf = f"  ZERO-ROWS {d['zero_rows']}" if d["zero_rows"] else ""
            print(f"      {d['meeting']}  actual={d['actual']}  forecast={d['forecast']}  previous={d['previous']}  "
                  f"rows={d['n_rows']} cov={cov}{zf}")
        print(f"      all rows with actual==0.0: {x['zero_rows_all'][-6:]}  | archive rows actual 0.0 & previous>0: "
              f"{x['range_zero_prev_pos']}/{x['range_json_rows']} (archive last {x['range_json_last']})")


def _m(m: Meeting) -> str:
    return f"{m.decision:%m-%d}" + ("P" if m.projections else "") + ("C" if m.presser else "")


def report_a5(probes: list[Probe], today: date) -> None:
    print("\n[A5] OFFICIAL CALENDAR 2026-2027  (P=projections that day, C=press conference; basis in brackets)")
    for p in probes:
        s = p.series
        if s is None:
            print(f"  {p.bank} {p.source:20} FAIL {p.code:8} {p.note}")
            continue
        if s.kind == "quiet":
            print(f"  {p.bank} quiet-period [{s.note}]: {s.extra['rule']}\n        {s.extra['url']}")
            continue
        ms = s.extra["meetings"]
        for yr in (2026, 2027):
            row = [m for m in ms if m.decision.year == yr]
            print(f"  {p.bank} {p.source:20} {yr} ({len(row)}): " + " ".join(_m(m) for m in row) + f"   [{row[0].basis if row else ''}]")
    for p in probes:
        if p.series is None and p.source == "quiet_period_rules":
            print(f"  {p.bank} quiet-period {p.note}")


BENCH_PREF = {"USD": "SOFR"}
WINDOW_TOL_DAYS = 3     # an IMM window starting <=3d before the meeting still captures ~all of its effect
COARSE_PROXIES = {"treasury_par_curve", "ecb_yc_forward"}   # bills / AAA-govt curve: no per-meeting resolution, not OIS
EFF_LAG_DAYS = {"USD": 1, "EUR": 6, "GBP": 0, "JPY": 0, "CAD": 0, "AUD": 1, "NZD": 0, "CHF": 1}

LICENSE = {   # source -> (verdict, evidence). 'ok' | 'flag' (unverified/restricted) | 'blocked'
    "boe_ois_curve": ("flag", "BoE data reusable under OGL v3.0, but third-party (Bloomberg-sourced) series are excluded; OIS pages cite "
                              "'Bloomberg Finance L.P. and Bank calculations' -> get written confirmation before publishing"),
    "atlantafed_mpt": ("blocked", "workbook LICENSE box: 'Use of this data is permitted for personal and educational purposes only'; "
                                  "CME data used under CME permission; site terms bar publishing without authorization"),
    "treasury_par_curve": ("ok", "US Treasury data — US-government work, public domain"),
    "ecb_yc_forward": ("ok", "ECB statistics reusable with source acknowledgement (not re-fetched this session)"),
    "asx_markit_json": ("flag", "undocumented site API behind asx.com.au; ASX terms page not retrievable (404 on guessed URLs) -> unverified; "
                                "exchange data is normally licensed, redistribution needs an ASX licence"),
    "asx_history_mirror": ("flag", "third-party GitHub repo with no licence file; data is derived from ASX"),
    "mx_corra_expectations": ("flag", "Bourse de Montréal copyright; 'for general information purposes only'; TMX Datalinx sells MX data "
                                      "'for a single end user, no redistribution rights'"),
    "jpx_settlement_csv": ("flag", "JPX terms page not retrievable (403/404); JPX sells history via J-Quants DataCube -> unverified"),
}


def _third_wed(y: int, m: int) -> date:
    d = date(y, m, 1)
    return d + timedelta(days=(2 - d.weekday()) % 7 + 14)


def _add_months(y: int, m: int, k: int) -> tuple[int, int]:
    yy, mm = divmod(m - 1 + k, 12)
    return y + yy, mm + 1


def a3_windows(src: str, bank: str, x: dict) -> list[dict]:
    """3M reference windows [start, end) with the implied average overnight rate in percent."""
    w = []
    if src == "atlantafed_mpt":
        cur = x["windows"]
        for i, (a, v) in enumerate(cur):
            w.append({"start": a, "end": cur[i + 1][0] if i + 1 < len(cur) else a + timedelta(days=91), "rate": v / 100.0})
    elif src == "mx_corra_expectations":
        for c in x["contracts"]["CRA"]:
            y, m = c["month"].year, c["month"].month
            y2, m2 = _add_months(y, m, 3)
            w.append({"start": _third_wed(y, m), "end": _third_wed(y2, m2), "rate": c["implied"]})
    elif src == "jpx_settlement_csv":
        for c in x["contracts"]:
            y, m = int(c["month"][:4]), int(c["month"][4:])
            y2, m2 = _add_months(y, m, 3)
            w.append({"start": _third_wed(y, m), "end": _third_wed(y2, m2), "rate": c["implied"]})
    elif src == "asx_markit_json" and bank == "NZD":
        for c in x["contracts"]:      # ASX factsheet: expiry = first Wednesday after the 9th of the settlement month
            d0 = date(c["expiry"].year, c["expiry"].month, 10)
            start = d0 + timedelta(days=(2 - d0.weekday()) % 7)
            w.append({"start": start, "end": start + timedelta(days=91), "rate": c["implied"], "label": c["symbol"]})
    return sorted(w, key=lambda r: r["start"])


def a3_horizon_end(src: str, bank: str, x: dict, series: CbSeries) -> Optional[date]:
    if src == "atlantafed_mpt":
        return x.get("horizon")
    if src == "boe_ois_curve":
        return x["asof"] + timedelta(days=int(x["horizon_months"] * 30.44))
    if src == "asx_markit_json":
        return x["contracts"][-1]["expiry"]
    if src == "jpx_settlement_csv":
        return x["contracts"][-1]["expiry"]
    if src == "mx_corra_expectations":
        return max(_third_wed(c["month"].year, c["month"].month) for c in x["contracts"]["CRA"])
    if src == "treasury_par_curve":
        return x["asof"] + timedelta(days=730)
    if src == "ecb_yc_forward":
        return x["asof"] + timedelta(days=int(max(x["curve"]) * 30.44))
    if src == "asx_history_mirror":
        return x.get("horizon")
    return None


def assess_a3(probes: list[Probe], today: date) -> None:
    """Pre-registered criteria (1)-(6); fills Probe.status/code/flags in place."""
    for p in probes:
        s = p.series
        if s is None or p.source in {t[1] for t in BLOCKED_TARGETS}:
            continue
        x = s.extra
        ms = official_meetings(p.bank)
        need = max((m.decision for m in ms if m.decision.year == 2027), default=None)
        hz = a3_horizon_end(p.source, p.bank, x, s)
        n_hist = x.get("n_dates") or x.get("n_scrapes") or _hist_bd(s.points)
        lag = s.lag_bd(today) if s.latest_date else None
        lic = LICENSE.get(p.source, ("ok", ""))
        crit = {
            "1_keyless": True,
            "2_eod_lag<=1": lag is not None and lag <= MAX_LAG_BD,
            "3_horizon>=last2027": bool(hz and need and hz >= need),
            "4_hist>=30bd": n_hist >= MIN_HIST_BD,
            "5_deterministic": True,
            "6_license": lic[0] == "ok",
        }
        codes = []
        if not crit["2_eod_lag<=1"]:
            codes.append("STALE")
        if not crit["3_horizon>=last2027"]:
            codes.append("SHORT_HORIZON")
        if not crit["4_hist>=30bd"]:
            codes.append("NO_HISTORY")
        if not crit["6_license"]:
            codes.append("LICENSE")
        p.status = "OK" if not codes else "PARTIAL"
        p.code = codes[0] if codes else ""
        if p.source in COARSE_PROXIES:
            p.status, p.code = "PARTIAL", p.code or "NONE"
            codes = codes + ["PROXY"]
        p.flags = codes + ([f"hist={n_hist}" + ("" if n_hist >= 250 else "<250")] if "NO_HISTORY" not in codes else [])
        p.series.extra["criteria"] = crit
        p.series.extra["license"] = lic
        p.series.extra["horizon_end"] = hz
        p.series.extra["need"] = need


def _a1_context(banks: list[str], today: date) -> dict:
    """policy level now + overnight-vs-policy spread (bp) per bank, from the A1 sources."""
    out = {}
    got = {(p.bank, p.source): p.series for p in probe_section("a1", banks, today) if p.series}
    for b in banks:
        prim = got.get((b, A1_PRIMARY[b]))
        if prim is None:
            continue
        lo = prim.extra.get("lower")
        pol = prim.latest_value if not lo else (prim.latest_value + _at(lo, prim.latest_date)) / 2
        bm = prim.extra.get("bench") or {}
        bench = bm.get(BENCH_PREF.get(b, ""), None) or next((v for v in bm.values() if v), None)
        spread = None
        if bench:
            ref = _at(prim.points, bench[-1][0])
            ref = ref if not lo else (ref + _at(lo, bench[-1][0])) / 2
            spread = (bench[-1][1] - ref) * 100
        out[b] = {"policy": pol, "spread_bp": spread}
    if "JPY" in banks:      # BIS guideline (lag) vs BoJ call rate on the same date; true policy after 18-Sep hike = 1.25 (FF/BoJ)
        call = got.get(("JPY", "boj_timeseries_api"))
        bis = got.get(("JPY", "bis_cbpol"))
        if call and bis:
            d = min(call.latest_date, bis.latest_date)
            out["JPY"] = {"policy": 1.25, "spread_bp": (_at(call.points, d) - _at(bis.points, d)) * 100,
                          "note": "policy 1.25 = FF/BoJ 18-Sep decision (BIS/BoJ-API series still show 1.00)"}
    if "NZD" in banks and "NZD" not in out:
        bis = got.get(("NZD", "bis_cbpol"))
        if bis:
            out["NZD"] = {"policy": bis.latest_value, "spread_bp": None, "note": "BKBM/OCR spread unavailable (RBNZ B2 blocked)"}
    return out


def report_a3_implied(probes: list[Probe], today: date) -> None:
    banks = sorted({p.bank for p in probes if p.series is not None})
    ctx = _a1_context(banks, today)
    got = {(p.bank, p.source): p.series for p in probes if p.series}
    print("\n[A3] IMPLIED POLICY-EQUIVALENT AFTER THE LAST 2026 / LAST 2027 MEETING (arithmetic only; assumptions inline)")
    print("     policy-equivalent = market-implied overnight rate - (current overnight-policy spread); delta vs policy now")
    for bank in banks:
        c = ctx.get(bank)
        ms = [m for m in official_meetings(bank) if m.decision > today]
        if not c or not ms:
            continue
        m26 = max((m for m in ms if m.decision.year == 2026), key=lambda m: m.decision, default=None)
        m27 = max((m for m in ms if m.decision.year == 2027), key=lambda m: m.decision, default=None)
        sp = (c["spread_bp"] or 0.0) / 100.0
        print(f"  {bank}: policy now {c['policy']:.3f}  overnight-policy spread "
              f"{'n/a' if c['spread_bp'] is None else format(c['spread_bp'], '+.1f') + 'bp'} {c.get('note', '')}")
        for src in ("atlantafed_mpt", "mx_corra_expectations", "jpx_settlement_csv", "asx_markit_json"):
            s = got.get((bank, src))
            if s is None:
                continue
            x = s.extra
            if src == "asx_markit_json" and bank == "AUD":
                mon = {(cn["expiry"].year, cn["expiry"].month): cn["implied"] for cn in x["contracts"]}
                eff = [m.decision + timedelta(days=EFF_LAG_DAYS[bank]) for m in ms]
                ch = implied_1m_chain(mon, eff, c["policy"] + sp)
                res = {e: v for e, v, _ in ch if v is not None}
                for lbl, mm in (("last-2026", m26), ("last-2027", m27)):
                    e = mm.decision + timedelta(days=EFF_LAG_DAYS[bank])
                    if e in res:
                        print(f"      [{src} IB 1M chain] {lbl} {mm.decision}: {res[e] - sp:.3f}  ({(res[e] - sp - c['policy']) * 100:+.0f}bp)  "
                              "[assumes one meeting/month, constant spread; late-month meetings noisy]")
                continue
            if src == "mx_corra_expectations":
                coa = {(cn["month"].year, cn["month"].month): cn["implied"] for cn in x["contracts"]["COA"]}
                eff = [m.decision + timedelta(days=EFF_LAG_DAYS[bank]) for m in ms]
                ch = implied_1m_chain(coa, eff, c["policy"] + sp)
                for e, v, why in ch:
                    print(f"      [{src} COA 1M chain] meeting eff {e}: " + (f"{v - sp:.3f} ({(v - sp - c['policy']) * 100:+.0f}bp)" if v is not None else f"n/a {why}"))
                print(f"      [{src} COA] horizon ends {max(coa)} -> SHORT_HORIZON for later meetings; CRA below")
            w = a3_windows(src, bank, x)
            for lbl, mm in (("last-2026", m26), ("last-2027", m27)):
                if mm is None:
                    continue
                eff = mm.decision + timedelta(days=EFF_LAG_DAYS[bank])
                win, gap = pick_window(w, eff)
                if win is None:
                    print(f"      [{src}] {lbl} {mm.decision}: no window starts after the meeting (SHORT_HORIZON)")
                    continue
                nin = sum(1 for m in official_meetings(bank) if win["start"] <= m.decision < win["end"] and m is not mm)
                v = win["rate"] - sp
                print(f"      [{src} window {win['start']}..{win['end']}] {lbl} {mm.decision}: {v:.3f}  ({(v - c['policy']) * 100:+.0f}bp)"
                      f"  [{nin} further meeting(s) inside the window -> upper bound of the cumulative move]"
                      + ("  [BKBM-OCR spread unknown: level is BKBM, not OCR]" if bank == "NZD" else ""))
        s = got.get((bank, "boe_ois_curve"))
        if s is not None:
            x = s.extra
            for lbl, mm in (("last-2026", m26), ("last-2027", m27)):
                t = (mm.decision - x["asof"]).days / 30.44
                v = fwd_at(x["months"], x["curve"], t) - sp
                print(f"      [boe_ois_curve fwd @{t:.1f}m] {lbl} {mm.decision}: {v:.3f}  ({(v - c['policy']) * 100:+.0f}bp)  [fitted spline: smooths steps]")
        s = got.get((bank, "ecb_yc_forward"))
        if s is not None:
            x = s.extra
            ks = sorted(x["curve"])
            for lbl, mm in (("last-2026", m26), ("last-2027", m27)):
                t = (mm.decision - x["asof"]).days / 30.44
                v = float(np.interp(t, ks, [x["curve"][k] for k in ks])) - sp
                print(f"      [ecb_yc_forward AAA govt fwd @{t:.1f}m] {lbl} {mm.decision}: {v:.3f}  ({(v - c['policy']) * 100:+.0f}bp)  [govt curve, not OIS: proxy]")
    # two sources on one currency -> difference at the last 2026 meeting
    print("\n[A3] CROSS-SOURCE DIFFERENCE at the last 2026 meeting (same currency, two sources)")
    us_m = next((m for m in official_meetings("USD") if m.decision.year == 2026 and m.decision == max(x.decision for x in official_meetings("USD") if x.decision.year == 2026)), None)
    mpt, tr = got.get(("USD", "atlantafed_mpt")), got.get(("USD", "treasury_par_curve"))
    if us_m and mpt and tr:
        eff = us_m.decision + timedelta(days=1)
        win, _ = pick_window(a3_windows("atlantafed_mpt", "USD", mpt.extra), eff)
        cv = tr.extra["curve"]
        f = ((1 + cv["6 Mo"] / 100 * 0.5) / (1 + cv["3 Mo"] / 100 * 0.25) - 1) / 0.25 * 100
        if win:
            print(f"  USD: Atlanta MPT window {win['start']}: {win['rate']:.3f}  vs Treasury 3M->6M bill forward: {f:.3f}  -> diff {(f - win['rate']) * 100:+.0f}bp "
                  "(bills vs SOFR basis + bond-equivalent conventions)")
    cad = got.get(("CAD", "mx_corra_expectations"))
    if cad:
        c = ctx.get("CAD", {"policy": 2.25, "spread_bp": 4.0})
        cm = max((m for m in official_meetings("CAD") if m.decision.year == 2026), key=lambda m: m.decision)
        coa = {(cn["month"].year, cn["month"].month): cn["implied"] for cn in cad.extra["contracts"]["COA"]}
        ch = implied_1m_chain(coa, [m.decision for m in official_meetings("CAD") if m.decision > today], c["policy"] + (c["spread_bp"] or 0) / 100)
        coa_v = next((v for e, v, _ in ch if e == cm.decision and v is not None), None)
        win, _ = pick_window(a3_windows("mx_corra_expectations", "CAD", cad.extra), cm.decision + timedelta(days=EFF_LAG_DAYS["CAD"]))
        if coa_v is not None and win:
            print(f"  CAD: COA 1M chain post-{cm.decision}: {coa_v:.3f}  vs CRA window {win['start']}: {win['rate']:.3f}  -> diff {(win['rate'] - coa_v) * 100:+.0f}bp "
                  "(different windows; CRA also spans the Jan-27 meeting)")


def report_a3_criteria(probes: list[Probe]) -> None:
    print("\n[A3] PRE-REGISTERED CRITERIA  1 keyless | 2 EOD lag<=1bd | 3 horizon>=last-2027 meeting | 4 hist>=30bd | 5 deterministic | 6 licence")
    for p in probes:
        s = p.series
        if s is None or "criteria" not in s.extra:
            continue
        c = s.extra["criteria"]
        row = " ".join("Y" if v else "N" for v in c.values())
        print(f"  {p.bank} {p.source:24} {row}   -> {p.status:7} {','.join(p.flags)}  horizon_end={s.extra['horizon_end']} need>={s.extra['need']}")
        print(f"        licence[{s.extra['license'][0]}]: {s.extra['license'][1]}")


def report_matrix(allp: list[Probe], today: date, banks: list[str]) -> None:
    print("\n" + "=" * 96)
    print("COVERAGE MATRIX (probe-derived; analyst notes in docs/spikes/cb_sources_numeric.md)")
    print("=" * 96)
    cell = {}
    for b in banks:
        ps = [p for p in allp if p.bank == b]
        # A1
        prim = next((p for p in ps if p.source == A1_PRIMARY[b] and p.section == "a1"), None)
        off = next((p for p in ps if p.section == "a1" and p.series and p.source not in ("bis_cbpol", "boj_timeseries_api", "rbnz_site")
                    and p.source == {"USD": "fred_target", "EUR": "ecb_data_portal", "GBP": "boe_iadb", "CAD": "boc_valet",
                                     "AUD": "rba_f1", "CHF": "snb_data_portal"}.get(b, "")), None)
        if off is not None:
            lag = off.series.lag_bd(today)
            cell[(b, "a1")] = ("OK", "") if lag <= MAX_LAG_BD else ("PARTIAL", "STALE")
        elif b == "JPY":
            cell[(b, "a1")] = ("PARTIAL", "NONE")     # API has call rate + BLR, guideline only in PDF/BIS(lag)
        else:
            cell[(b, "a1")] = ("FAIL", next((p.code for p in ps if p.source == "rbnz_site"), "NONE"))
        # A2
        f = next((p for p in ps if p.section == "a2" and p.source == "fed_sep_html"), None)
        n = next((p for p in ps if p.section == "a2" and p.source == "no_own_path_evidence"), None)
        if f is not None:
            cell[(b, "a2")] = ("OK", "") if f.series else ("FAIL", f.code)
        elif b == "NZD":
            cell[(b, "a2")] = ("FAIL", "BOTWALL")
        else:
            cell[(b, "a2")] = ("OK", "N/A-no-path" + ("" if n is not None and n.series else "*"))
        # A3
        a3 = [p for p in ps if p.section == "a3" and p.series and p.source not in {t[1] for t in BLOCKED_TARGETS}]
        best = min(a3, key=lambda p: (p.status != "OK", len(p.flags))) if a3 else None
        if best is None:
            code = next((p.code for p in ps if p.section == "a3" and p.code), "NONE")
            cell[(b, "a3")] = ("FAIL", code)
        else:
            cell[(b, "a3")] = (best.status, "+".join(f for f in best.flags if f.isupper() and not f.startswith("HIST")) or best.code)
        # A4
        ff = next((p for p in ps if p.section == "a4" and p.series), None)
        if ff is not None:
            d = ff.series.extra["decisions"]
            full = sum(1 for x in d if x["n_rows"] and pd.notna(x["actual"]) and pd.notna(x["forecast"]) and pd.notna(x["previous"]))
            zero = any(x["zero_rows"] for x in d)
            cell[(b, "a4")] = ("OK", "") if full == len(d) and not zero else (("PARTIAL", f"{full}/{len(d)}" + (" zero-row" if zero else "")) if full else ("FAIL", "NONE"))
        # A5
        cal = next((p for p in ps if p.section == "a5" and p.series and p.series.kind == "calendar"), None)
        q = next((p for p in ps if p.section == "a5" and p.series and p.series.kind == "quiet"), None)
        if cal is not None:
            ms = cal.series.extra["meetings"]
            n26 = sum(1 for m in ms if m.decision.year == 2026)
            n27 = sum(1 for m in ms if m.decision.year == 2027)
            manual = any(m.basis == "manual" for m in ms)
            ffb = any(m.basis.startswith("ff") for m in official_meetings(b))
            full = n26 >= 7 and n27 >= 7 and q is not None and not manual
            cell[(b, "a5")] = ("OK", "") if full and not ffb and n26 == 8 and n27 == 8 else \
                ("PARTIAL", ("manual " if manual else "") + f"{n26}+{n27} mtgs" + ("" if q else " no-quiet-rule"))
    print(f"  {'':5}{'A1 policy+4':<17}{'A2 own path':<19}{'A3 implied':<38}{'A4 FF consensus':<22}{'A5 calendar'}")
    for b in banks:
        row = []
        for sec in SECTIONS:
            st, code = cell.get((b, sec), ("-", ""))
            row.append(f"{st}{'/' + code if code else ''}")
        print(f"  {b:5}{row[0]:<17}{row[1]:<19}{row[2]:<38}{row[3]:<22}{row[4]}")


REPORTERS = {"a1": report_a1, "a2": report_a2, "a3": report_a3, "a4": report_a4, "a5": report_a5,
             "b1": report_b1, "b2": report_b2, "b3": report_b3, "b4": report_b4, "b5": report_b5, "b6": report_b6, "b7": report_b7, "b8": report_b8, "b9": report_b9,
             "a3bis": report_a3bis, "e": report_e}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Read-only Central Banks numeric-source + calendar spike.")
    ap.add_argument("--bank", action="append", choices=list(BANKS))
    ap.add_argument("--section", action="append", choices=list(SECTIONS))
    ap.add_argument("--json-out", help="optional: dump probe verdicts to this path")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.WARNING, format="%(asctime)s %(levelname)s [%(name)s] %(message)s")
    today = date.today()
    banks = args.bank or list(BANKS)
    print("=" * 96)
    print(f"CENTRAL BANKS DATA SPIKE — numeric sources + calendar · {today.isoformat()} · banks={','.join(banks)}")
    print("=" * 96)
    allp: list[Probe] = []
    for sec in args.section or list(SECTIONS):
        ps = probe_section(sec, banks, today)
        allp += ps
        if sec == "a3":
            assess_a3(ps, today)
        REPORTERS[sec](ps, today)
        if sec == "a3":
            report_a3_criteria(ps)
            report_a3_implied(ps, today)
            _print_negative(ps)
    if set(CORE_SECTIONS) <= set(args.section or SECTIONS):
        report_matrix(allp, today, banks)
    if args.json_out:
        Path(args.json_out).write_text(json.dumps(
            [{"bank": p.bank, "section": p.section, "source": p.source, "status": p.status, "code": p.code,
              "note": p.note, "cdn": p.cdn} for p in allp], indent=1, default=str))
    print("\n[STOP] Spike only — no parquet, no adapters, no UI, no pipeline changes.\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
